"""Bounded, deterministic parsers for the standalone extraction baseline."""

from __future__ import annotations

import csv
import hashlib
import json
import posixpath
import re
import time
import zipfile
from collections import Counter
from dataclasses import asdict, dataclass, replace
from pathlib import Path, PurePosixPath
from typing import Any, BinaryIO
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from openpyxl.styles.numbers import BUILTIN_FORMATS
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
from pypdf import PdfReader

from doclens_ai.extraction.models import (
    DocumentBlock,
    DocumentFormat,
    ErrorCategory,
    ExtractionResult,
    ExtractionStatus,
    ExtractionWarning,
    NoteReference,
    SourceLocation,
    SpreadsheetCell,
    SpreadsheetSheet,
    SpreadsheetTable,
    TableBlock,
    TableCell,
    TableRow,
    TextBlock,
    WarningCode,
    WarningSeverity,
)

SCHEMA_VERSION = "1.1"
_SUPPORTED_SUFFIXES = {
    ".pdf": DocumentFormat.PDF,
    ".docx": DocumentFormat.DOCX,
    ".xlsx": DocumentFormat.XLSX,
}
_OFFICE_CONTENT_TYPES = {
    DocumentFormat.DOCX: "word/document.xml",
    DocumentFormat.XLSX: "xl/workbook.xml",
}
_SPREADSHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_CELL_REFERENCE = re.compile(r"^\$?([A-Za-z]{1,3})\$?([1-9][0-9]{0,6})$")


@dataclass(frozen=True, slots=True)
class ExtractionLimits:
    """Hard resource ceilings applied before and during parsing."""

    max_files: int = 10_000
    max_file_bytes: int = 50 * 1024 * 1024
    max_office_entries: int = 20_000
    max_office_entry_bytes: int = 100 * 1024 * 1024
    max_office_uncompressed_bytes: int = 500 * 1024 * 1024
    max_office_compression_ratio: float = 200.0
    max_pdf_pages: int = 2_000
    min_pdf_non_whitespace_chars_per_page: int = 20
    max_extracted_characters: int = 5_000_000
    max_document_blocks: int = 100_000
    max_table_cells: int = 1_000_000
    max_xlsx_sheets: int = 256
    max_xlsx_stored_cells: int = 1_000_000
    max_xlsx_merged_ranges: int = 100_000
    max_xlsx_tables: int = 10_000


class ExtractionConfigurationError(ValueError):
    """Raised for invalid corpus-runner configuration."""


class _ExtractionFailure(Exception):
    def __init__(self, category: ErrorCategory) -> None:
        super().__init__(category.value)
        self.category = category


def extract_file(
    source: str | Path,
    *,
    limits: ExtractionLimits | None = None,
) -> ExtractionResult:
    """Extract one document without allowing failures to escape file isolation."""

    active_limits = limits or ExtractionLimits()
    path = Path(source)
    document_format = _SUPPORTED_SUFFIXES.get(path.suffix.lower())
    source_size: int | None = None
    source_sha256: str | None = None

    try:
        if path.is_symlink():
            raise _ExtractionFailure(ErrorCategory.SYMLINK_NOT_ALLOWED)
        if not path.is_file():
            raise _ExtractionFailure(ErrorCategory.FILE_UNREADABLE)

        source_size = path.stat().st_size
        if source_size == 0:
            raise _ExtractionFailure(ErrorCategory.EMPTY_FILE)
        if source_size > active_limits.max_file_bytes:
            raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)
        source_sha256 = _sha256(path)
        if document_format is None:
            raise _ExtractionFailure(ErrorCategory.UNSUPPORTED_FORMAT)

        if document_format is DocumentFormat.PDF:
            return _extract_pdf(path, source_sha256, source_size, active_limits)
        if document_format is DocumentFormat.DOCX:
            _validate_office_archive(path, document_format, active_limits)
            return _extract_docx(path, source_sha256, source_size, active_limits)

        _validate_office_archive(path, document_format, active_limits)
        return _extract_xlsx(path, source_sha256, source_size, active_limits)
    except _ExtractionFailure as exc:
        return _failure_result(
            document_format=document_format,
            source_sha256=source_sha256,
            source_size=source_size,
            category=exc.category,
        )
    except (PermissionError, OSError):
        return _failure_result(
            document_format=document_format,
            source_sha256=source_sha256,
            source_size=source_size,
            category=ErrorCategory.FILE_UNREADABLE,
        )
    except (ElementTree.ParseError, KeyError, ValueError, zipfile.BadZipFile):
        return _failure_result(
            document_format=document_format,
            source_sha256=source_sha256,
            source_size=source_size,
            category=ErrorCategory.CORRUPT_DOCUMENT,
        )
    except ImportError:
        return _failure_result(
            document_format=document_format,
            source_sha256=source_sha256,
            source_size=source_size,
            category=ErrorCategory.MISSING_DEPENDENCY,
        )
    except Exception:  # noqa: BLE001 - isolation is an explicit public guarantee
        return _failure_result(
            document_format=document_format,
            source_sha256=source_sha256,
            source_size=source_size,
            category=ErrorCategory.EXTRACTION_FAILED,
        )


def extract_corpus(
    input_dir: str | Path,
    output_dir: str | Path,
    *,
    limits: ExtractionLimits | None = None,
    measure_latency: bool = True,
) -> dict[str, Any]:
    """Run isolated extraction and persist only privacy-safe aggregate metadata."""

    active_limits = limits or ExtractionLimits()
    source_root = Path(input_dir).resolve()
    result_root = Path(output_dir).resolve()
    if not source_root.is_dir():
        raise ExtractionConfigurationError("input directory must exist")
    if result_root == source_root or source_root in result_root.parents:
        raise ExtractionConfigurationError("output directory must be outside input directory")

    files: list[Path] = []
    for path in source_root.rglob("*"):
        if path.is_file():
            files.append(path)
            if len(files) > active_limits.max_files:
                raise ExtractionConfigurationError("input directory exceeds the file limit")
    files.sort(key=lambda item: item.relative_to(source_root).as_posix().casefold())
    result_root.mkdir(parents=True, exist_ok=True)

    records: list[dict[str, Any]] = []
    for index, path in enumerate(files, start=1):
        started_ns = time.perf_counter_ns()
        result = extract_file(path, limits=active_limits)
        elapsed_ms = (time.perf_counter_ns() - started_ns) / 1_000_000
        record = _safe_summary(result, document_id=f"document-{index:06d}")
        record["latency_ms"] = round(elapsed_ms, 3) if measure_latency else None
        records.append(record)

    deterministic_records = [
        {key: value for key, value in record.items() if key != "latency_ms"}
        for record in records
    ]
    fingerprint = hashlib.sha256(
        json.dumps(deterministic_records, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()
    aggregate = _aggregate_summaries(
        records,
        fingerprint=fingerprint,
        limits=active_limits,
    )
    payload = {
        "schema_version": SCHEMA_VERSION,
        "aggregate": aggregate,
        "documents": records,
    }

    (result_root / "metrics.json").write_text(
        json.dumps(payload, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    _write_manifest(result_root / "manifest.csv", records)
    _write_report(result_root / "report.md", aggregate)
    return payload


def _extract_pdf(
    path: Path,
    source_sha256: str,
    source_size: int,
    limits: ExtractionLimits,
) -> ExtractionResult:
    try:
        reader = PdfReader(path, strict=False)
    except Exception as exc:  # noqa: BLE001 - parser exceptions vary by input
        raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT) from exc
    if reader.is_encrypted:
        raise _ExtractionFailure(ErrorCategory.ENCRYPTED_DOCUMENT)
    if len(reader.pages) > limits.max_pdf_pages:
        raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)

    blocks: list[DocumentBlock] = []
    warnings = [
        ExtractionWarning(
            code=WarningCode.PDF_TABLE_STRUCTURE_NOT_PRESERVED,
            severity=WarningSeverity.INFO,
        )
    ]
    character_count = 0
    for page_index, page in enumerate(reader.pages, start=1):
        source = SourceLocation(page_number=page_index, block_index=page_index - 1)
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - preserve other pages on malformed content
            text = ""
            warnings.append(
                ExtractionWarning(
                    code=WarningCode.PAGE_EXTRACTION_FAILED,
                    severity=WarningSeverity.REVIEW,
                    source=source,
                )
            )
        character_count += len(text)
        if character_count > limits.max_extracted_characters:
            raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)
        blocks.append(TextBlock(order=page_index - 1, text=text, source=source))
        if len("".join(text.split())) < limits.min_pdf_non_whitespace_chars_per_page:
            warnings.append(
                ExtractionWarning(
                    code=WarningCode.INSUFFICIENT_TEXT,
                    severity=WarningSeverity.REVIEW,
                    source=source,
                )
            )

    return _success_result(
        document_format=DocumentFormat.PDF,
        source_sha256=source_sha256,
        source_size=source_size,
        blocks=tuple(blocks),
        warnings=tuple(warnings),
    )


def _extract_docx(
    path: Path,
    source_sha256: str,
    source_size: int,
    limits: ExtractionLimits,
) -> ExtractionResult:
    try:
        document = Document(path)
    except Exception as exc:  # noqa: BLE001 - python-docx exceptions are not stable
        raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT) from exc

    blocks: list[DocumentBlock] = []
    warnings: list[ExtractionWarning] = []
    table_index = 0
    total_characters = 0
    total_table_grid_positions = 0
    unsupported_body_content_count = 0
    main_part_uri = str(document.part.partname)
    body_root_path = (list(document.element).index(document.element.body),)
    for child_index, child in enumerate(document.element.body.iterchildren()):
        if isinstance(child, CT_P):
            text = _docx_paragraph_text(child)
            source = SourceLocation(
                block_index=len(blocks),
                part_uri=main_part_uri,
                xml_path=body_root_path + (child_index,),
            )
            blocks.append(TextBlock(order=len(blocks), text=text, source=source))
            total_characters += len(text)
        elif isinstance(child, CT_Tbl):
            table = Table(child, document._body)  # noqa: SLF001
            table_block, grid_count, character_count, nested_count = _docx_table_block(
                table,
                order=len(blocks),
                table_index=table_index,
                max_grid_positions=limits.max_table_cells - total_table_grid_positions,
                part_uri=main_part_uri,
                xml_path=body_root_path + (child_index,),
            )
            blocks.append(table_block)
            total_table_grid_positions += grid_count
            total_characters += character_count
            table_index += 1
            if nested_count:
                warnings.append(
                    ExtractionWarning(
                        code=WarningCode.NESTED_TABLES_NOT_EXTRACTED,
                        severity=WarningSeverity.REVIEW,
                        count=nested_count,
                        source=table_block.source,
                    )
                )
        elif child.tag.rsplit("}", maxsplit=1)[-1] != "sectPr":
            unsupported_body_content_count += 1

        if len(blocks) > limits.max_document_blocks:
            raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)
        if total_characters > limits.max_extracted_characters:
            raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)
        if total_table_grid_positions > limits.max_table_cells:
            raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)

    if unsupported_body_content_count:
        warnings.append(
            ExtractionWarning(
                WarningCode.UNSUPPORTED_BODY_CONTENT_NOT_EXTRACTED,
                WarningSeverity.REVIEW,
                unsupported_body_content_count,
            )
        )
    warnings.extend(_docx_unsupported_content_warnings(path, document))
    # Reuse the bounded relationship-driven reader. Its unresolved independent
    # locator review remains explicit; extraction does not promote this evidence.
    from doclens_ai.extraction.docx_footnote_candidate import extract_footnotes

    ancillary = extract_footnotes(path)
    if (
        ancillary.status is ExtractionStatus.FAILED
        or ancillary.source_sha256 != source_sha256
        or ancillary.source_size_bytes != source_size
        or ancillary.main_part_uri != main_part_uri
    ):
        raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
    references = tuple(
        NoteReference(
            SourceLocation(
                block_index=ref.baseline_block_index,
                part_uri=main_part_uri,
                xml_path=body_root_path + ref.body_xml_path,
                note_id=ref.note_id,
                relationship_id=ancillary.footnote_relationship_id,
            ),
            ancillary.footnote_part_uri,
            ref.note_id,
        )
        for ref in ancillary.references
    )
    for note in ancillary.notes:
        for paragraph in note.paragraphs:
            blocks.append(
                TextBlock(
                    order=len(blocks),
                    text=paragraph.text,
                    source=SourceLocation(
                        part_uri=note.part_uri,
                        xml_path=(note.note_element_index, paragraph.child_index),
                        note_id=note.note_id,
                        relationship_id=ancillary.footnote_relationship_id,
                    ),
                )
            )
            total_characters += len(paragraph.text)
    if ancillary.notes or ancillary.issues:
        warnings.append(
            ExtractionWarning(WarningCode.FOOTNOTE_PROVENANCE_UNREVIEWED, WarningSeverity.REVIEW)
        )
    ancillary_blocks = _docx_headers_and_footers(
        path,
        document,
        start_order=len(blocks),
        limits=replace(limits, max_table_cells=limits.max_table_cells - total_table_grid_positions),
    )
    if ancillary_blocks:
        blocks.extend(ancillary_blocks)
        warnings = [
            warning
            for warning in warnings
            if warning.code
            not in (WarningCode.HEADERS_NOT_EXTRACTED, WarningCode.FOOTERS_NOT_EXTRACTED)
        ]
        warnings.append(
            ExtractionWarning(WarningCode.ANCILLARY_PROVENANCE_UNREVIEWED, WarningSeverity.REVIEW)
        )
        total_characters += sum(
            len(block.text)
            if isinstance(block, TextBlock)
            else sum(len(cell.text) for row in block.rows for cell in row.cells)
            for block in ancillary_blocks
        )
    if (
        len(blocks) > limits.max_document_blocks
        or total_characters > limits.max_extracted_characters
    ):
        raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)
    return _success_result(
        document_format=DocumentFormat.DOCX,
        source_sha256=source_sha256,
        source_size=source_size,
        blocks=tuple(blocks),
        warnings=tuple(warnings),
        note_references=references,
    )


def _docx_table_block(
    table: Table,
    *,
    order: int,
    table_index: int,
    max_grid_positions: int,
    part_uri: str | None = None,
    xml_path: tuple[int, ...] | None = None,
) -> tuple[TableBlock, int, int, int]:
    positions_by_cell: dict[int, list[tuple[int, int, Any]]] = {}
    grid_position_count = 0
    previous_merges: set[tuple[int, int]] = set()
    for row_index, row in enumerate(table.rows):
        before, after = row.grid_cols_before, row.grid_cols_after
        if before < 0 or after < 0:
            raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
        # python-docx expands gridSpan and resolves vertical continuations in
        # row.cells. Validate declared topology and budgets before that allocation;
        # an orphan continuation would otherwise silently reuse unrelated text.
        column = before
        current_merges: set[tuple[int, int]] = set()
        for tc in row._tr.tc_lst:  # noqa: SLF001 - validate original OOXML topology
            span = tc.grid_span
            if span < 1:
                raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
            if grid_position_count + column - before + span > max_grid_positions:
                raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)
            if tc.vMerge not in (None, "restart", "continue"):
                raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
            if tc.vMerge == "continue" and (column, span) not in previous_merges:
                raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
            if tc.vMerge is not None:
                current_merges.add((column, span))
            column += span
        if column + after != len(table.columns):
            raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
        previous_merges = current_merges
        row_cells = row.cells
        if before + len(row_cells) + after != len(table.columns):
            raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
        for column_index, cell in enumerate(row_cells, start=before):
            grid_position_count += 1
            if grid_position_count > max_grid_positions:
                raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)
            positions_by_cell.setdefault(id(cell._tc), []).append(  # noqa: SLF001
                (row_index, column_index, cell)
            )

    logical_cells: list[TableCell] = []
    total_characters = 0
    nested_count = 0
    for positions in positions_by_cell.values():
        positions.sort(key=lambda item: (item[0], item[1]))
        row_indexes = [item[0] for item in positions]
        column_indexes = [item[1] for item in positions]
        row_index = min(row_indexes)
        column_index = min(column_indexes)
        cell = positions[0][2]
        text = "\n".join(_docx_paragraph_text(p._p) for p in cell.paragraphs)
        total_characters += len(text)
        nested_count += len(cell.tables)
        source = SourceLocation(
            block_index=order,
            table_index=table_index,
            row_index=row_index,
            column_index=column_index,
            part_uri=part_uri,
            xml_path=(
                xml_path
                + (
                    list(table._tbl).index(cell._tc.getparent()),
                    list(cell._tc.getparent()).index(cell._tc),
                )
                if xml_path is not None
                else None
            ),
        )
        logical_cells.append(
            TableCell(
                row_index=row_index,
                column_index=column_index,
                text=text,
                source=source,
                row_span=max(row_indexes) - row_index + 1,
                column_span=max(column_indexes) - column_index + 1,
            )
        )

    logical_cells.sort(key=lambda item: (item.row_index, item.column_index))
    rows = tuple(
        TableRow(
            row_index=row_index,
            cells=tuple(cell for cell in logical_cells if cell.row_index == row_index),
        )
        for row_index in sorted({cell.row_index for cell in logical_cells})
    )

    source = SourceLocation(
        block_index=order, table_index=table_index, part_uri=part_uri, xml_path=xml_path
    )
    return (
        TableBlock(
            order=order,
            table_index=table_index,
            rows=rows,
            source=source,
            header_row_indices=tuple(
                i
                for i, row in enumerate(table.rows)
                if any(
                    node.get(f"{{{_WORD_NS}}}val", "true") not in ("0", "false", "off")
                    for node in row._tr.findall(f"{{{_WORD_NS}}}trPr/{{{_WORD_NS}}}tblHeader")
                )
            ),
        ),
        grid_position_count,
        total_characters,
        nested_count,
    )


def _docx_headers_and_footers(
    path: Path, document: Any, *, start_order: int, limits: ExtractionLimits
) -> tuple[DocumentBlock, ...]:
    """Keep shared parts once, in their own XML order, without invented pages."""
    blocks: list[DocumentBlock] = []
    seen = set()
    total_grid = 0
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        for section in document.sections:
            for container in (
                section.header,
                section.first_page_header,
                section.even_page_header,
                section.footer,
                section.first_page_footer,
                section.even_page_footer,
            ):
                part_uri = str(container.part.partname)
                if part_uri in seen or part_uri.lstrip("/") not in names:
                    continue
                seen.add(part_uri)
                if not _container_has_text(container):
                    continue
                table_index = 0
                for child_index, child in enumerate(container._element):
                    order = start_order + len(blocks)
                    source = SourceLocation(part_uri=part_uri, xml_path=(child_index,))
                    if isinstance(child, CT_P):
                        blocks.append(TextBlock(order, _docx_paragraph_text(child), source))
                    elif isinstance(child, CT_Tbl):
                        table, grid, _, _ = _docx_table_block(
                            Table(child, container),
                            order=order,
                            table_index=table_index,
                            max_grid_positions=limits.max_table_cells - total_grid,
                            part_uri=part_uri,
                            xml_path=(child_index,),
                        )
                        # Body ordinals do not apply to a separately addressed part.
                        table = replace(
                            table,
                            source=replace(table.source, block_index=None),
                            rows=tuple(
                                replace(
                                    row,
                                    cells=tuple(
                                        replace(cell, source=replace(cell.source, block_index=None))
                                        for cell in row.cells
                                    ),
                                )
                                for row in table.rows
                            ),
                        )
                        blocks.append(table)
                        table_index += 1
                        total_grid += grid
                    if len(blocks) > limits.max_document_blocks:
                        raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)
    return tuple(blocks)


def _docx_unsupported_content_warnings(
    path: Path,
    document: Any,
) -> list[ExtractionWarning]:
    warnings: list[ExtractionWarning] = []
    header_parts: set[str] = set()
    footer_parts: set[str] = set()
    header_count = 0
    footer_count = 0
    for section in document.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            part_name = str(header.part.partname)
            if part_name not in header_parts and _container_has_text(header):
                header_parts.add(part_name)
                header_count += 1
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            part_name = str(footer.part.partname)
            if part_name not in footer_parts and _container_has_text(footer):
                footer_parts.add(part_name)
                footer_count += 1
    if header_count:
        warnings.append(
            ExtractionWarning(
                WarningCode.HEADERS_NOT_EXTRACTED,
                WarningSeverity.REVIEW,
                header_count,
            )
        )
    if footer_count:
        warnings.append(
            ExtractionWarning(
                WarningCode.FOOTERS_NOT_EXTRACTED,
                WarningSeverity.REVIEW,
                footer_count,
            )
        )

    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        for part_name, warning_code in (
            ("word/endnotes.xml", WarningCode.ENDNOTES_NOT_EXTRACTED),
            ("word/comments.xml", WarningCode.COMMENTS_NOT_EXTRACTED),
        ):
            if part_name in names and _xml_has_text(archive.read(part_name)):
                warnings.append(ExtractionWarning(warning_code, WarningSeverity.REVIEW))

        root = _parse_xml(archive.read(str(document.part.partname).lstrip("/")))
        for tag, code in (
            ("endnoteReference", WarningCode.ENDNOTES_NOT_EXTRACTED),
            ("commentReference", WarningCode.COMMENTS_NOT_EXTRACTED),
        ):
            references = root.findall(f".//{{{_WORD_NS}}}{tag}")
            if references and not any(warning.code is code for warning in warnings):
                warnings.append(ExtractionWarning(code, WarningSeverity.REVIEW, len(references)))
        drawing_count = len(root.findall(f".//{{{_WORD_NS}}}drawing")) + len(
            root.findall(f".//{{{_WORD_NS}}}pict")
        )
        text_box_count = len(root.findall(f".//{{{_WORD_NS}}}txbxContent"))
        unsupported_inline = sum(
            len(root.findall(f".//{{{_WORD_NS}}}{tag}"))
            for tag in ("ins", "del", "sdt", "fldSimple", "fldChar", "sym", "altChunk")
        )
        if unsupported_inline:
            warnings.append(
                ExtractionWarning(
                    WarningCode.UNSUPPORTED_BODY_CONTENT_NOT_EXTRACTED,
                    WarningSeverity.REVIEW,
                    unsupported_inline,
                )
            )
        if drawing_count:
            warnings.append(
                ExtractionWarning(
                    WarningCode.DRAWINGS_NOT_EXTRACTED,
                    WarningSeverity.REVIEW,
                    drawing_count,
                )
            )
        if text_box_count:
            warnings.append(
                ExtractionWarning(
                    WarningCode.TEXT_BOXES_NOT_EXTRACTED,
                    WarningSeverity.REVIEW,
                    text_box_count,
                )
            )
    return warnings


def _docx_paragraph_text(paragraph: Any) -> str:
    """Render supported OOXML text controls without changing Unicode characters."""
    controls = {
        "tab": "\t",
        "ptab": "\t",
        "cr": "\n",
        "noBreakHyphen": "\u2011",
        "softHyphen": "\u00ad",
    }
    chunks = []
    stack = [paragraph]
    while stack:
        node = stack.pop()
        if node.tag == f"{{{_WORD_NS}}}t":
            chunks.append(node.text or "")
        elif node.tag in {f"{{{_WORD_NS}}}{tag}" for tag in controls}:
            chunks.append(controls[node.tag.rsplit("}", 1)[-1]])
        elif node.tag == f"{{{_WORD_NS}}}br":
            if node.get(f"{{{_WORD_NS}}}type", "textWrapping") == "textWrapping":
                chunks.append("\n")
        elif node.tag in {f"{{{_WORD_NS}}}{tag}" for tag in ("p", "r", "hyperlink")}:
            stack.extend(reversed(list(node)))
    return "".join(chunks)


def _extract_xlsx(
    path: Path,
    source_sha256: str,
    source_size: int,
    limits: ExtractionLimits,
) -> ExtractionResult:
    warnings: list[ExtractionWarning] = []
    sheets: list[SpreadsheetSheet] = []
    stored_cell_count = 0
    character_count = 0
    table_count = 0
    merged_range_count = 0

    with zipfile.ZipFile(path) as archive:
        shared_strings = _xlsx_shared_strings(archive, limits)
        workbook = _parse_xml(archive.read("xl/workbook.xml"))
        workbook_relationships = _read_relationships(archive, "xl/_rels/workbook.xml.rels")
        styles = _xlsx_number_formats(archive, workbook_relationships)
        properties = workbook.find(f"{{{_SPREADSHEET_NS}}}workbookPr")
        if properties is not None and properties.get("date1904", "0") not in (
            "0",
            "1",
            "true",
            "false",
        ):
            raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
        date_system = (
            "1904"
            if properties is not None and properties.get("date1904") in ("1", "true")
            else "1900"
        )
        sheet_elements = workbook.findall(f".//{{{_SPREADSHEET_NS}}}sheet")
        if len(sheet_elements) > limits.max_xlsx_sheets:
            raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)

        for sheet_index, sheet_element in enumerate(sheet_elements):
            missing_formula_definition_count = 0
            missing_formula_cache_count = 0
            relationship_id = sheet_element.get(f"{{{_OFFICE_REL_NS}}}id")
            if not relationship_id or relationship_id not in workbook_relationships:
                raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
            relationship_type = workbook_relationships[relationship_id][1]
            sheet_path = _resolve_relationship_target(
                "xl/workbook.xml", workbook_relationships[relationship_id][0]
            )
            sheet_root = _parse_xml(archive.read(sheet_path))
            if relationship_type in {
                f"{_OFFICE_REL_NS}/dialogsheet",
                f"{_OFFICE_REL_NS}/chartsheet",
            }:
                if sheet_root.tag != f"{{{_SPREADSHEET_NS}}}{relationship_type.rsplit('/', 1)[-1]}":
                    raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
                raise _ExtractionFailure(ErrorCategory.UNSUPPORTED_FORMAT)
            if (
                relationship_type != f"{_OFFICE_REL_NS}/worksheet"
                or sheet_root.tag != f"{{{_SPREADSHEET_NS}}}worksheet"
            ):
                raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
            cell_elements = sheet_root.findall(f".//{{{_SPREADSHEET_NS}}}c")
            stored_cell_count += len(cell_elements)
            if stored_cell_count > limits.max_xlsx_stored_cells:
                raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)

            cells: list[SpreadsheetCell] = []
            seen_coordinates: set[str] = set()
            for cell_element in cell_elements:
                coordinate = cell_element.get("r")
                if not coordinate or not _CELL_REFERENCE.match(coordinate):
                    raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
                row_index, column_index = coordinate_to_tuple(coordinate.replace("$", ""))
                coordinate = coordinate.replace("$", "").upper()
                if coordinate in seen_coordinates:
                    raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
                seen_coordinates.add(coordinate)
                if row_index > 1_048_576 or column_index > 16_384:
                    raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
                formula_element = cell_element.find(f"{{{_SPREADSHEET_NS}}}f")
                value_element = cell_element.find(f"{{{_SPREADSHEET_NS}}}v")
                formula = None
                if formula_element is not None:
                    formula = formula_element.text or None
                    if formula is None:
                        missing_formula_definition_count += 1
                value = _xlsx_cell_value(cell_element, value_element, shared_strings)
                cached_value = value if formula_element is not None else None
                if formula_element is not None and cached_value is None:
                    missing_formula_cache_count += 1
                if formula_element is None and value is None:
                    continue
                style_index = int(cell_element.get("s", "0"))
                if not 0 <= style_index < len(styles):
                    raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
                format_id, number_format = styles[style_index]
                if number_format is None:
                    warnings.append(
                        ExtractionWarning(
                            WarningCode.NUMBER_FORMAT_UNAVAILABLE, WarningSeverity.REVIEW
                        )
                    )
                if cell_element.get("t") == "e":
                    warnings.append(
                        ExtractionWarning(
                            WarningCode.SPREADSHEET_ERROR_VALUE, WarningSeverity.REVIEW
                        )
                    )
                character_count += len(value or "") + len(formula or "") + len(number_format or "")
                if formula_element is not None:
                    character_count += sum(
                        len(key) + len(attribute)
                        for key, attribute in formula_element.attrib.items()
                    )
                if character_count > limits.max_extracted_characters:
                    raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)
                source = SourceLocation(
                    sheet_index=sheet_index,
                    sheet_name=sheet_element.get("name", ""),
                    cell_coordinate=coordinate,
                    row_index=row_index - 1,
                    column_index=column_index - 1,
                )
                cells.append(
                    SpreadsheetCell(
                        coordinate=coordinate,
                        row_index=row_index - 1,
                        column_index=column_index - 1,
                        value=value,
                        formula=formula,
                        cached_value=cached_value,
                        data_type=cell_element.get("t", "n"),
                        source=source,
                        raw_value=value_element.text if value_element is not None else None,
                        number_format=number_format,
                        number_format_id=format_id,
                        style_index=style_index,
                        has_formula=formula_element is not None,
                        formula_attributes=(
                            tuple(sorted(formula_element.attrib.items()))
                            if formula_element is not None
                            else ()
                        ),
                    )
                )

            cells.sort(key=lambda item: (item.row_index, item.column_index, item.coordinate))
            sheet_source = SourceLocation(
                sheet_index=sheet_index,
                sheet_name=sheet_element.get("name", ""),
            )
            if missing_formula_definition_count:
                warnings.append(
                    ExtractionWarning(
                        WarningCode.FORMULA_DEFINITION_UNAVAILABLE,
                        WarningSeverity.REVIEW,
                        missing_formula_definition_count,
                        sheet_source,
                    )
                )
            if missing_formula_cache_count:
                warnings.append(
                    ExtractionWarning(
                        WarningCode.FORMULA_CACHED_VALUE_MISSING,
                        WarningSeverity.REVIEW,
                        missing_formula_cache_count,
                        sheet_source,
                    )
                )
            merged_ranges = tuple(
                element.get("ref", "")
                for element in sheet_root.findall(f".//{{{_SPREADSHEET_NS}}}mergeCell")
            )
            for merged in merged_ranges:
                _xlsx_range(merged)
            merged_range_count += len(merged_ranges)
            if merged_range_count > limits.max_xlsx_merged_ranges:
                raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)

            tables, relationship_warnings = _xlsx_tables_and_warnings(
                archive,
                sheet_path,
                sheet_root,
            )
            table_count += len(tables)
            if table_count > limits.max_xlsx_tables:
                raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)
            warnings.extend(relationship_warnings)
            if not _xlsx_structure_consistent(cells, tables, merged_ranges):
                warnings.append(
                    ExtractionWarning(
                        WarningCode.AMBIGUOUS_SPREADSHEET_STRUCTURE,
                        WarningSeverity.REVIEW,
                        source=sheet_source,
                    )
                )
            else:
                cells = _xlsx_table_memberships(cells, tables)
            sheets.append(
                SpreadsheetSheet(
                    index=sheet_index,
                    name=sheet_element.get("name", ""),
                    state=sheet_element.get("state", "visible"),
                    cells=tuple(cells),
                    merged_ranges=merged_ranges,
                    tables=tables,
                    date_system=date_system,
                )
            )

    return _success_result(
        document_format=DocumentFormat.XLSX,
        source_sha256=source_sha256,
        source_size=source_size,
        sheets=tuple(sheets),
        warnings=tuple(warnings),
    )


def _xlsx_shared_strings(
    archive: zipfile.ZipFile,
    limits: ExtractionLimits,
) -> tuple[str, ...]:
    if "xl/sharedStrings.xml" not in archive.namelist():
        return ()
    root = _parse_xml(archive.read("xl/sharedStrings.xml"))
    values: list[str] = []
    character_count = 0
    for item in root.findall(f"{{{_SPREADSHEET_NS}}}si"):
        value = _xlsx_string(item)
        values.append(value)
        if len(values) > limits.max_xlsx_stored_cells:
            raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)
        character_count += len(value)
        if character_count > limits.max_extracted_characters:
            raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)
    return tuple(values)


def _xlsx_string(container: ElementTree.Element) -> str:
    # Phonetic rPh/t annotations are not the cell's native text.
    return "".join(
        (child.text or "")
        if child.tag == f"{{{_SPREADSHEET_NS}}}t"
        else "".join(node.text or "" for node in child.findall(f"{{{_SPREADSHEET_NS}}}t"))
        for child in container
        if child.tag in (f"{{{_SPREADSHEET_NS}}}t", f"{{{_SPREADSHEET_NS}}}r")
    )


def _xlsx_number_formats(
    archive: zipfile.ZipFile, relationships: dict[str, tuple[str, str]]
) -> tuple[tuple[int, str | None], ...]:
    targets = [target for target, kind in relationships.values() if kind.endswith("/styles")]
    if not targets:
        return ((0, "General"),)
    if len(targets) != 1:
        raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
    root = _parse_xml(archive.read(_resolve_relationship_target("xl/workbook.xml", targets[0])))
    custom: dict[int, str] = {}
    for node in root.findall(f"{{{_SPREADSHEET_NS}}}numFmts/{{{_SPREADSHEET_NS}}}numFmt"):
        key = int(node.get("numFmtId", "-1"))
        value = node.get("formatCode")
        # Explicit definitions also carry localized formats with built-in IDs (<164).
        if key < 0 or key in custom or value is None:
            raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
        custom[key] = value
    formats = []
    for node in root.findall(f"{{{_SPREADSHEET_NS}}}cellXfs/{{{_SPREADSHEET_NS}}}xf"):
        key = int(node.get("numFmtId", "0"))
        value = custom.get(key, BUILTIN_FORMATS.get(key))
        if key < 0:
            raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
        formats.append((key, value))
    if not formats:
        raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
    return tuple(formats)


def _xlsx_range(reference: str) -> tuple[int, int, int, int]:
    bounds = range_boundaries(reference)
    if (
        any(value is None for value in bounds)
        or not 1 <= bounds[0] <= bounds[2] <= 16_384
        or not 1 <= bounds[1] <= bounds[3] <= 1_048_576
    ):
        raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
    return bounds


def _xlsx_structure_consistent(
    cells: list[SpreadsheetCell], tables: tuple[SpreadsheetTable, ...], merges: tuple[str, ...]
) -> bool:
    by_coordinate = {(cell.column_index + 1, cell.row_index + 1): cell for cell in cells}
    for table in tables:
        left, top, _, _ = _xlsx_range(table.reference)
        if table.header_row_count:
            for offset, name in enumerate(table.columns):
                cell = by_coordinate.get((left + offset, top))
                if cell is None or cell.formula is not None or cell.value != name:
                    return False
    # Sweep row boundaries rather than expand large sparse ranges into grid cells.
    regions = [(bounds, "merge") for bounds in map(_xlsx_range, merges)]
    regions += [(_xlsx_range(table.reference), "table") for table in tables]
    regions.sort(key=lambda item: item[0][1])
    active: list[tuple[tuple[int, int, int, int], str]] = []
    for bounds, kind in regions:
        left, top, right, _ = bounds
        active = [(other, tag) for other, tag in active if other[3] >= top]
        if any(left <= other[2] and other[0] <= right and tag == kind for other, tag in active):
            return False
        active.append((bounds, kind))
    merge_regions = sorted(map(_xlsx_range, merges), key=lambda bounds: bounds[1])
    active_merges = []
    index = 0
    for cell in cells:
        row, col = cell.row_index + 1, cell.column_index + 1
        while index < len(merge_regions) and merge_regions[index][1] <= row:
            active_merges.append(merge_regions[index])
            index += 1
        active_merges = [bounds for bounds in active_merges if bounds[3] >= row]
        if any(
            left <= col <= right and (col, row) != (left, top)
            for left, top, right, _ in active_merges
        ):
            return False  # Non-anchor stored content in a merged range is ambiguous.
    return True


def _xlsx_table_memberships(
    cells: list[SpreadsheetCell], tables: tuple[SpreadsheetTable, ...]
) -> list[SpreadsheetCell]:
    regions = sorted(
        ((_xlsx_range(table.reference), index) for index, table in enumerate(tables)),
        key=lambda item: item[0][1],
    )
    active = []
    next_region = 0
    output = []
    for cell in cells:
        row, col = cell.row_index + 1, cell.column_index + 1
        while next_region < len(regions) and regions[next_region][0][1] <= row:
            active.append(regions[next_region])
            next_region += 1
        active = [(bounds, index) for bounds, index in active if bounds[3] >= row]
        membership = [index for (left, _, right, _), index in active if left <= col <= right]
        output.append(
            replace(
                cell, source=replace(cell.source, table_index=membership[0] if membership else None)
            )
        )
    return output


def _xlsx_cell_value(
    cell: ElementTree.Element,
    value_element: ElementTree.Element | None,
    shared_strings: tuple[str, ...],
) -> str | None:
    data_type = cell.get("t")
    if data_type == "inlineStr":
        inline = cell.find(f"{{{_SPREADSHEET_NS}}}is")
        return _xlsx_string(inline) if inline is not None else None
    if value_element is None or value_element.text is None:
        if value_element is not None and data_type == "str":
            return ""
        return None
    raw_value = value_element.text
    if data_type == "s":
        if not re.fullmatch(r"[0-9]+", raw_value) or not 0 <= int(raw_value) < len(shared_strings):
            raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
        return shared_strings[int(raw_value)]
    if data_type == "b":
        if raw_value not in ("0", "1", "true", "false"):
            raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
        return "true" if raw_value in ("1", "true") else "false"
    if data_type not in (None, "n", "d", "str", "e"):
        raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
    if data_type in (None, "n") and not re.fullmatch(
        r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[Ee][+-]?[0-9]+)?", raw_value
    ):
        raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
    return raw_value


def _xlsx_tables_and_warnings(
    archive: zipfile.ZipFile,
    sheet_path: str,
    sheet_root: ElementTree.Element,
) -> tuple[tuple[SpreadsheetTable, ...], list[ExtractionWarning]]:
    relationship_path = posixpath.join(
        posixpath.dirname(sheet_path),
        "_rels",
        f"{posixpath.basename(sheet_path)}.rels",
    )
    relationships = _read_relationships(archive, relationship_path)
    tables: list[SpreadsheetTable] = []
    warnings: list[ExtractionWarning] = []

    for table_part in sheet_root.findall(f".//{{{_SPREADSHEET_NS}}}tablePart"):
        relationship_id = table_part.get(f"{{{_OFFICE_REL_NS}}}id")
        if not relationship_id or relationship_id not in relationships:
            raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
        target, relationship_type = relationships[relationship_id]
        if not relationship_type.endswith("/table"):
            raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
        table_path = _resolve_relationship_target(sheet_path, target)
        table_root = _parse_xml(archive.read(table_path))
        reference = table_root.get("ref")
        if not reference:
            raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
        min_col, min_row, max_col, max_row = _xlsx_range(reference)
        columns = tuple(
            column.get("name", "")
            for column in table_root.findall(f".//{{{_SPREADSHEET_NS}}}tableColumn")
        )
        header_rows = int(table_root.get("headerRowCount", "1"))
        totals_rows = int(table_root.get("totalsRowCount", "0"))
        if (
            len(columns) != max_col - min_col + 1
            or header_rows not in (0, 1)
            or totals_rows not in (0, 1)
            or header_rows + totals_rows > max_row - min_row + 1
        ):
            raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
        tables.append(
            SpreadsheetTable(
                name=table_root.get("name"),
                display_name=table_root.get("displayName"),
                reference=reference,
                columns=columns,
                header_row_count=header_rows,
                totals_row_count=totals_rows,
            )
        )

    relationship_types = [relationship_type for _, relationship_type in relationships.values()]
    drawing_count = sum(item.endswith("/drawing") for item in relationship_types)
    comment_count = sum(item.endswith("/comments") for item in relationship_types)
    if drawing_count:
        warnings.append(
            ExtractionWarning(
                WarningCode.SPREADSHEET_DRAWINGS_NOT_EXTRACTED,
                WarningSeverity.REVIEW,
                drawing_count,
            )
        )
    if comment_count:
        warnings.append(
            ExtractionWarning(
                WarningCode.SPREADSHEET_COMMENTS_NOT_EXTRACTED,
                WarningSeverity.REVIEW,
                comment_count,
            )
        )
    return tuple(tables), warnings


def _read_relationships(
    archive: zipfile.ZipFile,
    relationship_path: str,
) -> dict[str, tuple[str, str]]:
    if relationship_path not in archive.namelist():
        return {}
    root = _parse_xml(archive.read(relationship_path))
    if root.tag != f"{{{_PACKAGE_REL_NS}}}Relationships":
        raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
    result = {}
    seen = set()
    for relationship in root.findall(f"{{{_PACKAGE_REL_NS}}}Relationship"):
        key, target, kind = (relationship.get(k) for k in ("Id", "Target", "Type"))
        if not key or key in seen or not target or not kind:
            raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
        seen.add(key)
        mode = relationship.get("TargetMode", "Internal")
        if mode not in ("Internal", "External"):
            raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
        if mode == "External":
            # These relationships are never followed. Table/sheet lookup fails closed.
            continue
        result[key] = (target, kind)
    return result


def _resolve_relationship_target(source_path: str, target: str) -> str:
    url = urlsplit(target)
    if url.scheme or url.netloc or url.query or url.fragment:
        raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
    target = unquote(url.path)
    if not target or any(c in target for c in ("\\", "\x00", ":")):
        raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
    if target.startswith("/"):
        candidate = target.lstrip("/")
    else:
        candidate = posixpath.normpath(posixpath.join(posixpath.dirname(source_path), target))
    if candidate.startswith("../") or candidate == "..":
        raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
    return candidate


def _validate_office_archive(
    path: Path,
    document_format: DocumentFormat,
    limits: ExtractionLimits,
) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > limits.max_office_entries:
                raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)
            names = {entry.filename for entry in entries}
            if len({name.lower() for name in names}) != len(entries):
                raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
            total_uncompressed = 0
            for entry in entries:
                _validate_archive_member_name(entry.filename)
                if entry.file_size > limits.max_office_entry_bytes:
                    raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)
                total_uncompressed += entry.file_size
                if total_uncompressed > limits.max_office_uncompressed_bytes:
                    raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)
                if entry.file_size and entry.compress_size == 0:
                    raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)
                if entry.compress_size:
                    ratio = entry.file_size / entry.compress_size
                    if ratio > limits.max_office_compression_ratio:
                        raise _ExtractionFailure(ErrorCategory.LIMIT_EXCEEDED)
            if document_format is DocumentFormat.DOCX:
                main_links = [
                    target
                    for target, kind in _read_relationships(archive, "_rels/.rels").values()
                    if kind == f"{_OFFICE_REL_NS}/officeDocument"
                ]
                if (
                    len(main_links) != 1
                    or _resolve_relationship_target("", main_links[0]) not in names
                ):
                    raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
            elif _OFFICE_CONTENT_TYPES[document_format] not in names:
                raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)
            # Validate before python-docx or any other consumer parses package XML.
            for entry in entries:
                if entry.filename.lower().endswith((".xml", ".rels")):
                    _parse_xml(archive.read(entry))
    except zipfile.BadZipFile as exc:
        raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT) from exc


def _validate_archive_member_name(name: str) -> None:
    path = PurePosixPath(name.replace("\\", "/"))
    if path.is_absolute() or ".." in path.parts:
        raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)


def _container_has_text(container: Any) -> bool:
    # Inspect source paragraphs without materializing table grids. This probe
    # runs before the bounded table reader, including for headers and footers.
    return any(
        _docx_paragraph_text(paragraph).strip()
        for paragraph in container._element.iter(f"{{{_WORD_NS}}}p")
    )


def _xml_has_text(data: bytes) -> bool:
    root = _parse_xml(data)
    return any((node.text or "").strip() for node in root.findall(f".//{{{_WORD_NS}}}t"))


class _NoDoctypeTreeBuilder(ElementTree.TreeBuilder):
    def doctype(self, name: str, pubid: str | None, system: str | None) -> None:
        raise _ExtractionFailure(ErrorCategory.CORRUPT_DOCUMENT)


def _parse_xml(data: bytes) -> ElementTree.Element:
    return ElementTree.fromstring(
        data, parser=ElementTree.XMLParser(target=_NoDoctypeTreeBuilder())
    )


def _success_result(
    *,
    document_format: DocumentFormat,
    source_sha256: str,
    source_size: int,
    blocks: tuple[DocumentBlock, ...] = (),
    sheets: tuple[SpreadsheetSheet, ...] = (),
    warnings: tuple[ExtractionWarning, ...] = (),
    note_references: tuple[NoteReference, ...] = (),
) -> ExtractionResult:
    text_values = [block.text for block in blocks if isinstance(block, TextBlock)]
    text_values.extend(
        cell.text
        for block in blocks
        if isinstance(block, TableBlock)
        for row in block.rows
        for cell in row.cells
    )
    text_values.extend(
        value or ""
        for sheet in sheets
        for cell in sheet.cells
        for value in (cell.value, cell.formula, cell.cached_value)
    )
    if any(
        character == "\ufffd" or 0xD800 <= ord(character) <= 0xDFFF
        for value in text_values
        for character in value
    ):
        warnings += (ExtractionWarning(WarningCode.CORRUPTED_UNICODE, WarningSeverity.REVIEW),)
    status = (
        ExtractionStatus.NEEDS_REVIEW
        if any(warning.severity is WarningSeverity.REVIEW for warning in warnings)
        else ExtractionStatus.SUCCESS
    )
    return ExtractionResult(
        schema_version=SCHEMA_VERSION,
        document_format=document_format,
        status=status,
        source_sha256=source_sha256,
        source_size_bytes=source_size,
        blocks=blocks,
        sheets=sheets,
        warnings=warnings,
        note_references=note_references,
    )


def _failure_result(
    *,
    document_format: DocumentFormat | None,
    source_sha256: str | None,
    source_size: int | None,
    category: ErrorCategory,
) -> ExtractionResult:
    return ExtractionResult(
        schema_version=SCHEMA_VERSION,
        document_format=document_format,
        status=ExtractionStatus.FAILED,
        source_sha256=source_sha256,
        source_size_bytes=source_size,
        error_category=category,
    )


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        _update_hash(digest, source)
    return digest.hexdigest()


def _update_hash(digest: Any, source: BinaryIO) -> None:
    for chunk in iter(lambda: source.read(1024 * 1024), b""):
        digest.update(chunk)


def _safe_summary(result: ExtractionResult, *, document_id: str) -> dict[str, Any]:
    table_blocks = [block for block in result.blocks if isinstance(block, TableBlock)]
    text_blocks = [block for block in result.blocks if isinstance(block, TextBlock)]
    warning_codes = sorted({warning.code.value for warning in result.warnings})
    return {
        "document_id": document_id,
        "format": result.document_format.value if result.document_format else None,
        "sha256": result.source_sha256,
        "size_bytes": result.source_size_bytes,
        "status": result.status.value,
        "error_category": result.error_category.value if result.error_category else None,
        "warning_count": sum(warning.count for warning in result.warnings),
        "warning_codes": warning_codes,
        "text_block_count": len(text_blocks),
        "table_count": len(table_blocks) + sum(len(sheet.tables) for sheet in result.sheets),
        "table_cell_count": sum(
            len(row.cells) for table in table_blocks for row in table.rows
        ),
        "sheet_count": len(result.sheets),
        "spreadsheet_cell_count": sum(len(sheet.cells) for sheet in result.sheets),
        "text_character_count": sum(len(block.text) for block in text_blocks),
    }


def _aggregate_summaries(
    records: list[dict[str, Any]],
    *,
    fingerprint: str,
    limits: ExtractionLimits,
) -> dict[str, Any]:
    status_counts = Counter(record["status"] for record in records)
    format_counts = Counter(record["format"] or "unsupported" for record in records)
    error_counts = Counter(
        record["error_category"] for record in records if record["error_category"]
    )
    warning_counts: Counter[str] = Counter()
    for record in records:
        warning_counts.update(record["warning_codes"])
    latencies = [record["latency_ms"] for record in records if record["latency_ms"] is not None]
    return {
        "document_count": len(records),
        "status_counts": dict(sorted(status_counts.items())),
        "format_counts": dict(sorted(format_counts.items())),
        "error_counts": dict(sorted(error_counts.items())),
        "warning_document_counts": dict(sorted(warning_counts.items())),
        "total_size_bytes": sum(record["size_bytes"] or 0 for record in records),
        "total_text_blocks": sum(record["text_block_count"] for record in records),
        "total_tables": sum(record["table_count"] for record in records),
        "total_sheets": sum(record["sheet_count"] for record in records),
        "total_spreadsheet_cells": sum(
            record["spreadsheet_cell_count"] for record in records
        ),
        "latency_ms_total": round(sum(latencies), 3) if latencies else None,
        "deterministic_fingerprint": fingerprint,
        "limits": asdict(limits),
    }


def _write_manifest(path: Path, records: list[dict[str, Any]]) -> None:
    fieldnames = (
        "document_id",
        "format",
        "sha256",
        "size_bytes",
        "status",
        "error_category",
        "warning_count",
        "warning_codes",
        "text_block_count",
        "table_count",
        "table_cell_count",
        "sheet_count",
        "spreadsheet_cell_count",
        "text_character_count",
        "latency_ms",
    )
    with path.open("w", encoding="utf-8", newline="") as destination:
        writer = csv.DictWriter(destination, fieldnames=fieldnames)
        writer.writeheader()
        for record in records:
            row = dict(record)
            row["warning_codes"] = "|".join(record["warning_codes"])
            writer.writerow({key: row[key] for key in fieldnames})


def _write_report(path: Path, aggregate: dict[str, Any]) -> None:
    lines = [
        "# Structured extraction run",
        "",
        "This report contains aggregate metadata only; extracted content is not persisted.",
        "",
        f"- Documents: {aggregate['document_count']}",
        f"- Statuses: {json.dumps(aggregate['status_counts'], sort_keys=True)}",
        f"- Formats: {json.dumps(aggregate['format_counts'], sort_keys=True)}",
        f"- Total size (bytes): {aggregate['total_size_bytes']}",
        f"- Total extraction latency (ms): {aggregate['latency_ms_total']}",
        f"- Deterministic fingerprint: `{aggregate['deterministic_fingerprint']}`",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")

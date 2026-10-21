"""Deterministic offline chunking strategies and privacy-safe evaluation."""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import time
from bisect import bisect_right
from collections import Counter
from collections.abc import Iterable
from dataclasses import asdict, dataclass, fields, replace
from pathlib import Path
from typing import Any

from openpyxl.styles.numbers import is_date_format
from openpyxl.utils.cell import range_boundaries

from doclens_ai.chunking.models import (
    Chunk,
    ChunkingConfig,
    ChunkingOutcome,
    ChunkingReason,
    ChunkingStatus,
    ChunkingStrategy,
    ChunkSourceSpan,
    NormalizationRange,
    TextBlockOffsetMap,
    config_as_dict,
)
from doclens_ai.extraction.models import (
    DocumentFormat,
    ErrorCategory,
    ExtractionResult,
    ExtractionStatus,
    SourceLocation,
    SpreadsheetCell,
    TableBlock,
    TextBlock,
)
from doclens_ai.extraction.parsers import ExtractionLimits, extract_file


@dataclass(frozen=True, slots=True)
class _SourceUnit:
    unit_id: str
    text: str
    source: SourceLocation
    scope_key: str
    group_key: str
    context_ids: tuple[str, ...] = ()
    text_block_map: TextBlockOffsetMap | None = None


@dataclass(frozen=True, slots=True)
class _RenderedSpan:
    unit: _SourceUnit
    rendered_start: int
    rendered_end: int
    source_start: int
    source_end: int


@dataclass(frozen=True, slots=True)
class _RenderedText:
    text: str
    spans: tuple[_RenderedSpan, ...]


@dataclass(frozen=True, slots=True)
class _StructuralGroup:
    scope_key: str
    group_key: str
    units: tuple[_SourceUnit, ...]
    rendered: _RenderedText
    context_units: tuple[_SourceUnit, ...] = ()


class EvaluationConfigurationError(ValueError):
    """Raised when an offline evaluation request is unsafe or invalid."""


class _ChunkingFailure(Exception):
    def __init__(self, reason: ChunkingReason) -> None:
        super().__init__(reason.value)
        self.reason = reason


def chunk_document(
    extraction: ExtractionResult,
    *,
    document_id: str,
    config: ChunkingConfig,
) -> ChunkingOutcome:
    """Chunk one extraction result while containing all document-local failures."""

    config_fingerprint = config.fingerprint()
    try:
        if extraction.status is ExtractionStatus.FAILED:
            return _empty_outcome(
                extraction,
                document_id=document_id,
                config=config,
                status=ChunkingStatus.SKIPPED,
                reason=ChunkingReason.EXTRACTION_FAILED,
            )
        if extraction.status is ExtractionStatus.NEEDS_REVIEW and not config.allow_needs_review:
            return _empty_outcome(
                extraction,
                document_id=document_id,
                config=config,
                status=ChunkingStatus.SKIPPED,
                reason=ChunkingReason.EXTRACTION_NEEDS_REVIEW,
            )
        if (
            not document_id.strip()
            or extraction.source_sha256 is None
            or extraction.document_format is None
        ):
            return _empty_outcome(
                extraction,
                document_id=document_id,
                config=config,
                status=ChunkingStatus.SKIPPED,
                reason=ChunkingReason.MISSING_DOCUMENT_IDENTITY,
            )

        units = _build_source_units(extraction)
        if len({unit.unit_id for unit in units}) != len(units):
            raise _ChunkingFailure(ChunkingReason.CHUNKING_FAILED)
        normalized_source_characters = sum(len(unit.text) for unit in units)
        if not units or normalized_source_characters == 0:
            return _empty_outcome(
                extraction,
                document_id=document_id,
                config=config,
                status=ChunkingStatus.SKIPPED,
                reason=ChunkingReason.EMPTY_CONTENT,
            )
        if normalized_source_characters > config.max_source_characters:
            raise _ChunkingFailure(ChunkingReason.LIMIT_EXCEEDED)

        if config.strategy is ChunkingStrategy.FIXED_SIZE:
            rendered_chunks = _fixed_size_chunks(units, config)
        else:
            rendered_chunks = _structure_aware_chunks(units, config)
        if len(rendered_chunks) > config.max_chunks_per_document:
            raise _ChunkingFailure(ChunkingReason.LIMIT_EXCEEDED)

        chunks = tuple(
            _make_chunk(
                rendered,
                ordinal=ordinal,
                document_id=document_id,
                source_sha256=extraction.source_sha256,
                document_format=extraction.document_format,
                config=config,
                split_oversized_element=split_oversized_element,
            )
            for ordinal, (rendered, split_oversized_element) in enumerate(rendered_chunks)
        )
        if any(len(chunk.text) > config.max_characters for chunk in chunks):
            raise _ChunkingFailure(ChunkingReason.LIMIT_EXCEEDED)

        covered_source_characters, emitted_source_characters = _coverage(chunks)
        if covered_source_characters != normalized_source_characters:
            raise _ChunkingFailure(ChunkingReason.CHUNKING_FAILED)
        if any(not chunk.provenance for chunk in chunks):
            raise _ChunkingFailure(ChunkingReason.CHUNKING_FAILED)

        return ChunkingOutcome(
            document_id=document_id,
            source_sha256=extraction.source_sha256,
            document_format=extraction.document_format,
            strategy=config.strategy,
            config_version=config.version,
            config_fingerprint=config_fingerprint,
            status=ChunkingStatus.CHUNKED,
            chunks=chunks,
            normalized_source_characters=normalized_source_characters,
            covered_source_characters=covered_source_characters,
            emitted_source_characters=emitted_source_characters,
        )
    except _ChunkingFailure as exc:
        return _empty_outcome(
            extraction,
            document_id=document_id,
            config=config,
            status=ChunkingStatus.FAILED,
            reason=exc.reason,
        )
    except Exception:  # noqa: BLE001 - document isolation is a public guarantee
        return _empty_outcome(
            extraction,
            document_id=document_id,
            config=config,
            status=ChunkingStatus.FAILED,
            reason=ChunkingReason.CHUNKING_FAILED,
        )


def chunk_documents(
    documents: Iterable[tuple[str, ExtractionResult]],
    *,
    config: ChunkingConfig,
) -> tuple[ChunkingOutcome, ...]:
    """Chunk documents independently in caller-provided deterministic order."""

    return tuple(
        chunk_document(extraction, document_id=document_id, config=config)
        for document_id, extraction in documents
    )


def evaluate_corpus(
    input_dir: str | Path,
    output_dir: str | Path,
    *,
    max_characters: int = 1_200,
    overlap_characters: int = 120,
    allow_needs_review: bool = False,
    measure_latency: bool = True,
    extraction_limits: ExtractionLimits | None = None,
) -> dict[str, Any]:
    """Extract, compare both strategies, and persist metadata-only reports."""

    source_root = Path(input_dir).resolve()
    result_root = Path(output_dir).resolve()
    if not source_root.is_dir():
        raise EvaluationConfigurationError("input directory must exist")
    if result_root == source_root or source_root in result_root.parents:
        raise EvaluationConfigurationError("output directory must be outside input directory")

    limits = extraction_limits or ExtractionLimits()
    files: list[Path] = []
    for path in source_root.rglob("*"):
        if path.is_file():
            files.append(path)
            if len(files) > limits.max_files:
                raise EvaluationConfigurationError("input directory exceeds the file limit")
    files.sort(key=lambda path: path.relative_to(source_root).as_posix().casefold())

    configs = tuple(
        ChunkingConfig(
            strategy=strategy,
            max_characters=max_characters,
            overlap_characters=overlap_characters,
            allow_needs_review=allow_needs_review,
        )
        for strategy in ChunkingStrategy
    )
    outcomes_by_strategy: dict[ChunkingStrategy, list[ChunkingOutcome]] = {
        config.strategy: [] for config in configs
    }
    records: list[dict[str, Any]] = []

    for index, path in enumerate(files, start=1):
        document_id = f"document-{index:06d}"
        extraction_started_ns = time.perf_counter_ns()
        try:
            extraction = extract_file(path, limits=limits)
        except Exception:  # noqa: BLE001 - isolate unexpected extractor failures per file
            extraction = ExtractionResult(
                schema_version="1.0",
                document_format={
                    ".pdf": DocumentFormat.PDF,
                    ".docx": DocumentFormat.DOCX,
                    ".xlsx": DocumentFormat.XLSX,
                }.get(path.suffix.lower()),
                status=ExtractionStatus.FAILED,
                source_sha256=None,
                source_size_bytes=None,
                error_category=ErrorCategory.EXTRACTION_FAILED,
            )
        extraction_latency_ms = (time.perf_counter_ns() - extraction_started_ns) / 1_000_000
        strategy_records: dict[str, Any] = {}

        for config in configs:
            chunking_started_ns = time.perf_counter_ns()
            outcome = chunk_document(extraction, document_id=document_id, config=config)
            chunking_latency_ms = (time.perf_counter_ns() - chunking_started_ns) / 1_000_000
            outcomes_by_strategy[config.strategy].append(outcome)
            strategy_records[config.strategy.value] = _safe_outcome_summary(
                outcome,
                config=config,
                latency_ms=round(chunking_latency_ms, 3) if measure_latency else None,
            )

        records.append(
            {
                "document_id": document_id,
                "format": (
                    extraction.document_format.value if extraction.document_format else None
                ),
                "source_sha256": extraction.source_sha256,
                "source_size_bytes": extraction.source_size_bytes,
                "extraction_status": extraction.status.value,
                "extraction_error_category": (
                    extraction.error_category.value if extraction.error_category else None
                ),
                "extraction_warning_codes": sorted(
                    {warning.code.value for warning in extraction.warnings}
                ),
                "extraction_latency_ms": (
                    round(extraction_latency_ms, 3) if measure_latency else None
                ),
                "strategies": strategy_records,
            }
        )

    strategy_aggregates = {
        config.strategy.value: _aggregate_strategy(
            outcomes_by_strategy[config.strategy],
            config=config,
            records=records,
        )
        for config in configs
    }
    deterministic_payload = {
        "documents": [_without_runtime_metrics(record) for record in records],
        "strategies": {
            name: _without_runtime_metrics(aggregate)
            for name, aggregate in strategy_aggregates.items()
        },
    }
    deterministic_report_fingerprint = hashlib.sha256(
        json.dumps(deterministic_payload, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()
    payload = {
        "schema_version": "1.0",
        "document_count": len(records),
        "deterministic_report_fingerprint": deterministic_report_fingerprint,
        "strategies": strategy_aggregates,
        "documents": records,
    }

    result_root.mkdir(parents=True, exist_ok=True)
    (result_root / "metrics.json").write_text(
        json.dumps(payload, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    _write_manifest(result_root / "manifest.csv", records)
    _write_report(result_root / "report.md", payload)
    return payload


def _same_note_id(left: str, right: str) -> bool:
    # OOXML decimal IDs have numeric identity; retain both original lexical forms.
    return (
        re.fullmatch(r"[+-]?[0-9]{1,10}", left) is not None
        and re.fullmatch(r"[+-]?[0-9]{1,10}", right) is not None
        and int(left) == int(right)
    )


def _build_source_units(extraction: ExtractionResult) -> tuple[_SourceUnit, ...]:
    if extraction.document_format is DocumentFormat.XLSX:
        return _xlsx_source_units(extraction)

    units: list[_SourceUnit] = []
    for block in extraction.blocks:
        reference_units = []
        if block.source.block_index is not None:
            for index, ref in enumerate(extraction.note_references):
                if (
                    ref.source.block_index == block.source.block_index
                    and ref.source.part_uri == block.source.part_uri
                ):
                    scope = f"docx-part:{ref.source.part_uri}"
                    reference_units.append(
                        _SourceUnit(
                            unit_id=f"docx-note-reference:{index}",
                            text=f"footnote_reference={ref.note_id}\ttarget_part={ref.target_part_uri}",
                            source=ref.source,
                            scope_key=scope,
                            group_key=f"{scope}:block:{ref.source.block_index}",
                        )
                    )
        if isinstance(block, TextBlock):
            text, text_block_map = _normalize_text_block(block.text)
            if not text:
                units.extend(reference_units)
                continue
            if extraction.document_format is DocumentFormat.PDF:
                page_number = block.source.page_number
                if page_number is None:
                    raise _ChunkingFailure(ChunkingReason.CHUNKING_FAILED)
                scope_key = f"pdf-page:{page_number}"
                group_key = scope_key
                unit_id = f"page:{page_number}:block:{block.order}"
            else:
                scope_key = f"docx-part:{block.source.part_uri or 'body'}"
                group_key = f"{scope_key}:block:{block.order}"
                unit_id = group_key
            units.append(
                _SourceUnit(
                    unit_id=unit_id,
                    text=text,
                    source=block.source,
                    scope_key=scope_key,
                    group_key=group_key,
                    text_block_map=text_block_map,
                    context_ids=tuple(
                        f"docx-note-reference:{index}"
                        for index, ref in enumerate(extraction.note_references)
                        if block.source.note_id is not None
                        and _same_note_id(ref.note_id, block.source.note_id)
                        and ref.target_part_uri == block.source.part_uri
                    ),
                )
            )
        elif isinstance(block, TableBlock):
            scope_key = (
                f"{extraction.document_format.value}:part:{block.source.part_uri}:"
                f"page:{block.source.page_number}:block:{block.source.block_index}:"
                f"table:{block.table_index}"
            )
            header_ids = tuple(
                f"{scope_key}:row:{cell.row_index}:column:{cell.column_index}"
                for row in block.rows
                if row.row_index in block.header_row_indices
                for cell in row.cells
            )
            for row in block.rows:
                group_key = f"{scope_key}:row:{row.row_index}"
                for cell in row.cells:
                    if any(
                        getattr(cell.source, key) != getattr(block.source, key)
                        for key in ("part_uri", "page_number", "block_index", "table_index")
                    ) or (
                        cell.source.row_index != cell.row_index
                        or cell.source.column_index != cell.column_index
                    ):
                        raise _ChunkingFailure(ChunkingReason.CHUNKING_FAILED)
                    value = _normalize_text(cell.text)
                    span = (
                        f"{cell.row_span}x{cell.column_span}"
                        if cell.row_span > 1 or cell.column_span > 1
                        else "1x1"
                    )
                    text = f"r{cell.row_index + 1}c{cell.column_index + 1}[{span}]\t{value}"
                    units.append(
                        _SourceUnit(
                            unit_id=f"{scope_key}:row:{cell.row_index}:column:{cell.column_index}",
                            text=text,
                            source=cell.source,
                            scope_key=scope_key,
                            group_key=group_key,
                            context_ids=(
                                header_ids if row.row_index not in block.header_row_indices else ()
                            ),
                        )
                    )
        else:
            raise TypeError("unsupported extraction block")
        units.extend(reference_units)
    return tuple(units)


def _xlsx_source_units(extraction: ExtractionResult) -> tuple[_SourceUnit, ...]:
    units: list[_SourceUnit] = []
    for sheet in extraction.sheets:
        scope_key = f"xlsx-sheet:{sheet.index}"
        for range_index, merged_range in enumerate(sheet.merged_ranges):
            units.append(
                _SourceUnit(
                    unit_id=f"sheet:{sheet.index}:merge:{range_index}",
                    text=f"merged_range={merged_range}",
                    source=SourceLocation(
                        sheet_index=sheet.index,
                        sheet_name=sheet.name,
                        cell_coordinate=merged_range,
                    ),
                    scope_key=scope_key,
                    group_key=f"{scope_key}:metadata",
                )
            )
        for table_index, table in enumerate(sheet.tables):
            table_name = table.display_name or table.name or "<unnamed>"
            columns = "|".join(table.columns)
            units.append(
                _SourceUnit(
                    unit_id=f"sheet:{sheet.index}:table:{table_index}",
                    text=(f"table={table_name}\treference={table.reference}\tcolumns={columns}"),
                    source=SourceLocation(
                        sheet_index=sheet.index,
                        sheet_name=sheet.name,
                        table_index=table_index,
                    ),
                    scope_key=scope_key,
                    group_key=f"{scope_key}:metadata",
                )
            )
        ranges = [range_boundaries(table.reference) for table in sheet.tables]
        cells_by_row: dict[int, list[SpreadsheetCell]] = {}
        for cell in sheet.cells:
            cells_by_row.setdefault(cell.row_index + 1, []).append(cell)
        table_headers = [
            tuple(
                f"sheet:{sheet.index}:cell:{cell.coordinate}"
                for cell in cells_by_row.get(min_row, ())
                if min_col <= cell.column_index + 1 <= max_col
            )
            for min_col, min_row, max_col, _ in ranges
        ]
        merge_anchors = {
            merged.split(":", 1)[0]: f"sheet:{sheet.index}:merge:{index}"
            for index, merged in enumerate(sheet.merged_ranges)
        }
        for cell in sheet.cells:
            context_ids = []
            table_indices = []
            for index, (min_col, min_row, max_col, max_row) in enumerate(ranges):
                if (
                    min_col <= cell.column_index + 1 <= max_col
                    and min_row <= cell.row_index + 1 <= max_row
                ):
                    table_indices.append(index)
                    context_ids.append(f"sheet:{sheet.index}:table:{index}")
                    if sheet.tables[index].header_row_count and cell.row_index + 1 > min_row:
                        context_ids.extend(table_headers[index])
            if len(table_indices) > 1:
                raise _ChunkingFailure(ChunkingReason.CHUNKING_FAILED)
            # A merge is a relationship, not an invitation to duplicate its value.
            if cell.coordinate in merge_anchors:
                context_ids.append(merge_anchors[cell.coordinate])
            unit = _xlsx_cell_unit(sheet.index, sheet.name, cell)
            if cell.number_format is not None:
                unit = replace(unit, text=f"{unit.text}\tnumber_format={cell.number_format}")
            elif cell.number_format_id is not None:
                unit = replace(unit, text=f"{unit.text}\tnumber_format_id={cell.number_format_id}")
            if cell.data_type is not None:
                unit = replace(unit, text=f"{unit.text}\tdata_type={cell.data_type}")
            if cell.data_type == "n" and cell.number_format and is_date_format(cell.number_format):
                unit = replace(unit, text=f"{unit.text}\tdate_system={sheet.date_system}")
            if cell.raw_value is not None and cell.raw_value != cell.value:
                unit = replace(unit, text=f"{unit.text}\traw_value={cell.raw_value}")
            if cell.has_formula and cell.formula is None and cell.cached_value is None:
                unit = replace(unit, text=f"{unit.text}\tformula=<unavailable>\tcache=<missing>")
            if cell.formula_attributes:
                attributes = json.dumps(dict(cell.formula_attributes), sort_keys=True)
                unit = replace(unit, text=f"{unit.text}\tformula_attributes={attributes}")
            if table_indices:
                unit = replace(unit, source=replace(unit.source, table_index=table_indices[0]))
            units.append(replace(unit, context_ids=tuple(context_ids)))
    return tuple(units)


def _xlsx_cell_unit(
    sheet_index: int,
    sheet_name: str,
    cell: SpreadsheetCell,
) -> _SourceUnit:
    if cell.formula is not None:
        cached = cell.cached_value if cell.cached_value is not None else "<missing>"
        text = f"{cell.coordinate}\tvalue={cached}\tformula={cell.formula}"
    elif cell.cached_value is not None:
        text = (
            f"{cell.coordinate}\tvalue={cell.cached_value}"
            "\tformula=<unavailable-shared>"
        )
    else:
        text = f"{cell.coordinate}\tvalue={cell.value or ''}"
    return _SourceUnit(
        unit_id=f"sheet:{sheet_index}:cell:{cell.coordinate}",
        # Whitespace inside stored values and quoted formula strings is literal.
        # Prose normalization would rewrite CR/CRLF and strip factual content.
        text=text,
        source=SourceLocation(
            sheet_index=sheet_index,
            sheet_name=sheet_name,
            cell_coordinate=cell.coordinate,
            row_index=cell.row_index,
            column_index=cell.column_index,
        ),
        scope_key=f"xlsx-sheet:{sheet_index}",
        group_key=f"xlsx-sheet:{sheet_index}:row:{cell.row_index}",
    )


def _fixed_size_chunks(
    units: tuple[_SourceUnit, ...],
    config: ChunkingConfig,
) -> list[tuple[_RenderedText, bool]]:
    rendered = _render_units(units, separator="\n\n")
    chunks: list[tuple[_RenderedText, bool]] = []
    for start, end in _windows(
        len(rendered.text),
        max_characters=config.max_characters,
        overlap_characters=config.overlap_characters,
    ):
        sliced = _slice_rendered(rendered, start, end)
        if sliced.spans:
            chunks.append((sliced, False))
    return chunks


def _structure_aware_chunks(
    units: tuple[_SourceUnit, ...],
    config: ChunkingConfig,
) -> list[tuple[_RenderedText, bool]]:
    groups = _structural_groups(units)
    chunks: list[tuple[_RenderedText, bool]] = []
    current: list[_StructuralGroup] = []
    current_scope: str | None = None

    for group in groups:
        if current and group.scope_key != current_scope:
            chunks.append((_render_groups(current), False))
            current = []
        current_scope = group.scope_key

        if len(group.rendered.text) > config.max_characters:
            if current:
                chunks.append((_render_groups(current), False))
                current = []
            for start, end in _windows(
                len(group.rendered.text),
                max_characters=config.max_characters,
                overlap_characters=config.overlap_characters,
            ):
                sliced = _slice_rendered(group.rendered, start, end)
                if sliced.spans:
                    chunks.append((sliced, True))
            continue

        candidate = [*current, group]
        if current and len(_render_groups(candidate).text) > config.max_characters:
            chunks.append((_render_groups(current), False))
            current = _tail_groups(current, config.overlap_characters)
            while current and len(_render_groups([*current, group]).text) > config.max_characters:
                current.pop(0)
            current.append(group)
        else:
            current = candidate

    if current:
        chunks.append((_render_groups(current), False))
    return chunks


def _structural_groups(units: tuple[_SourceUnit, ...]) -> tuple[_StructuralGroup, ...]:
    groups: list[_StructuralGroup] = []
    current: list[_SourceUnit] = []
    current_key: tuple[str, str] | None = None
    for unit in units:
        key = (unit.scope_key, unit.group_key)
        if current and key != current_key:
            groups.append(
                _StructuralGroup(
                    scope_key=current[0].scope_key,
                    group_key=current[0].group_key,
                    units=tuple(current),
                    rendered=_render_units(tuple(current), separator="\t"),
                )
            )
            current = []
        current_key = key
        current.append(unit)
    if current:
        groups.append(
            _StructuralGroup(
                scope_key=current[0].scope_key,
                group_key=current[0].group_key,
                units=tuple(current),
                rendered=_render_units(tuple(current), separator="\t"),
            )
        )
    by_id = {unit.unit_id: unit for unit in units}
    enriched = []
    for group in groups:
        ids = tuple(dict.fromkeys(key for unit in group.units for key in unit.context_ids))
        context = tuple(by_id[key] for key in ids)
        group = replace(group, context_units=context)
        enriched.append(replace(group, rendered=_render_groups([group])))
    return tuple(enriched)


def _tail_groups(
    groups: list[_StructuralGroup],
    overlap_characters: int,
) -> list[_StructuralGroup]:
    if overlap_characters == 0:
        return []
    tail: list[_StructuralGroup] = []
    for group in reversed(groups):
        candidate = [group, *tail]
        if len(_render_groups(candidate).text) > overlap_characters:
            break
        tail = candidate
    return tail


def _render_groups(groups: list[_StructuralGroup]) -> _RenderedText:
    units: list[_SourceUnit] = []
    separators: list[str] = []
    present = {unit.unit_id for group in groups for unit in group.units}
    for group in groups:
        for unit in group.context_units:
            if unit.unit_id not in present:
                if units:
                    separators.append("\n\n")
                units.append(unit)
                present.add(unit.unit_id)
    for group in groups:
        for unit_index, unit in enumerate(group.units):
            if units:
                separators.append("\t" if unit_index else "\n\n")
            units.append(unit)
    return _render_units_with_separators(tuple(units), tuple(separators))


def _render_units(units: tuple[_SourceUnit, ...], *, separator: str) -> _RenderedText:
    return _render_units_with_separators(
        units,
        tuple(separator for _ in range(max(0, len(units) - 1))),
    )


def _render_units_with_separators(
    units: tuple[_SourceUnit, ...],
    separators: tuple[str, ...],
) -> _RenderedText:
    parts: list[str] = []
    spans: list[_RenderedSpan] = []
    position = 0
    for index, unit in enumerate(units):
        if index:
            separator = separators[index - 1]
            parts.append(separator)
            position += len(separator)
        start = position
        parts.append(unit.text)
        position += len(unit.text)
        spans.append(
            _RenderedSpan(
                unit=unit,
                rendered_start=start,
                rendered_end=position,
                source_start=0,
                source_end=len(unit.text),
            )
        )
    return _RenderedText(text="".join(parts), spans=tuple(spans))


def _windows(
    length: int,
    *,
    max_characters: int,
    overlap_characters: int,
) -> tuple[tuple[int, int], ...]:
    windows: list[tuple[int, int]] = []
    start = 0
    while start < length:
        end = min(length, start + max_characters)
        windows.append((start, end))
        if end == length:
            break
        start = end - overlap_characters
    return tuple(windows)


def _slice_rendered(rendered: _RenderedText, start: int, end: int) -> _RenderedText:
    spans: list[_RenderedSpan] = []
    for span in rendered.spans:
        intersection_start = max(start, span.rendered_start)
        intersection_end = min(end, span.rendered_end)
        if intersection_start >= intersection_end:
            continue
        unit_offset_start = intersection_start - span.rendered_start
        unit_offset_end = intersection_end - span.rendered_start
        spans.append(
            _RenderedSpan(
                unit=span.unit,
                rendered_start=intersection_start - start,
                rendered_end=intersection_end - start,
                source_start=span.source_start + unit_offset_start,
                source_end=span.source_start + unit_offset_end,
            )
        )
    return _RenderedText(text=rendered.text[start:end], spans=tuple(spans))


def _make_chunk(
    rendered: _RenderedText,
    *,
    ordinal: int,
    document_id: str,
    source_sha256: str,
    document_format: DocumentFormat,
    config: ChunkingConfig,
    split_oversized_element: bool,
) -> Chunk:
    provenance = tuple(
        ChunkSourceSpan(
            unit_id=span.unit.unit_id,
            source=span.unit.source,
            source_start=span.source_start,
            source_end=span.source_end,
            source_length=len(span.unit.text),
            text_block_map=_slice_text_block_map(span),
        )
        for span in rendered.spans
    )
    identity_payload = {
        "config": config.fingerprint(),
        "document_id": document_id,
        "format": document_format.value,
        "ordinal": ordinal,
        "provenance": [_source_span_identity(item) for item in provenance],
        "source_sha256": source_sha256,
        "text": rendered.text,
    }
    digest = hashlib.sha256(
        json.dumps(identity_payload, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()
    return Chunk(
        chunk_id=f"chunk-{digest}",
        document_id=document_id,
        source_sha256=source_sha256,
        document_format=document_format,
        strategy=config.strategy,
        config_version=config.version,
        config_fingerprint=config.fingerprint(),
        ordinal=ordinal,
        text=rendered.text,
        provenance=provenance,
        split_oversized_element=split_oversized_element,
    )


def _source_span_identity(span: ChunkSourceSpan) -> dict[str, Any]:
    return {
        "source": {field.name: getattr(span.source, field.name) for field in fields(span.source)},
        "source_end": span.source_end,
        "source_length": span.source_length,
        "source_start": span.source_start,
        "unit_id": span.unit_id,
        **({"text_block_map": asdict(span.text_block_map)} if span.text_block_map else {}),
    }


def _coverage(chunks: tuple[Chunk, ...]) -> tuple[int, int]:
    intervals: dict[str, list[tuple[int, int]]] = {}
    emitted = 0
    for chunk in chunks:
        for span in chunk.provenance:
            intervals.setdefault(span.unit_id, []).append((span.source_start, span.source_end))
            emitted += span.source_end - span.source_start
    covered = 0
    for unit_intervals in intervals.values():
        unit_intervals.sort()
        current_start, current_end = unit_intervals[0]
        for start, end in unit_intervals[1:]:
            if start > current_end:
                covered += current_end - current_start
                current_start, current_end = start, end
            else:
                current_end = max(current_end, end)
        covered += current_end - current_start
    return covered, emitted


def _empty_outcome(
    extraction: ExtractionResult,
    *,
    document_id: str,
    config: ChunkingConfig,
    status: ChunkingStatus,
    reason: ChunkingReason,
) -> ChunkingOutcome:
    return ChunkingOutcome(
        document_id=document_id,
        source_sha256=extraction.source_sha256,
        document_format=extraction.document_format,
        strategy=config.strategy,
        config_version=config.version,
        config_fingerprint=config.fingerprint(),
        status=status,
        reason=reason,
    )


def _normalize_text(value: str) -> str:
    return value.replace("\r\n", "\n").replace("\r", "\n").strip()


def _normalize_text_block(value: str) -> tuple[str, TextBlockOffsetMap]:
    """Keep the existing normalization and map surviving Python code points."""
    text = _normalize_text(value)
    start = len(value) - len(value.lstrip())
    end = len(value.rstrip())
    ranges: list[NormalizationRange] = []
    original_cursor, normalized_cursor = start, 0
    for match in re.finditer(r"\r\n|\r", value[start:end]):
        left, right = start + match.start(), start + match.end()
        if left > original_cursor:
            width = left - original_cursor
            ranges.append(
                NormalizationRange(
                    normalized_cursor,
                    normalized_cursor + width,
                    original_cursor,
                    left,
                )
            )
            normalized_cursor += width
        ranges.append(NormalizationRange(normalized_cursor, normalized_cursor + 1, left, right))
        normalized_cursor += 1
        original_cursor = right
    if end > original_cursor:
        ranges.append(
            NormalizationRange(
                normalized_cursor,
                normalized_cursor + end - original_cursor,
                original_cursor,
                end,
            )
        )
    return text, TextBlockOffsetMap(
        original_text_sha256=hashlib.sha256(value.encode()).hexdigest(),
        original_text_length=len(value),
        chunk_start=0,
        chunk_end=len(text),
        ranges=tuple(ranges),
    )


def _slice_text_block_map(span: _RenderedSpan) -> TextBlockOffsetMap | None:
    mapping = span.unit.text_block_map
    if mapping is None:
        return None
    ranges = []
    # Binary search avoids rescanning an oversized block's whole map per chunk.
    first = bisect_right(mapping.ranges, span.source_start, key=lambda item: item.normalized_end)
    for index in range(first, len(mapping.ranges)):
        item = mapping.ranges[index]
        left, right = (
            max(item.normalized_start, span.source_start),
            min(
                item.normalized_end,
                span.source_end,
            ),
        )
        if left >= right:
            break
        if item.original_end - item.original_start == item.normalized_end - item.normalized_start:
            original_start = item.original_start + left - item.normalized_start
            original_end = original_start + right - left
        else:
            # A CRLF pair is one indivisible normalized character.
            original_start, original_end = item.original_start, item.original_end
        ranges.append(NormalizationRange(left, right, original_start, original_end))
    return replace(
        mapping, chunk_start=span.rendered_start, chunk_end=span.rendered_end, ranges=tuple(ranges)
    )


def _safe_outcome_summary(
    outcome: ChunkingOutcome,
    *,
    config: ChunkingConfig,
    latency_ms: float | None,
) -> dict[str, Any]:
    lengths = [len(chunk.text) for chunk in outcome.chunks]
    provenance_chunks = sum(bool(chunk.provenance) for chunk in outcome.chunks)
    return {
        "status": outcome.status.value,
        "reason": outcome.reason.value if outcome.reason else None,
        "config_fingerprint": outcome.config_fingerprint,
        "chunk_count": len(outcome.chunks),
        "chunk_fingerprint": outcome.deterministic_fingerprint(),
        "length_distribution": _length_distribution(lengths),
        "oversized_chunk_count": sum(length > config.max_characters for length in lengths),
        "oversized_element_split_chunk_count": sum(
            chunk.split_oversized_element for chunk in outcome.chunks
        ),
        "normalized_source_characters": outcome.normalized_source_characters,
        "covered_source_characters": outcome.covered_source_characters,
        "emitted_source_characters": outcome.emitted_source_characters,
        "normalized_content_coverage": _ratio(
            outcome.covered_source_characters,
            outcome.normalized_source_characters,
        ),
        "provenance_chunk_count": provenance_chunks,
        "provenance_coverage": _ratio(provenance_chunks, len(outcome.chunks)),
        "latency_ms": latency_ms,
    }


def _aggregate_strategy(
    outcomes: list[ChunkingOutcome],
    *,
    config: ChunkingConfig,
    records: list[dict[str, Any]],
) -> dict[str, Any]:
    chunks = [chunk for outcome in outcomes for chunk in outcome.chunks]
    status_counts = Counter(outcome.status.value for outcome in outcomes)
    skipped_reason_counts = Counter(
        outcome.reason.value
        for outcome in outcomes
        if outcome.status is ChunkingStatus.SKIPPED and outcome.reason is not None
    )
    failed_reason_counts = Counter(
        outcome.reason.value
        for outcome in outcomes
        if outcome.status is ChunkingStatus.FAILED and outcome.reason is not None
    )
    source_characters = sum(outcome.normalized_source_characters for outcome in outcomes)
    covered_characters = sum(outcome.covered_source_characters for outcome in outcomes)
    emitted_characters = sum(outcome.emitted_source_characters for outcome in outcomes)
    provenance_chunks = sum(bool(chunk.provenance) for chunk in chunks)
    strategy_latencies = [
        record["strategies"][config.strategy.value]["latency_ms"] for record in records
    ]
    measured_latencies = [value for value in strategy_latencies if value is not None]
    return {
        "config": config_as_dict(config),
        "document_status_counts": dict(sorted(status_counts.items())),
        "skipped_reason_counts": dict(sorted(skipped_reason_counts.items())),
        "failed_reason_counts": dict(sorted(failed_reason_counts.items())),
        "chunk_count": len(chunks),
        "chunk_length_distribution": _length_distribution(
            [len(chunk.text) for chunk in chunks]
        ),
        "oversized_chunk_count": sum(
            len(chunk.text) > config.max_characters for chunk in chunks
        ),
        "oversized_element_split_chunk_count": sum(
            chunk.split_oversized_element for chunk in chunks
        ),
        "normalized_source_characters": source_characters,
        "covered_source_characters": covered_characters,
        "emitted_source_characters": emitted_characters,
        "normalized_content_coverage": _ratio(covered_characters, source_characters),
        "source_character_duplication_ratio": _ratio(
            emitted_characters,
            covered_characters,
        ),
        "provenance_chunk_count": provenance_chunks,
        "provenance_coverage": _ratio(provenance_chunks, len(chunks)),
        "deterministic_chunk_fingerprint": hashlib.sha256(
            json.dumps(
                [outcome.deterministic_fingerprint() for outcome in outcomes],
                separators=(",", ":"),
            ).encode()
        ).hexdigest(),
        "latency_ms_total": (
            round(sum(measured_latencies), 3) if measured_latencies else None
        ),
    }


def _length_distribution(values: list[int]) -> dict[str, float | int | None]:
    if not values:
        return {"minimum": None, "mean": None, "p50": None, "p95": None, "maximum": None}
    ordered = sorted(values)
    return {
        "minimum": ordered[0],
        "mean": round(sum(ordered) / len(ordered), 3),
        "p50": ordered[max(0, math.ceil(0.50 * len(ordered)) - 1)],
        "p95": ordered[max(0, math.ceil(0.95 * len(ordered)) - 1)],
        "maximum": ordered[-1],
    }


def _ratio(numerator: int, denominator: int) -> float | None:
    if denominator == 0:
        return None
    return round(numerator / denominator, 6)


def _without_runtime_metrics(value: Any) -> Any:
    if isinstance(value, dict):
        return {
            key: _without_runtime_metrics(item)
            for key, item in value.items()
            if key not in {"latency_ms", "latency_ms_total", "extraction_latency_ms"}
        }
    if isinstance(value, list):
        return [_without_runtime_metrics(item) for item in value]
    return value


def _write_manifest(path: Path, records: list[dict[str, Any]]) -> None:
    fieldnames = [
        "document_id",
        "format",
        "source_sha256",
        "source_size_bytes",
        "extraction_status",
        "extraction_error_category",
        "extraction_warning_codes",
        "extraction_latency_ms",
    ]
    for strategy in ChunkingStrategy:
        prefix = strategy.value
        fieldnames.extend(
            [
                f"{prefix}_status",
                f"{prefix}_reason",
                f"{prefix}_chunk_count",
                f"{prefix}_chunk_fingerprint",
                f"{prefix}_normalized_content_coverage",
                f"{prefix}_provenance_coverage",
                f"{prefix}_oversized_chunk_count",
                f"{prefix}_latency_ms",
            ]
        )
    with path.open("w", encoding="utf-8", newline="") as destination:
        writer = csv.DictWriter(destination, fieldnames=fieldnames)
        writer.writeheader()
        for record in records:
            row: dict[str, Any] = {
                key: record[key]
                for key in fieldnames[:8]
                if key != "extraction_warning_codes"
            }
            row["extraction_warning_codes"] = "|".join(
                record["extraction_warning_codes"]
            )
            for strategy in ChunkingStrategy:
                prefix = strategy.value
                summary = record["strategies"][prefix]
                for field_name in fieldnames[8 + list(ChunkingStrategy).index(strategy) * 8 :][
                    :8
                ]:
                    metric = field_name.removeprefix(f"{prefix}_")
                    row[field_name] = summary[metric]
            writer.writerow(row)


def _write_report(path: Path, payload: dict[str, Any]) -> None:
    lines = [
        "# Offline chunking evaluation",
        "",
        "This report contains metadata only; chunk and source content are not persisted.",
        "",
        f"- Documents: {payload['document_count']}",
        f"- Deterministic report fingerprint: `{payload['deterministic_report_fingerprint']}`",
        "",
    ]
    for strategy in ChunkingStrategy:
        aggregate = payload["strategies"][strategy.value]
        status_summary = json.dumps(aggregate["document_status_counts"], sort_keys=True)
        length_summary = json.dumps(aggregate["chunk_length_distribution"], sort_keys=True)
        lines.extend(
            [
                f"## {strategy.value}",
                "",
                f"- Document statuses: {status_summary}",
                f"- Chunks: {aggregate['chunk_count']}",
                f"- Length distribution: {length_summary}",
                f"- Normalized content coverage: {aggregate['normalized_content_coverage']}",
                f"- Provenance coverage: {aggregate['provenance_coverage']}",
                f"- Oversized chunks: {aggregate['oversized_chunk_count']}",
                f"- Chunk fingerprint: `{aggregate['deterministic_chunk_fingerprint']}`",
                "",
            ]
        )
    path.write_text("\n".join(lines), encoding="utf-8")

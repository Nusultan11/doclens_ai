"""Deterministic, privacy-safe profiling for an explicitly supplied local corpus."""

from __future__ import annotations

import csv
import hashlib
import json
import logging
import math
import os
import time
import zipfile
from collections import Counter
from collections.abc import Callable
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

SUPPORTED_FORMATS = {".pdf": "pdf", ".docx": "docx", ".xlsx": "xlsx"}
MANIFEST_FIELDS = (
    "document_id",
    "format",
    "file_size_bytes",
    "sha256",
    "duplicate_of",
    "extraction_success",
    "error_category",
    "extraction_latency_ms",
    "extracted_character_count",
    "empty_text",
    "pdf_page_count",
    "pdf_pages_with_text",
    "pdf_pages_without_text",
    "pdf_possible_image_page_count",
    "possible_scanned_pdf",
    "docx_paragraph_count",
    "docx_table_count",
    "docx_table_cell_count",
    "docx_populated_table_cell_count",
    "xlsx_sheet_count",
    "xlsx_populated_cell_count",
    "xlsx_rows_with_data",
)


class ProfilingError(ValueError):
    """Raised when the corpus or output configuration cannot be profiled safely."""


class DocumentProcessingError(Exception):
    def __init__(self, category: str, partial_metrics: dict[str, Any] | None = None) -> None:
        super().__init__(category)
        self.category = category
        self.partial_metrics = partial_metrics or {}


@dataclass(frozen=True)
class ProfileLimits:
    max_files: int = 10_000
    max_file_size_bytes: int = 50 * 1024 * 1024
    max_pdf_pages: int = 2_000
    max_archive_entries: int = 10_000
    max_archive_uncompressed_bytes: int = 200 * 1024 * 1024
    max_archive_entry_bytes: int = 100 * 1024 * 1024
    max_archive_compression_ratio: float = 1_000.0
    max_extracted_characters: int = 5_000_000
    max_xlsx_sheets: int = 256
    max_xlsx_cells_scanned: int = 1_000_000


def _blank_profile(document_id: str, file_format: str) -> dict[str, Any]:
    return {
        "document_id": document_id,
        "format": file_format,
        "file_size_bytes": None,
        "sha256": None,
        "duplicate_of": None,
        "extraction_success": False,
        "error_category": None,
        "extraction_latency_ms": None,
        "extracted_character_count": None,
        "empty_text": None,
        "pdf_page_count": None,
        "pdf_pages_with_text": None,
        "pdf_pages_without_text": None,
        "pdf_possible_image_page_count": None,
        "pdf_page_error_count": None,
        "possible_scanned_pdf": None,
        "docx_paragraph_count": None,
        "docx_table_count": None,
        "docx_table_cell_count": None,
        "docx_populated_table_cell_count": None,
        "xlsx_sheet_count": None,
        "xlsx_populated_cell_count": None,
        "xlsx_rows_with_data": None,
        "xlsx_sheet_populated_cell_counts": None,
    }


def _discover_files(input_dir: Path, output_dir: Path, limits: ProfileLimits) -> list[Path]:
    files: list[Path] = []
    output_resolved = output_dir.resolve()
    for root, directory_names, file_names in os.walk(input_dir, followlinks=False):
        root_path = Path(root)
        directory_names[:] = sorted(directory_names, key=lambda value: (value.casefold(), value))
        file_names.sort(key=lambda value: (value.casefold(), value))
        for file_name in file_names:
            path = root_path / file_name
            try:
                if path.resolve().is_relative_to(output_resolved):
                    continue
            except OSError:
                pass
            files.append(path)
            if len(files) > limits.max_files:
                raise ProfilingError(f"corpus exceeds the {limits.max_files}-file limit")
    return sorted(
        files,
        key=lambda path: (
            path.relative_to(input_dir).as_posix().casefold(),
            path.relative_to(input_dir).as_posix(),
        ),
    )


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while chunk := source.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _check_office_archive(path: Path, limits: ProfileLimits) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > limits.max_archive_entries:
                raise DocumentProcessingError("limit_exceeded")
            total_size = 0
            for entry in entries:
                total_size += entry.file_size
                if entry.file_size > limits.max_archive_entry_bytes:
                    raise DocumentProcessingError("limit_exceeded")
                if total_size > limits.max_archive_uncompressed_bytes:
                    raise DocumentProcessingError("limit_exceeded")
                if entry.file_size and (
                    entry.compress_size == 0
                    or entry.file_size / entry.compress_size
                    > limits.max_archive_compression_ratio
                ):
                    raise DocumentProcessingError("limit_exceeded")
    except zipfile.BadZipFile as exc:
        raise DocumentProcessingError("corrupt_document") from exc


def _extract_pdf(path: Path, limits: ProfileLimits) -> dict[str, Any]:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as exc:
        raise DocumentProcessingError("missing_dependency") from exc

    pypdf_logger = logging.getLogger("pypdf")
    previous_level = pypdf_logger.level
    try:
        pypdf_logger.setLevel(logging.CRITICAL)
        reader = PdfReader(path, strict=False)
        if reader.is_encrypted:
            raise DocumentProcessingError("encrypted_document")
        page_count = len(reader.pages)
        base = {"pdf_page_count": page_count}
        if page_count > limits.max_pdf_pages:
            raise DocumentProcessingError("limit_exceeded", base)

        characters = 0
        pages_with_text = 0
        possible_image_pages = 0
        page_errors = 0
        for page in reader.pages:
            try:
                text = page.extract_text() or ""
            except Exception:  # a bad page must not stop the corpus
                text = ""
                page_errors += 1
            characters += len(text)
            if characters > limits.max_extracted_characters:
                raise DocumentProcessingError("limit_exceeded", base)
            if len("".join(text.split())) >= 20:
                pages_with_text += 1
            else:
                possible_image_pages += 1

        metrics = {
            **base,
            "extracted_character_count": characters,
            "empty_text": characters == 0,
            "pdf_pages_with_text": pages_with_text,
            "pdf_pages_without_text": page_count - pages_with_text,
            "pdf_possible_image_page_count": possible_image_pages,
            "pdf_page_error_count": page_errors,
            "possible_scanned_pdf": bool(
                page_count and possible_image_pages / page_count >= 0.8
            ),
        }
        if page_errors:
            raise DocumentProcessingError("extraction_failed", metrics)
        return metrics
    except DocumentProcessingError:
        raise
    except (PdfReadError, OSError, ValueError) as exc:
        raise DocumentProcessingError("corrupt_document") from exc
    finally:
        pypdf_logger.setLevel(previous_level)


def _extract_docx(path: Path, limits: ProfileLimits) -> dict[str, Any]:
    try:
        from docx import Document
        from docx.opc.exceptions import PackageNotFoundError
    except ImportError as exc:
        raise DocumentProcessingError("missing_dependency") from exc

    _check_office_archive(path, limits)
    try:
        document = Document(path)
        character_count = 0
        for paragraph in document.paragraphs:
            character_count += len(paragraph.text)
            if character_count > limits.max_extracted_characters:
                raise DocumentProcessingError("limit_exceeded")

        table_cell_count = 0
        populated_cell_count = 0
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    table_cell_count += 1
                    character_count += len(cell.text)
                    populated_cell_count += bool(cell.text.strip())
                    if character_count > limits.max_extracted_characters:
                        raise DocumentProcessingError("limit_exceeded")

        return {
            "extracted_character_count": character_count,
            "empty_text": character_count == 0,
            "docx_paragraph_count": len(document.paragraphs),
            "docx_table_count": len(document.tables),
            "docx_table_cell_count": table_cell_count,
            "docx_populated_table_cell_count": populated_cell_count,
        }
    except DocumentProcessingError:
        raise
    except (PackageNotFoundError, KeyError, OSError, ValueError) as exc:
        raise DocumentProcessingError("corrupt_document") from exc


def _extract_xlsx(path: Path, limits: ProfileLimits) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:
        raise DocumentProcessingError("missing_dependency") from exc

    _check_office_archive(path, limits)
    workbook = None
    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        sheet_count = len(workbook.worksheets)
        base = {"xlsx_sheet_count": sheet_count}
        if sheet_count > limits.max_xlsx_sheets:
            raise DocumentProcessingError("limit_exceeded", base)

        scanned_cells = 0
        populated_cells = 0
        rows_with_data = 0
        character_count = 0
        sheet_counts: list[int] = []
        for sheet in workbook.worksheets:
            sheet_populated = 0
            for row in sheet.iter_rows():
                row_has_data = False
                for cell in row:
                    scanned_cells += 1
                    if scanned_cells > limits.max_xlsx_cells_scanned:
                        raise DocumentProcessingError("limit_exceeded", base)
                    value = cell.value
                    if value is not None and value != "":
                        row_has_data = True
                        populated_cells += 1
                        sheet_populated += 1
                        character_count += len(str(value))
                        if character_count > limits.max_extracted_characters:
                            raise DocumentProcessingError("limit_exceeded", base)
                rows_with_data += row_has_data
            sheet_counts.append(sheet_populated)

        return {
            **base,
            "extracted_character_count": character_count,
            "empty_text": populated_cells == 0,
            "xlsx_populated_cell_count": populated_cells,
            "xlsx_rows_with_data": rows_with_data,
            "xlsx_sheet_populated_cell_counts": sheet_counts,
        }
    except DocumentProcessingError:
        raise
    except (InvalidFileException, KeyError, OSError, ValueError, zipfile.BadZipFile) as exc:
        raise DocumentProcessingError("corrupt_document") from exc
    finally:
        if workbook is not None:
            workbook.close()


EXTRACTORS: dict[str, Callable[[Path, ProfileLimits], dict[str, Any]]] = {
    "pdf": _extract_pdf,
    "docx": _extract_docx,
    "xlsx": _extract_xlsx,
}


def _profile_file(
    path: Path,
    document_id: str,
    limits: ProfileLimits,
    measure_latency: bool,
    clock_ns: Callable[[], int],
) -> dict[str, Any]:
    file_format = SUPPORTED_FORMATS.get(path.suffix.lower(), "unsupported")
    profile = _blank_profile(document_id, file_format)
    if path.is_symlink():
        profile["error_category"] = "symlink_not_allowed"
        return profile
    try:
        size = path.stat().st_size
    except OSError:
        profile["error_category"] = "file_unreadable"
        return profile
    profile["file_size_bytes"] = size
    if size == 0:
        profile["sha256"] = hashlib.sha256(b"").hexdigest()
        profile["error_category"] = "empty_file"
        return profile
    if size > limits.max_file_size_bytes:
        profile["error_category"] = "limit_exceeded"
        return profile
    try:
        profile["sha256"] = _sha256(path)
    except OSError:
        profile["error_category"] = "file_unreadable"
        return profile
    if file_format == "unsupported":
        profile["error_category"] = "unsupported_format"
        return profile

    started = clock_ns() if measure_latency else None
    try:
        profile.update(EXTRACTORS[file_format](path, limits))
        profile["extraction_success"] = True
    except DocumentProcessingError as exc:
        profile.update(exc.partial_metrics)
        profile["error_category"] = exc.category
    except Exception:
        profile["error_category"] = "extraction_failed"
    finally:
        if started is not None:
            profile["extraction_latency_ms"] = round((clock_ns() - started) / 1_000_000, 3)
    return profile


def _mark_duplicates(profiles: list[dict[str, Any]]) -> None:
    first_by_hash: dict[str, str] = {}
    for profile in profiles:
        digest = profile["sha256"]
        if digest is None:
            continue
        if digest in first_by_hash:
            profile["duplicate_of"] = first_by_hash[digest]
        else:
            first_by_hash[digest] = profile["document_id"]


def _rate(numerator: int, denominator: int) -> dict[str, int | float | None]:
    return {
        "numerator": numerator,
        "denominator": denominator,
        "percentage": round(numerator * 100 / denominator, 2) if denominator else None,
    }


def _distribution(values: list[int | float]) -> dict[str, Any]:
    ordered = sorted(values)
    result: dict[str, Any] = {
        "sample_count": len(ordered),
        "minimum": ordered[0] if ordered else None,
        "maximum": ordered[-1] if ordered else None,
        "mean": round(sum(ordered) / len(ordered), 3) if ordered else None,
        "p50": None,
        "p95": None,
        "percentile_method": "nearest_rank",
    }
    if len(ordered) < 5:
        result["percentiles_unavailable_reason"] = "fewer_than_5_samples"
        return result
    result["p50"] = ordered[math.ceil(0.50 * len(ordered)) - 1]
    result["p95"] = ordered[math.ceil(0.95 * len(ordered)) - 1]
    return result


def _aggregate(profiles: list[dict[str, Any]]) -> dict[str, Any]:
    supported = [profile for profile in profiles if profile["format"] != "unsupported"]
    successful = [profile for profile in supported if profile["extraction_success"]]
    successfully_text_profiled = [
        profile for profile in successful if profile["empty_text"] is not None
    ]
    successful_pdfs = [profile for profile in successful if profile["format"] == "pdf"]
    scanned_pdfs = [profile for profile in successful_pdfs if profile["possible_scanned_pdf"]]
    hashed = [profile for profile in profiles if profile["sha256"] is not None]
    duplicates = [profile for profile in hashed if profile["duplicate_of"] is not None]
    error_counts = Counter(
        profile["error_category"] for profile in profiles if profile["error_category"]
    )
    format_counts = Counter(profile["format"] for profile in profiles)
    success_by_format = {
        file_format: _rate(
            sum(
                profile["extraction_success"]
                for profile in profiles
                if profile["format"] == file_format
            ),
            format_counts[file_format],
        )
        for file_format in sorted(format_counts)
        if file_format != "unsupported"
    }
    return {
        "document_count": len(profiles),
        "format_counts": dict(sorted(format_counts.items())),
        "supported_document_count": len(supported),
        "unsupported_document_count": len(profiles) - len(supported),
        "hashed_document_count": len(hashed),
        "successful_extraction_count": len(successful),
        "extraction_success_rate": _rate(len(successful), len(supported)),
        "extraction_success_rate_by_format": success_by_format,
        "empty_text_rate": _rate(
            sum(bool(profile["empty_text"]) for profile in successfully_text_profiled),
            len(successfully_text_profiled),
        ),
        "possible_scanned_pdf_rate": _rate(len(scanned_pdfs), len(successful_pdfs)),
        "duplicate_copy_rate": _rate(len(duplicates), len(hashed)),
        "error_category_counts": dict(sorted(error_counts.items())),
        "file_size_bytes": _distribution(
            [
                profile["file_size_bytes"]
                for profile in profiles
                if profile["file_size_bytes"] is not None
            ]
        ),
        "extracted_character_count": _distribution(
            [
                profile["extracted_character_count"]
                for profile in successful
                if profile["extracted_character_count"] is not None
            ]
        ),
        "pdf_page_count": _distribution(
            [profile["pdf_page_count"] for profile in successful_pdfs]
        ),
        "extraction_latency_ms": _distribution(
            [
                profile["extraction_latency_ms"]
                for profile in supported
                if profile["extraction_latency_ms"] is not None
            ]
        ),
    }


def _fingerprint(profiles: list[dict[str, Any]], aggregate: dict[str, Any]) -> str:
    stable_profiles = [
        {key: value for key, value in profile.items() if key != "extraction_latency_ms"}
        for profile in profiles
    ]
    stable_aggregate = {
        key: value for key, value in aggregate.items() if key != "extraction_latency_ms"
    }
    payload = json.dumps(
        {"documents": stable_profiles, "aggregate": stable_aggregate},
        ensure_ascii=False,
        separators=(",", ":"),
        sort_keys=True,
    ).encode()
    return hashlib.sha256(payload).hexdigest()


def _write_json(path: Path, result: dict[str, Any]) -> None:
    content = json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    _write_text(path, content)


def _write_manifest(path: Path, profiles: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as destination:
        writer = csv.DictWriter(destination, fieldnames=MANIFEST_FIELDS, lineterminator="\n")
        writer.writeheader()
        for profile in profiles:
            writer.writerow({field: profile[field] for field in MANIFEST_FIELDS})


def _format_rate(rate: dict[str, Any]) -> str:
    percentage = "n/a" if rate["percentage"] is None else f"{rate['percentage']:.2f}%"
    return f"{rate['numerator']}/{rate['denominator']} ({percentage})"


def _format_distribution(distribution: dict[str, Any]) -> str:
    if not distribution["sample_count"]:
        return "n=0; no measurements"
    p50 = "n/a" if distribution["p50"] is None else distribution["p50"]
    p95 = "n/a" if distribution["p95"] is None else distribution["p95"]
    return (
        f"n={distribution['sample_count']}; min={distribution['minimum']}; "
        f"mean={distribution['mean']}; max={distribution['maximum']}; "
        f"p50={p50}; p95={p95}"
    )


def _write_report(path: Path, result: dict[str, Any]) -> None:
    aggregate = result["aggregate"]
    error_rows = (
        [f"- `{key}`: {value}" for key, value in aggregate["error_category_counts"].items()]
        or ["- None"]
    )
    lines = [
        "# Document EDA report",
        "",
        (
            "This report contains aggregate metadata only. It does not contain filenames, "
            "paths, or extracted content."
        ),
        "",
        "## Coverage",
        "",
        f"- Files discovered: {aggregate['document_count']}",
        f"- Supported files: {aggregate['supported_document_count']}",
        f"- Unsupported files: {aggregate['unsupported_document_count']}",
        f"- Successful extractions: {aggregate['successful_extraction_count']}",
        f"- SHA-256 fingerprints available: {aggregate['hashed_document_count']}",
        f"- Deterministic content fingerprint: `{result['deterministic_content_sha256']}`",
        "",
        "## Format distribution",
        "",
        *[
            f"- `{file_format}`: {count}"
            for file_format, count in aggregate["format_counts"].items()
        ],
        "",
        "## Rates",
        "",
        f"- Extraction success: {_format_rate(aggregate['extraction_success_rate'])}",
        f"- Empty extracted text: {_format_rate(aggregate['empty_text_rate'])}",
        f"- Possible scanned PDFs: {_format_rate(aggregate['possible_scanned_pdf_rate'])}",
        f"- Duplicate copies: {_format_rate(aggregate['duplicate_copy_rate'])}",
        "",
        "Extraction success by supported format:",
        "",
        *[
            f"- `{file_format}`: {_format_rate(rate)}"
            for file_format, rate in aggregate["extraction_success_rate_by_format"].items()
        ],
        "",
        "## Distributions",
        "",
        f"- File size in bytes: {_format_distribution(aggregate['file_size_bytes'])}",
        (
            "- Extracted characters: "
            f"{_format_distribution(aggregate['extracted_character_count'])}"
        ),
        f"- PDF pages: {_format_distribution(aggregate['pdf_page_count'])}",
        (
            "- Extraction latency in milliseconds: "
            f"{_format_distribution(aggregate['extraction_latency_ms'])}"
        ),
        "",
        "## Errors",
        "",
        *error_rows,
        "",
        "## Interpretation",
        "",
        "- Percentiles are omitted for distributions with fewer than five samples.",
        (
            "- A PDF is a possible scan when at least 80% of its pages have fewer than "
            "20 non-whitespace extracted characters."
        ),
        (
            "- Timing is operational and nondeterministic; it is excluded from the content "
            "fingerprint."
        ),
        (
            "- Parser success does not establish semantic extraction quality or corpus "
            "representativeness."
        ),
        (
            "- OCR, production ingestion, storage, database writes, and external services are "
            "outside this baseline."
        ),
        "",
    ]
    _write_text(path, "\n".join(lines))


def _write_text(path: Path, content: str) -> None:
    path.write_text(content, encoding="utf-8", newline="\n")


def profile_corpus(
    input_dir: Path | str,
    output_dir: Path | str,
    *,
    limits: ProfileLimits | None = None,
    measure_latency: bool = True,
    clock_ns: Callable[[], int] = time.perf_counter_ns,
) -> dict[str, Any]:
    """Profile an explicit local corpus and write privacy-safe reports."""

    resolved_input = Path(input_dir).resolve()
    resolved_output = Path(output_dir).resolve()
    if not resolved_input.exists():
        raise ProfilingError("input directory does not exist")
    if not resolved_input.is_dir():
        raise ProfilingError("input path is not a directory")
    if resolved_input == resolved_output:
        raise ProfilingError("output directory must differ from input directory")

    active_limits = limits or ProfileLimits()
    paths = _discover_files(resolved_input, resolved_output, active_limits)
    profiles = [
        _profile_file(
            path,
            f"document-{index:06d}",
            active_limits,
            measure_latency,
            clock_ns,
        )
        for index, path in enumerate(paths, start=1)
    ]
    _mark_duplicates(profiles)
    aggregate = _aggregate(profiles)
    result = {
        "schema_version": 1,
        "runtime_metrics_included": measure_latency,
        "limits": asdict(active_limits),
        "documents": profiles,
        "aggregate": aggregate,
    }
    result["deterministic_content_sha256"] = _fingerprint(profiles, aggregate)

    resolved_output.mkdir(parents=True, exist_ok=True)
    _write_json(resolved_output / "metrics.json", result)
    _write_manifest(resolved_output / "manifest.csv", profiles)
    _write_report(resolved_output / "eda_report.md", result)
    return result

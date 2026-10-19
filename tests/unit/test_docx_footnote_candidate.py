"""Public synthetic OPC packages only; no private corpus or Gold dependencies."""

from __future__ import annotations

import json
import sys
import zipfile
from contextlib import nullcontext
from dataclasses import replace

import pytest

from doclens_ai.extraction import docx_footnote_candidate as candidate
from doclens_ai.extraction.docx_footnote_candidate import FootnoteLimits, extract_footnotes
from doclens_ai.extraction.fidelity_gold import FidelityRecord, fingerprint
from doclens_ai.extraction.parsers import extract_file
from tools.docx_footnote_candidate import benchmark

W = candidate.W[1:-1]
R = candidate.REL[:-1]
P = candidate.PKG[1:-1]


def paragraph(text="Synthetic note 007.50", extra=""):
    return f"<w:p><w:r><w:t>{text}</w:t></w:r>{extra}</w:p>"


def note(key="7", text="Synthetic note 007.50", *, extra="", kind=""):
    return f'<w:footnote w:id="{key}" {kind}>{paragraph(text)}{extra}</w:footnote>'


def package(tmp_path, *, notes=None, references=("7",), moved=False, changes=None):
    path = tmp_path / "synthetic.docx"
    main = "parts/body.xml" if moved else "word/document.xml"
    target = "../notes/actual.xml" if moved else "footnotes.xml"
    part = "notes/actual.xml" if moved else "word/footnotes.xml"
    main_rels = "parts/_rels/body.xml.rels" if moved else "word/_rels/document.xml.rels"
    body = paragraph(
        "Synthetic body 42",
        "".join(f'<w:r><w:footnoteReference w:id="{r}"/></w:r>' for r in references),
    )
    content = {
        "[Content_Types].xml": (
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" '
            'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            f'<Override PartName="/{main}" ContentType="application/vnd.openxmlformats-'
            'officedocument.wordprocessingml.document.main+xml"/>'
            f'<Override PartName="/{part}" ContentType="application/vnd.openxmlformats-'
            'officedocument.wordprocessingml.footnotes+xml"/></Types>'
        ),
        "_rels/.rels": (
            f'<Relationships xmlns="{P}"><Relationship Id="main" '
            f'Type="{candidate.REL}officeDocument" Target="{main}"/></Relationships>'
        ),
        main: (
            f'<w:document xmlns:w="{W}" xmlns:r="{R}">'
            f"<w:body>{body}<w:sectPr/></w:body></w:document>"
        ),
        main_rels: (
            f'<Relationships xmlns="{P}"><Relationship Id="notes" '
            f'Type="{candidate.REL}footnotes" Target="{target}"/></Relationships>'
        ),
        part: (
            f'<w:footnotes xmlns:w="{W}" xmlns:r="{R}">'
            f"{note() if notes is None else notes}</w:footnotes>"
        ),
    }
    content.update(changes or {})
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for name, value in content.items():
            if value is not None:
                z.writestr(name, value)
    return path


def test_no_footnotes(tmp_path):
    path = package(
        tmp_path,
        references=(),
        changes={"word/_rels/document.xml.rels": None, "word/footnotes.xml": None},
    )
    result = extract_footnotes(path)
    assert result.status.value == "success" and not result.notes and not result.references
    assert result.schema_version == "docx-footnotes-offline-v1"


@pytest.mark.parametrize("moved", [False, True])
def test_relationship_discovery_and_truthful_identity(tmp_path, moved):
    result = extract_footnotes(package(tmp_path, moved=moved))
    assert result.status.value == "needs_review"
    n = result.notes[0]
    assert n.note_id == "7" and n.note_element_index == 0 and n.gold_part_index is None
    assert n.part_uri == ("/notes/actual.xml" if moved else "/word/footnotes.xml")
    assert n.paragraphs[0].text == "Synthetic note 007.50" and n.complete
    assert result.references[0].baseline_block_index == 0
    assert result.references[0].paragraph_xml_path == (0,)


def test_multiple_notes_repeated_references_and_duplicate_text(tmp_path):
    path = package(
        tmp_path,
        notes=note("8") + note("7", extra=paragraph("Second paragraph")),
        references=("7", "8", "7"),
    )
    result = extract_footnotes(path)
    assert [n.note_id for n in result.notes] == ["8", "7"]
    assert [n.reference_ordinals for n in result.notes] == [(1,), (0, 2)]
    assert [r.ordinal for r in result.references] == [0, 1, 2]
    assert len(result.notes[1].paragraphs) == 2
    assert result.notes[0].paragraphs[0].text == result.notes[1].paragraphs[0].text
    assert len(set(r.body_xml_path for r in result.references)) == 3


@pytest.mark.parametrize("kind", ["separator", "continuationSeparator", "continuationNotice"])
def test_separator_exclusion_is_by_type_not_assumed_id(tmp_path, kind):
    result = extract_footnotes(
        package(tmp_path, notes=note("0", kind=f'w:type="{kind}"') + note("7"))
    )
    assert [n.note_id for n in result.notes] == ["7"]
    normal = extract_footnotes(package(tmp_path, notes=note("0"), references=("0",)))
    assert normal.notes[0].note_id == "0" and normal.notes[0].complete


def test_unicode_numbers_hyperlinks_and_controls(tmp_path):
    extra = (
        '<w:p><w:hyperlink r:id="h"><w:r><w:t>Link label</w:t><w:tab/>'
        "<w:t>Қазақ Русский 001.50</w:t><w:br/><w:t>tail</w:t></w:r></w:hyperlink></w:p>"
    )
    rels = (
        f'<Relationships xmlns="{P}"><Relationship Id="h" '
        f'Type="{candidate.REL}hyperlink" Target="https://example.invalid/never-fetched" '
        'TargetMode="External"/></Relationships>'
    )
    result = extract_footnotes(
        package(tmp_path, notes=note(extra=extra), changes={"word/_rels/footnotes.xml.rels": rels})
    )
    p = result.notes[0].paragraphs[1]
    assert p.text == "Link label\tҚазақ Русский 001.50\ntail"
    assert p.hyperlink_relationship_ids == ("h",) and p.complete
    assert "never-fetched" not in json.dumps(result.to_dict())


@pytest.mark.parametrize("extra", ["<w:tbl/>", "<w:sdt/>", "<w:altChunk/>"])
def test_unsupported_note_blocks_never_flattened(tmp_path, extra):
    result = extract_footnotes(package(tmp_path, notes=note(extra=extra)))
    assert result.status.value == "needs_review" and not result.notes[0].complete
    assert "unsupported_note_structure" in result.issues
    assert len(result.notes[0].paragraphs) == 1


@pytest.mark.parametrize(
    "extra",
    [
        "<w:r><w:instrText>PRIVATE_FIELD</w:instrText></w:r>",
        "<w:drawing/>",
        "<w:ins><w:r><w:t>Hidden revision</w:t></w:r></w:ins>",
        '<w:hyperlink r:id="missing"><w:r><w:t>Label</w:t></w:r></w:hyperlink>',
    ],
)
def test_complex_paragraphs_incomplete_without_false_success(tmp_path, extra):
    xml = f'<w:footnote w:id="7">{paragraph(extra=extra)}</w:footnote>'
    result = extract_footnotes(package(tmp_path, notes=xml))
    assert not result.notes[0].complete and not result.notes[0].paragraphs[0].complete
    assert "PRIVATE_FIELD" not in json.dumps(result.to_dict())


def test_dangling_orphan_and_separator_references_are_explicit(tmp_path):
    result = extract_footnotes(
        package(
            tmp_path,
            notes=note("7") + note("9") + note("0", kind='w:type="separator"'),
            references=("7", "99", "0"),
        )
    )
    assert {"dangling_reference", "unreferenced_note", "separator_reference"} <= set(result.issues)
    assert result.notes[0].complete and not result.notes[1].complete
    assert len(result.references) == 3


@pytest.mark.parametrize(
    "changes,issue",
    [
        ({"word/_rels/document.xml.rels": None}, "missing_footnote_relationship"),
        ({"word/footnotes.xml": None}, "invalid_package"),
        ({"word/_rels/document.xml.rels": "bad"}, "invalid_package"),
        ({"_rels/.rels": None}, "invalid_package"),
        (
            {
                "_rels/.rels": (
                    f'<Relationships xmlns="{P}"><Relationship Id="m" '
                    'Target="https://example.invalid/secret" '
                    f'Type="{candidate.REL}officeDocument" TargetMode="External"/></Relationships>'
                )
            },
            "invalid_main_relationship",
        ),
        (
            {
                "word/_rels/document.xml.rels": (
                    f'<Relationships xmlns="{P}"><Relationship Id="f" Target="../../secret.xml" '
                    f'Type="{candidate.REL}footnotes"/></Relationships>'
                )
            },
            "invalid_relationship",
        ),
    ],
)
def test_missing_malformed_external_relationships_fail_closed(tmp_path, changes, issue):
    result = extract_footnotes(package(tmp_path, changes=changes))
    assert result.status.value == "failed" and result.issues == (issue,) and not result.notes
    assert "secret" not in json.dumps(result.to_dict())


@pytest.mark.parametrize(
    "xml",
    [b'<!DOCTYPE x [<!ENTITY y "PRIVATE">]><x/>', "<!DOCTYPE x><x/>".encode("utf-16"), b"<bad>"],
)
def test_corrupt_and_entity_xml_never_echoed(tmp_path, xml):
    result = extract_footnotes(package(tmp_path, changes={"word/footnotes.xml": xml}))
    assert result.status.value == "failed" and not result.notes
    assert "PRIVATE" not in json.dumps(result.to_dict())


@pytest.mark.parametrize(
    "field,value",
    [
        ("max_file_bytes", 1),
        ("max_entries", 1),
        ("max_entry_bytes", 1),
        ("max_expanded_bytes", 1),
        ("max_ratio", 1),
        ("max_nodes", 1),
        ("max_depth", 1),
        ("max_notes", 1),
        ("max_references", 1),
        ("max_paragraphs", 1),
        ("max_characters", 1),
    ],
)
def test_all_resource_limits_reject_partial_output(tmp_path, field, value):
    path = package(
        tmp_path, notes=note("7", extra=paragraph("extra")) + note("8"), references=("7", "8")
    )
    result = extract_footnotes(path, limits=replace(FootnoteLimits(), **{field: value}))
    assert result.status.value == "failed" and not result.notes and not result.references


def test_logical_timeout_and_source_handles_cleaned(tmp_path, monkeypatch):
    path = package(tmp_path)
    clock = iter([0, 20])
    monkeypatch.setattr(candidate.time, "monotonic", lambda: next(clock))
    result = extract_footnotes(path)
    assert result.issues == ("timeout",) and not result.notes
    path.rename(tmp_path / "renamed.docx")


def test_duplicate_ids_and_corrupt_zip_fail_without_notes(tmp_path):
    result = extract_footnotes(package(tmp_path, notes=note("7") + note("07")))
    assert result.issues == ("duplicate_note_id",) and not result.notes
    path = tmp_path / "corrupt.docx"
    path.write_bytes(b"not a ZIP PRIVATE")
    assert extract_footnotes(path).issues == ("invalid_package",)


def test_raw_and_production_contract_are_unchanged(tmp_path):
    path = package(tmp_path)
    original = path.read_bytes()
    before = extract_file(path)
    result = extract_footnotes(path)
    assert result.notes[0].complete
    assert before.to_dict() == extract_file(path).to_dict()
    assert path.read_bytes() == original
    assert not hasattr(before, "notes") and not hasattr(result, "blocks")
    assert before.status.value == "needs_review"


def test_independent_provenance_rejects_wrong_reference_and_note_addresses(tmp_path):
    path = package(tmp_path)
    a = json.loads(json.dumps(extract_footnotes(path).to_dict()))
    raw = {"ancillary": a}
    assert benchmark.verify_ancillary(path, raw) == {
        "status": "EVALUATED",
        "notes": 1,
        "references": 1,
        "paragraphs": 1,
    }
    a["notes"][0]["note_id"] = "8"
    with pytest.raises(AssertionError):
        benchmark.verify_ancillary(path, raw)


def synthetic_gold(path):
    result = extract_file(path)
    data = {
        "schema_version": "1.0",
        "dataset_revision": "corpus-v1-fidelity-v1",
        "dataset_kind": "synthetic",
        "review_basis": "synthetic_fixture",
        "document_id": "c1-0001",
        "source_sha256": result.source_sha256,
        "source_size_bytes": result.source_size_bytes,
        "document_format": "docx",
        "review_status": "reviewed",
        "reviewer_count": 1,
        "reviewed_on": "2026-01-01",
        "anchors": [
            {
                "anchor_id": "a1",
                "text": "Synthetic note 007.50",
                "source": {"region": "footnote", "part_index": 0},
            }
        ],
        "order_pairs": [],
        "table_cells": [],
        "table_relationships": [],
        "cells": [],
    }
    return FidelityRecord.model_validate_json(json.dumps(data)), result


def test_recovery_separate_from_frozen_scores_and_gold_provenance(tmp_path):
    path = package(tmp_path)
    record, base = synthetic_gold(path)
    raw = {
        "baseline": base.to_dict(),
        "ancillary": json.loads(json.dumps(extract_footnotes(path).to_dict())),
        "baseline_after_fingerprint": fingerprint(base.to_dict()),
        "effective_status": "needs_review",
    }
    report = benchmark.summarize(record, raw)
    assert report["footnote_text"] == {"passed": 1, "evaluated": 1, "label_total": 1}
    assert report["frozen_candidate_scores"] == report["baseline_scores"]
    assert report["baseline_scores"]["metrics"]["required_text_recall"]["passed"] == 0
    assert report["reviewed_footnote_provenance"]["status"] == "NOT_EVALUATED"
    assert report["numeric_bearing"] == {"passed": 0, "evaluated": 1, "label_total": 1}
    assert report["supplemented_numeric_bearing"] == {"passed": 1, "evaluated": 1, "label_total": 1}
    assert "Synthetic note" not in json.dumps(report) and "/word/" not in json.dumps(report)
    raw["ancillary"]["notes"][0]["complete"] = False
    assert benchmark.summarize(record, raw)["footnote_text"]["passed"] == 0
    assert benchmark.summarize(record, None)["footnote_text"]["passed"] is None
    raw["effective_status"] = "success"
    assert benchmark.summarize(record, raw)["regressions"]


@pytest.mark.parametrize("mode", ["timeout", "memory_limit", "worker_failed"])
def test_isolated_guard_failure_and_cleanup(tmp_path, mode):
    guards = {
        **benchmark.GUARDS,
        "timeout_seconds": 0.05 if mode == "timeout" else 5,
        "max_process_bytes": 1 if mode == "memory_limit" else benchmark.GUARDS["max_process_bytes"],
    }
    command = [
        sys.executable,
        "-B",
        "-c",
        "import time; time.sleep(10)" if mode != "worker_failed" else "raise SystemExit(3)",
    ]
    runtime = benchmark.monitor(command, tmp_path, guards=guards)
    assert runtime["failure"] == mode and runtime["cleanup"]
    assert not (tmp_path / "result.json").exists() and not (tmp_path / "result.partial").exists()


def test_capture_budget_cleanup_and_worker_success(tmp_path, monkeypatch):
    monkeypatch.setitem(benchmark.GUARDS, "max_capture_bytes", 5)
    with pytest.raises(ValueError):
        benchmark.save(tmp_path / "result.json", {"private": "text"})
    assert not (tmp_path / "result.partial").exists() and not (tmp_path / "result.json").exists()
    monkeypatch.setitem(benchmark.GUARDS, "max_capture_bytes", 64 * 1024**2)
    assert benchmark.worker(package(tmp_path), tmp_path) == 0
    raw = json.loads((tmp_path / "result.json").read_text())
    assert raw["effective_status"] == "needs_review"
    assert raw["baseline_after_fingerprint"] == fingerprint(raw["baseline"])


@pytest.mark.parametrize("value", [0, -1, float("inf"), float("nan"), True, 1.5, "1"])
def test_invalid_limits_rejected(value):
    with pytest.raises(ValueError):
        FootnoteLimits(max_notes=value)


def test_body_table_reference_xml_paths_without_invented_cell_coordinates(tmp_path):
    body = (
        f'<w:document xmlns:w="{W}"><w:body><w:sdt/><w:p/>'
        '<w:tbl><w:tr><w:tc><w:p><w:r><w:footnoteReference w:id="7"/>'
        "</w:r></w:p></w:tc></w:tr></w:tbl><w:sectPr/></w:body></w:document>"
    )
    path = package(tmp_path, changes={"word/document.xml": body})
    result = extract_footnotes(path)
    ref = result.references[0]
    assert ref.body_child_index == 2 and ref.baseline_block_index == 1
    assert ref.body_xml_path == (2, 0, 0, 0, 0, 0)
    assert ref.paragraph_xml_path == (2, 0, 0, 0)
    assert not hasattr(ref, "cell_coordinate")
    assert benchmark.verify_ancillary(path, {"ancillary": result.to_dict()})["references"] == 1


@pytest.mark.parametrize(
    "variant", ["duplicate_member", "traversal", "external_notes", "duplicate_relation"]
)
def test_ambiguous_archive_or_note_relationships_fail(tmp_path, variant):
    if variant == "external_notes":
        rels = (
            f'<Relationships xmlns="{P}"><Relationship Id="f" '
            f'Type="{candidate.REL}footnotes" Target="https://example.invalid/notes" '
            'TargetMode="External"/></Relationships>'
        )
        path = package(tmp_path, changes={"word/_rels/document.xml.rels": rels})
    elif variant == "duplicate_relation":
        link = f'<Relationship Id="f" Type="{candidate.REL}footnotes" Target="footnotes.xml"/>'
        path = package(
            tmp_path,
            changes={
                "word/_rels/document.xml.rels": (
                    f'<Relationships xmlns="{P}">{link}{link}</Relationships>'
                )
            },
        )
    else:
        path = package(tmp_path)
        with zipfile.ZipFile(path, "a") as z:
            with pytest.warns(UserWarning) if variant == "duplicate_member" else nullcontext():
                z.writestr(
                    "word/document.xml" if variant == "duplicate_member" else "../escape", "x"
                )
    result = extract_footnotes(path)
    assert result.status.value == "failed" and not result.notes


def test_worker_logging_state_and_atomic_partial_failure(tmp_path, monkeypatch):
    import logging

    previous = logging.root.manager.disable
    assert benchmark.worker(package(tmp_path), tmp_path) == 0
    assert logging.root.manager.disable == previous
    (tmp_path / "result.json").unlink()

    def reject(*args):
        raise PermissionError("PRIVATE PATH")

    monkeypatch.setattr(benchmark.os, "replace", reject)
    assert benchmark.worker(tmp_path / "synthetic.docx", tmp_path) == 2
    assert logging.root.manager.disable == previous
    assert not (tmp_path / "result.partial").exists()
    assert not (tmp_path / "result.json").exists()


def test_monitor_does_not_start_process_when_journal_cannot_open(tmp_path, monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("Must open journal before spawning")

    monkeypatch.setattr(benchmark.subprocess, "Popen", forbidden)
    with pytest.raises(OSError):
        benchmark.monitor([sys.executable], tmp_path / "missing")


def test_restore_docx_preserves_tables_spans_warnings_and_source_locations(tmp_path):
    from docx import Document

    path = tmp_path / "synthetic-table.docx"
    doc = Document()
    doc.add_paragraph("Synthetic paragraph 001")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Synthetic merged 007.50"
    table.cell(1, 0).text = "Left"
    table.cell(1, 1).text = "Right"
    doc.sections[0].header.paragraphs[0].text = "Synthetic ancillary header"
    doc.save(path)
    base = extract_file(path)
    restored = benchmark.restore_docx(base.to_dict())
    assert restored == base
    assert restored.to_dict() == base.to_dict()
    assert restored.blocks[1].rows[0].cells[0].column_span == 2
    assert restored.warnings
    bad = base.to_dict()
    bad["blocks"][1]["kind"] = "unsupported"
    with pytest.raises(ValueError):
        benchmark.restore_docx(bad)


def test_independent_verifier_excludes_tab_stop_properties_from_text(tmp_path):
    xml = (
        '<w:footnote w:id="7"><w:p><w:pPr><w:tabs>'
        '<w:tab w:val="left" w:pos="1440"/></w:tabs></w:pPr>'
        "<w:r><w:t>Synthetic</w:t><w:tab/><w:t>42</w:t></w:r></w:p></w:footnote>"
    )
    path = package(tmp_path, notes=xml)
    result = extract_footnotes(path)
    assert result.notes[0].paragraphs[0].text == "Synthetic\t42"
    assert benchmark.verify_ancillary(path, {"ancillary": result.to_dict()})["paragraphs"] == 1


def test_missing_capture_retains_every_reviewed_denominator_as_unknown(tmp_path):
    record, _ = synthetic_gold(package(tmp_path))
    summary = benchmark.summarize(record, None)
    assert summary["reviewed_denominators"] == {
        "anchors": 1,
        "provenance": 1,
        "order": 0,
        "table_cells": 0,
        "relations": 0,
        "value_checks": 0,
        "footnotes": 1,
        "numeric_bearing": 1,
    }
    assert (
        summary["numeric_bearing"]
        == summary["supplemented_numeric_bearing"]
        == {"passed": None, "evaluated": 0, "label_total": 1}
    )
    assert summary["recovered"] is summary["lost"] is None
    assert summary["reviewed_footnote_provenance"]["label_total"] == 1
    assert summary["regressions"] and summary["status"] == "failed"

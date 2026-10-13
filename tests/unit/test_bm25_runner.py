from __future__ import annotations

import csv
import json
from dataclasses import asdict, replace
from pathlib import Path

import pytest
from docx import Document

from doclens_ai.chunking.core import chunk_document
from doclens_ai.chunking.models import Chunk, ChunkSourceSpan
from doclens_ai.evaluation.gold import (
    AdjudicationState,
    DatasetSplit,
    EvidenceAnchor,
    EvidenceLocator,
    GoldExample,
    Language,
    QuestionCategory,
    ReviewStatus,
    load_gold_jsonl,
    validate_gold_dataset,
)
from doclens_ai.evaluation.metrics import evaluate_retrieval, load_ranked_results_jsonl
from doclens_ai.extraction.models import (
    DocumentFormat,
    ErrorCategory,
    ExtractionResult,
    ExtractionStatus,
    SourceLocation,
    TextBlock,
)
from doclens_ai.retrieval_baseline import BM25Document, BM25Index
from doclens_ai.retrieval_baseline.__main__ import main
from doclens_ai.retrieval_baseline.runner import (
    FROZEN_CHUNKING,
    BaselineError,
    build_corpus,
    chunk_evidence,
    file_sha256,
    results_jsonl,
    retrieve_queries,
    run_baseline,
    verify_corpus,
)

REVISION = "synthetic-retrieval-gold-v1"


@pytest.fixture
def fixture_corpus(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.chdir(tmp_path)
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    for name, text in (
        ("z.docx", "orchard iris budget 2042"),
        ("a.docx", "orchard orchard iris budget 2042 extra"),
    ):
        doc = Document()
        doc.add_paragraph(text)
        doc.save(corpus / name)
    manifest = tmp_path / "manifest.csv"
    with manifest.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(
            stream, fieldnames=["document_id", "sha256", "format", "size_bytes"]
        )
        writer.writeheader()
        # Order intentionally disagrees with local filename order.
        for number, name in enumerate(("z.docx", "a.docx"), start=1):
            path = corpus / name
            writer.writerow(
                {
                    "document_id": f"fixture-{number}",
                    "sha256": file_sha256(path),
                    "format": "docx",
                    "size_bytes": path.stat().st_size,
                }
            )
    documents = verify_corpus(corpus, manifest, expected_source_count=2)
    examples = tuple(
        GoldExample(
            schema_version="1.0",
            question_id=f"fixture-q-{number}",
            dataset_revision=REVISION,
            synthetic_fixture=True,
            document_family_id=f"fixture-family-{number}",
            language=Language.EN,
            question="orchard budget",
            category=QuestionCategory.FACT,
            answerable=True,
            source_document_sha256=document.sha256,
            evidence=(
                EvidenceAnchor(
                    f"fixture-e-{number}",
                    document.sha256,
                    DocumentFormat.DOCX,
                    EvidenceLocator(block_index=0),
                    0,
                    7,
                ),
            ),
            split=DatasetSplit.DEV if number == 1 else DatasetSplit.TEST,
            review_status=ReviewStatus.REVIEWED,
            reviewer_count=1,
            adjudication_state=AdjudicationState.NOT_REQUIRED,
        )
        for number, document in enumerate(documents, start=1)
    )
    gold = tmp_path / "gold.jsonl"
    gold.write_text(
        "".join(json.dumps(item.to_dict()) + "\n" for item in examples), encoding="utf-8"
    )
    return corpus, manifest, gold, documents


def test_manifest_identity_not_filename_order(fixture_corpus) -> None:
    corpus, manifest, _, documents = fixture_corpus
    assert documents[0].path.name == "z.docx"
    before = build_corpus(documents)
    (corpus / "z.docx").rename(corpus / "renamed.docx")
    after = build_corpus(verify_corpus(corpus, manifest, expected_source_count=2))
    assert before.chunks == after.chunks
    assert before.metadata == after.metadata


@pytest.mark.parametrize("change", ["modified", "missing", "duplicate", "extra", "extension"])
def test_corpus_integrity_blockers(fixture_corpus, change: str) -> None:
    corpus, manifest, _, _ = fixture_corpus
    path = corpus / "a.docx"
    if change == "modified":
        path.write_bytes(b"changed bytes")
    elif change == "missing":
        path.unlink()
    elif change == "duplicate":
        (corpus / "duplicate.docx").write_bytes(path.read_bytes())
    elif change == "extra":
        (corpus / "unexpected.txt").write_text("unexpected")
    else:
        path.rename(corpus / "wrong.pdf")
    with pytest.raises(BaselineError):
        verify_corpus(corpus, manifest, expected_source_count=2)


@pytest.mark.parametrize(
    "fmt,location",
    [
        (DocumentFormat.PDF, SourceLocation(page_number=2, block_index=1)),
        (DocumentFormat.DOCX, SourceLocation(block_index=3)),
        (
            DocumentFormat.DOCX,
            SourceLocation(block_index=3, table_index=0, row_index=1, column_index=2),
        ),
        (
            DocumentFormat.XLSX,
            SourceLocation(
                sheet_index=0,
                sheet_name="Fixture",
                cell_coordinate="B2",
                row_index=2,
                column_index=2,
            ),
        ),
    ],
)
def test_provenance_coordinates_and_spans_are_copied(fmt, location) -> None:
    span = ChunkSourceSpan("unit-1", location, 10, 20, 40)
    chunk = Chunk(
        "chunk-1",
        "doc-1",
        "a" * 64,
        fmt,
        FROZEN_CHUNKING.strategy,
        FROZEN_CHUNKING.version,
        FROZEN_CHUNKING.fingerprint(),
        7,
        "fixture",
        (span, span),
    )
    evidence = chunk_evidence(chunk)
    assert len(evidence) == 1
    assert evidence[0].locator == EvidenceLocator(**asdict(location))
    assert (evidence[0].source_start, evidence[0].source_end) == (10, 20)
    assert evidence[0].document_sha256 == chunk.source_sha256
    assert evidence[0].document_format == fmt
    with pytest.raises(BaselineError, match="provenance"):
        chunk_evidence(replace(chunk, provenance=()))


def test_extraction_gate_and_failures_are_isolated(fixture_corpus, monkeypatch) -> None:
    _, _, _, documents = fixture_corpus
    success = ExtractionResult(
        "1.0",
        DocumentFormat.DOCX,
        ExtractionStatus.SUCCESS,
        documents[0].sha256,
        documents[0].size_bytes,
        blocks=(TextBlock(0, "orchard", SourceLocation(block_index=0)),),
    )
    failed = replace(
        success,
        status=ExtractionStatus.FAILED,
        source_sha256=documents[1].sha256,
        blocks=(),
        error_category=ErrorCategory.CORRUPT_DOCUMENT,
    )
    monkeypatch.setattr(
        "doclens_ai.retrieval_baseline.runner.extract_file",
        lambda path: success if path == documents[0].path else failed,
    )
    corpus = build_corpus(documents)
    assert len(corpus.chunks) == 1
    assert corpus.metadata["eligible_count"] == 1
    assert corpus.metadata["exclusions"][0]["reason"] == "extraction_failed"
    assert corpus.metadata["exclusions"][0]["error_category"] == "corrupt_document"
    review = replace(failed, status=ExtractionStatus.NEEDS_REVIEW, error_category=None)
    monkeypatch.setattr(
        "doclens_ai.retrieval_baseline.runner.extract_file",
        lambda path: success if path == documents[0].path else review,
    )
    corpus = build_corpus(documents)
    assert len(corpus.chunks) == 1
    assert corpus.metadata["exclusions"][0]["reason"] == "extraction_needs_review"


def test_zero_result_records_valid_contract_and_security(fixture_corpus, tmp_path) -> None:
    _, _, _, documents = fixture_corpus
    corpus = build_corpus(documents)
    chunks = {item.chunk_id: item for item in corpus.chunks}
    index = BM25Index(BM25Document(item.chunk_id, item.text) for item in corpus.chunks)
    records, latencies = retrieve_queries(
        index,
        chunks,
        [("fixture-q-1", "unknownterm"), ("fixture-q-2", "orchard")],
        dataset_revision=REVISION,
    )
    path = tmp_path / "ranked.jsonl"
    path.write_text(results_jsonl(records), encoding="utf-8")
    parsed = load_ranked_results_jsonl(path)
    assert len(latencies) == len(parsed) == 2
    assert parsed[0].results == ()
    assert parsed[1].results[0].runtime_chunk_id in chunks
    raw = json.loads(path.read_text().splitlines()[1])
    assert set(raw["results"][0]) == {"rank", "runtime_chunk_id", "evidence"}
    assert "orchard" not in path.read_text()
    assert all(item.authorized_workspace_id is None for item in parsed)


def test_two_runs_stable_fingerprints_and_no_gold_ranking_leakage(fixture_corpus, tmp_path) -> None:
    corpus, manifest, gold_path, _ = fixture_corpus
    gold = load_gold_jsonl(gold_path)
    outputs = [tmp_path / "evaluation-output" / name for name in ("run-1", "run-2", "run-3")]
    for output in outputs[:2]:
        run_baseline(
            corpus_dir=corpus,
            manifest_path=manifest,
            gold_path=gold_path,
            output_dir=output,
            expected_gold_fingerprint=gold.fingerprint,
            expected_source_count=2,
        )
    for name in ("ranked_results.jsonl", "quality_report.json", "experiment.json"):
        assert (outputs[0] / name).read_bytes() == (outputs[1] / name).read_bytes()
    report = json.loads((outputs[0] / "quality_report.json").read_text())
    assert report["security"]["status"] == "not_evaluated"
    assert report["security"]["passed"] is None
    ranked = load_ranked_results_jsonl(outputs[0] / "ranked_results.jsonl")
    # Different Gold document labels and dev/test assignments do not change either search.
    assert ranked[0].results == ranked[1].results
    changed = tuple(
        replace(
            item,
            source_document_sha256=gold.examples[1 - number].source_document_sha256,
            evidence=(),
            answerable=False,
            category=QuestionCategory.UNANSWERABLE,
            split=DatasetSplit.TEST if item.split is DatasetSplit.DEV else DatasetSplit.DEV,
        )
        for number, item in enumerate(gold.examples)
    )
    mutated = validate_gold_dataset(changed)
    gold_path.write_text(
        "".join(json.dumps(item.to_dict()) + "\n" for item in changed), encoding="utf-8"
    )
    run_baseline(
        corpus_dir=corpus,
        manifest_path=manifest,
        gold_path=gold_path,
        output_dir=outputs[2],
        expected_gold_fingerprint=mutated.fingerprint,
        expected_source_count=2,
    )
    assert (outputs[0] / "ranked_results.jsonl").read_bytes() == (
        outputs[2] / "ranked_results.jsonl"
    ).read_bytes()
    assert (
        evaluate_retrieval(mutated, ranked).to_dict()["overall"]["unanswerable"][
            "queries_with_results"
        ]
        == 2
    )
    metadata = (outputs[0] / "experiment.json").read_text() + (
        outputs[0] / "quality_report.json"
    ).read_text()
    for private in ("orchard", "iris", "a.docx", "z.docx", str(corpus), "budget"):
        assert private not in metadata
    with pytest.raises(BaselineError, match="already exists"):
        run_baseline(
            corpus_dir=corpus, manifest_path=manifest, gold_path=gold_path, output_dir=outputs[0]
        )


def test_gold_fingerprint_and_output_root_fail_closed(fixture_corpus, tmp_path) -> None:
    corpus, manifest, gold, _ = fixture_corpus
    with pytest.raises(BaselineError, match="fingerprint"):
        run_baseline(
            corpus_dir=corpus,
            manifest_path=manifest,
            gold_path=gold,
            output_dir=tmp_path / "evaluation-output/bad-gold",
            expected_source_count=2,
        )
    with pytest.raises(BaselineError, match="subdirectory"):
        run_baseline(
            corpus_dir=corpus,
            manifest_path=manifest,
            gold_path=gold,
            output_dir=tmp_path / "outside",
        )
    with pytest.raises(SystemExit) as exc:
        main(["--output-dir", str(tmp_path / "outside")])
    assert exc.value.code == 2


def test_provenance_from_existing_chunker_is_retained(fixture_corpus) -> None:
    _, _, _, documents = fixture_corpus
    extraction = ExtractionResult(
        "1.0",
        DocumentFormat.PDF,
        ExtractionStatus.SUCCESS,
        documents[0].sha256,
        documents[0].size_bytes,
        blocks=(TextBlock(0, "iris " * 400, SourceLocation(page_number=1, block_index=0)),),
    )
    chunks = chunk_document(extraction, document_id="fixture-1", config=FROZEN_CHUNKING).chunks
    assert len(chunks) > 1
    for chunk in chunks:
        assert chunk.config_fingerprint == FROZEN_CHUNKING.fingerprint()
        for evidence, span in zip(chunk_evidence(chunk), chunk.provenance, strict=True):
            assert evidence.source_start == span.source_start
            assert evidence.source_end == span.source_end

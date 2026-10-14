from __future__ import annotations

import csv
import hashlib
import json
from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from doclens_ai.extraction.fidelity_eval import evaluate, score_record, verify_corpus
from doclens_ai.extraction.fidelity_gold import FidelityRecord, GoldError, load_gold
from doclens_ai.extraction.models import (
    DocumentFormat,
    ExtractionResult,
    ExtractionStatus,
    SourceLocation,
    SpreadsheetCell,
    SpreadsheetSheet,
    TableBlock,
    TableCell,
    TableRow,
    TextBlock,
)
from doclens_ai.extraction.parsers import ExtractionLimits


def _gold(fmt: str = "pdf") -> dict:
    source = (
        {"region": "body", "page_number": 1}
        if fmt == "pdf"
        else (
            {"region": "body", "block_index": 0}
            if fmt == "docx"
            else {"region": "body", "sheet_index": 0, "cell_coordinate": "A1"}
        )
    )
    return {
        "schema_version": "1.0",
        "dataset_revision": "corpus-v1-fidelity-v1",
        "dataset_kind": "synthetic",
        "document_id": "c1-0001",
        "source_sha256": "a" * 64,
        "source_size_bytes": 100,
        "document_format": fmt,
        "review_status": "reviewed",
        "review_basis": "synthetic_fixture",
        "reviewer_count": 1,
        "reviewed_on": "2026-01-01",
        "anchors": [{"anchor_id": "a1", "text": "alpha", "source": source}],
        "order_pairs": [],
        "table_cells": [],
        "table_relationships": [],
        "cells": [],
    }


def _record(data: dict) -> FidelityRecord:
    return FidelityRecord.model_validate_json(json.dumps(data))


def _result(fmt: str = "pdf", **kwargs) -> ExtractionResult:
    return ExtractionResult(
        schema_version="1.0",
        document_format=DocumentFormat(fmt),
        status=ExtractionStatus.SUCCESS,
        source_sha256="a" * 64,
        source_size_bytes=100,
        **kwargs,
    )


@pytest.mark.parametrize(
    "field,value",
    [
        ("schema_version", "2.0"),
        ("dataset_revision", "unknown"),
        ("review_status", "draft"),
        ("reviewer_count", 0),
        ("reviewer_count", True),
        ("reviewer_count", "1"),
        ("review_basis", "original_document"),
        ("document_id", "private-filename.pdf"),
        ("source_sha256", "bad"),
        ("source_size_bytes", -1),
        ("reviewed_on", "3000-01-01"),
        ("unexpected", "SECRET"),
        ("document_format", "csv"),
        ("anchors", []),
    ],
)
def test_invalid_gold_is_rejected_without_echo(tmp_path: Path, field, value) -> None:
    data = _gold()
    data[field] = value
    path = tmp_path / "gold.jsonl"
    path.write_text(json.dumps(data), encoding="utf-8")
    with pytest.raises(GoldError) as exc:
        load_gold(path, {"c1-0001": ("a" * 64, 100, "pdf")}, allow_synthetic=True)
    assert "SECRET" not in str(exc.value)
    assert "private-filename" not in str(exc.value)


@pytest.mark.parametrize(
    "source",
    [
        {"region": "body"},
        {"region": "body", "page_number": 0},
        {"region": "body", "page_number": 1, "sheet_index": 0},
        {"region": "body", "page_number": 1, "row_index": 0},
    ],
)
def test_locator_shapes_rejected(source) -> None:
    data = _gold()
    data["anchors"][0]["source"] = source
    with pytest.raises(ValueError):
        _record(data)


def test_references_cycles_duplicates_and_blank_anchors_rejected() -> None:
    for edit in ("unknown", "cycle", "duplicate", "blank"):
        data = _gold()
        data["anchors"].append({**data["anchors"][0], "anchor_id": "a2", "text": "beta"})
        if edit == "unknown":
            data["order_pairs"] = [{"before": "a1", "after": "a3"}]
        elif edit == "cycle":
            data["order_pairs"] = [{"before": "a1", "after": "a2"}, {"before": "a2", "after": "a1"}]
        elif edit == "duplicate":
            data["anchors"][1]["anchor_id"] = "a1"
        else:
            data["anchors"][0]["text"] = " \n "
        with pytest.raises(ValueError):
            _record(data)


def test_source_mismatch_unreviewed_and_synthetic_opt_in(tmp_path: Path) -> None:
    path = tmp_path / "gold.jsonl"
    path.write_text(json.dumps(_gold()), encoding="utf-8")
    sources = {"c1-0001": ("a" * 64, 100, "pdf")}
    with pytest.raises(GoldError):
        load_gold(path, sources)
    assert len(load_gold(path, sources, allow_synthetic=True)) == 1
    with pytest.raises(GoldError):
        load_gold(path, {"c1-0001": ("b" * 64, 100, "pdf")}, allow_synthetic=True)
    with pytest.raises(GoldError):
        score_record(_record(_gold()), replace(_result(), source_sha256="b" * 64))


def test_pdf_recall_order_provenance_are_distinct() -> None:
    data = _gold()
    data["anchors"].extend(
        [
            {"anchor_id": "a2", "text": "beta", "source": {"region": "body", "page_number": 1}},
            {"anchor_id": "a3", "text": "missing", "source": {"region": "body", "page_number": 2}},
            {
                "anchor_id": "a4",
                "text": "elsewhere",
                "source": {"region": "body", "page_number": 1},
            },
        ]
    )
    data["order_pairs"] = [{"before": "a1", "after": "a2"}]
    result = _result(
        blocks=(
            TextBlock(0, " beta\nalpha ", SourceLocation(page_number=1)),
            TextBlock(1, "elsewhere", SourceLocation(page_number=2)),
        )
    )
    scored = score_record(_record(data), result)
    assert scored["metrics"]["required_text_recall"]["score"] == 0.75
    assert scored["missing_evidence_count"] == 1
    assert scored["metrics"]["provenance"]["score"] == 0.5
    assert scored["metrics"]["reading_order"]["score"] == 0


def test_pdf_text_is_not_structured_table_fidelity() -> None:
    data = _gold()
    data["anchors"][0]["source"].update(row_index=0, column_index=0, block_index=0, table_index=0)
    data["table_cells"] = [{"anchor_id": "a1"}]
    scored = score_record(
        _record(data), _result(blocks=(TextBlock(0, "alpha", SourceLocation(page_number=1)),))
    )
    assert scored["metrics"]["required_text_recall"]["score"] == 1
    assert scored["metrics"]["table_cells"]["score"] == 0


def test_docx_order_merged_cells_relationships_and_ancillary_loss() -> None:
    data = _gold("docx")
    data["anchors"] = [
        {
            "anchor_id": "a1",
            "text": "alpha",
            "source": {
                "region": "body",
                "block_index": 0,
                "table_index": 0,
                "row_index": 0,
                "column_index": 0,
            },
        },
        {
            "anchor_id": "a2",
            "text": "beta",
            "source": {
                "region": "body",
                "block_index": 0,
                "table_index": 0,
                "row_index": 0,
                "column_index": 2,
            },
        },
        {"anchor_id": "a3", "text": "header", "source": {"region": "header", "part_index": 0}},
    ]
    data["table_cells"] = [{"anchor_id": "a1", "column_span": 2}, {"anchor_id": "a2"}]
    data["table_relationships"] = [{"left": "a1", "right": "a2", "relationship": "same_row"}]
    data["order_pairs"] = [{"before": "a1", "after": "a2"}]
    cells = tuple(
        TableCell(
            0,
            col,
            text,
            SourceLocation(block_index=0, table_index=0, row_index=0, column_index=col),
            column_span=span,
        )
        for col, text, span in ((0, "alpha", 2), (2, "beta", 1))
    )
    scored = score_record(
        _record(data),
        _result(
            "docx",
            blocks=(
                TableBlock(
                    0, 0, (TableRow(0, cells),), SourceLocation(block_index=0, table_index=0)
                ),
            ),
        ),
    )
    for metric in ("reading_order", "table_cells", "table_relationships"):
        assert scored["metrics"][metric]["score"] == 1
    assert scored["missing_evidence_count"] == 1


def test_xlsx_formula_null_cache_and_unsupported_display() -> None:
    data = _gold("xlsx")
    data["anchors"][0]["text"] = "1+1"
    data["cells"] = [
        {
            "source": {"region": "body", "sheet_index": 0, "cell_coordinate": "A1"},
            "value": None,
            "formula": "1+1",
            "cached_value": None,
            "check_display": True,
            "displayed_value": "2.00",
        }
    ]
    source = SourceLocation(sheet_index=0, cell_coordinate="A1")
    result = _result(
        "xlsx",
        sheets=(
            SpreadsheetSheet(
                0,
                "SECRET_SHEET",
                "visible",
                (SpreadsheetCell("A1", 0, 0, None, "1+1", None, "n", source),),
                (),
                (),
            ),
        ),
    )
    scored = score_record(_record(data), result)
    for metric in ("stored_values", "formulas", "cached_values", "provenance"):
        assert scored["metrics"][metric]["score"] == 1
    assert scored["metrics"]["displayed_values"] == {
        "status": "NOT_EVALUATED",
        "passed": 0,
        "evaluated": 0,
        "unsupported": 1,
        "score": None,
    }


def _corpus(tmp_path: Path) -> tuple[Path, Path]:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    doc = Document()
    doc.add_paragraph("PRIVATE_CONTENT")
    doc.save(corpus / "PRIVATE_FILENAME.docx")
    book = Workbook()
    book.active.title = "PRIVATE_SHEET"
    book.active["A1"] = "PRIVATE_CELL"
    book.active["A2"] = "=1+1"
    book.save(corpus / "PRIVATE_WORKBOOK.xlsx")
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    with (corpus / "PRIVATE_PDF.pdf").open("wb") as stream:
        writer.write(stream)
    manifest = tmp_path / "manifest.csv"
    with manifest.open("w", encoding="utf-8", newline="") as stream:
        csv_writer = csv.writer(stream)
        csv_writer.writerow(["document_id", "sha256", "size_bytes", "format"])
        for index, path in enumerate(sorted(corpus.iterdir()), 1):
            csv_writer.writerow(
                [
                    f"c1-{index:04}",
                    hashlib.sha256(path.read_bytes()).hexdigest(),
                    path.stat().st_size,
                    path.suffix[1:],
                ]
            )
    return corpus, manifest


def test_real_pipeline_determinism_privacy_status_and_chunks(tmp_path: Path) -> None:
    corpus, manifest = _corpus(tmp_path)
    first = evaluate(corpus, manifest, tmp_path / "one", measure_latency=False)
    second = evaluate(corpus, manifest, tmp_path / "two", measure_latency=False)
    assert first == second
    assert first["automated_by_format"]["docx"]["statuses"]["success"] == 1
    assert first["automated_by_format"]["pdf"]["statuses"]["needs_review"] == 1
    assert first["automated_by_format"]["xlsx"]["statuses"]["needs_review"] == 1
    assert first["gold"]["reviewed_documents"] == 0
    assert all(
        metric["status"] == "NOT_EVALUATED"
        for metric in first["reviewed_fidelity"]["metrics"].values()
    )
    serialized = (tmp_path / "one/metrics.json").read_text()
    assert "PRIVATE" not in serialized and str(tmp_path) not in serialized
    assert sum(doc["chunking"]["status"] == "skipped" for doc in first["documents"]) == 2


@pytest.mark.parametrize("fault", ["change", "extra", "missing", "extension", "duplicate"])
def test_integrity_gate_before_output(tmp_path: Path, fault: str) -> None:
    corpus, manifest = _corpus(tmp_path)
    path = sorted(corpus.iterdir())[0]
    if fault == "change":
        path.write_bytes(b"changed")
    elif fault == "extra":
        (corpus / "extra.txt").write_text("extra")
    elif fault == "missing":
        path.unlink()
    elif fault == "extension":
        contents = path.read_bytes()
        path.unlink()
        path.with_suffix(".pdf").write_bytes(contents)
    else:
        (corpus / "copy.docx").write_bytes(path.read_bytes())
    with pytest.raises(ValueError):
        evaluate(corpus, manifest, tmp_path / "output")
    assert not (tmp_path / "output").exists()


def test_unexpected_parser_failure_is_isolated(tmp_path: Path, monkeypatch) -> None:
    from doclens_ai.extraction import fidelity_eval

    corpus, manifest = _corpus(tmp_path)
    original = fidelity_eval.extract_file

    def extract(path, **kwargs):
        if path.suffix == ".docx":
            raise RuntimeError("PRIVATE_FAILURE")
        return original(path, **kwargs)

    monkeypatch.setattr(fidelity_eval, "extract_file", extract)
    result = evaluate(corpus, manifest, tmp_path / "out")
    assert result["automated_by_format"]["docx"]["statuses"]["failed"] == 1
    assert len(result["documents"]) == 3
    assert "PRIVATE_FAILURE" not in json.dumps(result)


def test_bounds_and_overwrite_protection(tmp_path: Path) -> None:
    corpus, manifest = _corpus(tmp_path)
    with pytest.raises(ValueError, match="bound"):
        verify_corpus(corpus, {}, ExtractionLimits(max_file_bytes=1))
    with pytest.raises(ValueError):
        evaluate(corpus, manifest, corpus / "output")
    evaluate(corpus, manifest, tmp_path / "out")
    with pytest.raises(ValueError):
        evaluate(corpus, manifest, tmp_path / "out")


def test_draft_template_is_private_and_not_loadable(tmp_path: Path) -> None:
    corpus, manifest = _corpus(tmp_path)
    # Use an actual priority opaque ID to exercise template generation.
    manifest.write_text(manifest.read_text().replace("c1-0001", "c1-0006"))
    template = tmp_path / "local-evaluation/fidelity_draft.jsonl"
    evaluate(corpus, manifest, tmp_path / "out", template=template)
    draft = json.loads(template.read_text())
    assert draft["review_status"] == "draft" and draft["reviewer_count"] == 0
    assert draft["anchors"] == []
    with pytest.raises(GoldError):
        load_gold(template, {})


def test_experiments_only_return_metadata(tmp_path: Path) -> None:
    corpus, manifest = _corpus(tmp_path)
    text = manifest.read_text()
    for before, after in (("c1-0001", "c1-0006"), ("c1-0002", "c1-0041"), ("c1-0003", "c1-0067")):
        text = text.replace(before, after)
    manifest.write_text(text)
    report = evaluate(corpus, manifest, tmp_path / "out", experiments=True)
    # pypdf layout raises for an entirely blank page without /Contents;
    # this candidate failure must stay isolated rather than alter production behavior.
    assert [doc["experiment"]["status"] for doc in report["documents"]] == [
        "OBSERVED_NOT_REVIEWED",
        "EXPERIMENT_FAILED",
        "OBSERVED_NOT_REVIEWED",
    ]
    assert "PRIVATE" not in json.dumps(report)

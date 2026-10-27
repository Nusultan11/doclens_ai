from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
from openpyxl import Workbook
from openpyxl.worksheet.table import Table as XlsxTable
from openpyxl.worksheet.table import TableStyleInfo
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from doclens_ai.extraction import (
    ErrorCategory,
    ExtractionConfigurationError,
    ExtractionLimits,
    ExtractionStatus,
    TableBlock,
    TextBlock,
    WarningCode,
    extract_corpus,
    extract_file,
)


def _write_pdf(path: Path, page_texts: list[str]) -> None:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_reference = writer._add_object(font)  # noqa: SLF001
    for text in page_texts:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_reference}
                )
            }
        )
        stream = DecodedStreamObject()
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1"))
        page[NameObject("/Contents")] = writer._add_object(stream)  # noqa: SLF001
    with path.open("wb") as destination:
        writer.write(destination)


def _write_xlsx(path: Path, *, formula: bool = False, extreme: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Metrics"
    sheet.append(["left", "right", "total"])
    sheet.append([1, 2, "=A2+B2" if formula else 3])
    table = XlsxTable(displayName="MetricsTable", ref="A1:C2")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    if extreme:
        sheet.cell(row=1_048_576, column=16_384, value="edge")
    workbook.create_sheet("Second")["B2"] = "other"
    workbook.save(path)


def _set_formula_cache(path: Path, cached_value: str) -> Path:
    rewritten = path.with_suffix(".rewritten.xlsx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(
        rewritten, "w", compression=zipfile.ZIP_DEFLATED
    ) as destination:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                text = text.replace("<f>A2+B2</f><v></v>", f"<f>A2+B2</f><v>{cached_value}</v>")
                data = text.encode("utf-8")
            destination.writestr(item, data)
    return rewritten


def test_pdf_preserves_page_provenance_and_flags_low_text(tmp_path: Path) -> None:
    source = tmp_path / "sample.pdf"
    _write_pdf(source, ["A sufficiently long first page for extraction.", ""])

    result = extract_file(source)

    assert result.status is ExtractionStatus.NEEDS_REVIEW
    assert [block.source.page_number for block in result.blocks] == [1, 2]
    assert all(isinstance(block, TextBlock) for block in result.blocks)
    assert [warning.code for warning in result.warnings].count(
        WarningCode.INSUFFICIENT_TEXT
    ) == 1
    assert WarningCode.PDF_TABLE_STRUCTURE_NOT_PRESERVED in {
        warning.code for warning in result.warnings
    }


def test_docx_preserves_body_order_and_merged_cell_semantics(tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("before")
    table = document.add_table(rows=3, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "merged once"
    table.cell(0, 2).text = "right"
    table.cell(1, 0).text = "a"
    table.cell(1, 1).text = "b"
    vertical = table.cell(1, 2).merge(table.cell(2, 2))
    vertical.text = "vertical once"
    table.cell(2, 0).text = "d"
    table.cell(2, 1).text = "e"
    document.add_paragraph("after")
    document.save(source)

    result = extract_file(source)

    assert result.status is ExtractionStatus.SUCCESS
    assert [block.kind for block in result.blocks] == ["text", "table", "text"]
    assert result.blocks[0].text == "before"
    assert result.blocks[2].text == "after"
    extracted_table = result.blocks[1]
    assert isinstance(extracted_table, TableBlock)
    cells = [cell for row in extracted_table.rows for cell in row.cells]
    assert len(cells) == 7
    assert [cell.text for cell in cells].count("merged once") == 1
    merged_cell = next(cell for cell in cells if cell.text == "merged once")
    assert (merged_cell.row_span, merged_cell.column_span) == (1, 2)
    assert (merged_cell.row_index, merged_cell.column_index) == (0, 0)
    vertical_cell = next(cell for cell in cells if cell.text == "vertical once")
    assert [cell.text for cell in cells].count("vertical once") == 1
    assert (vertical_cell.row_span, vertical_cell.column_span) == (2, 1)
    assert (vertical_cell.row_index, vertical_cell.column_index) == (1, 2)


@pytest.mark.parametrize("node_kind", ["comment", "processing_instruction"])
@pytest.mark.parametrize("unsupported_element", [False, True])
def test_docx_non_element_body_nodes_preserve_text_tables_and_physical_xml_paths(
    tmp_path: Path, node_kind: str, unsupported_element: bool
) -> None:
    source = tmp_path / "non-elements.docx"
    document = Document()
    document.add_paragraph("before −001.50")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "merged header"
    table.cell(1, 0).text = "left"
    table.cell(1, 1).text = "003 −12.50"
    document.add_paragraph("after")
    if unsupported_element:
        table._tbl.addprevious(OxmlElement("w:customXml"))
    for element in list(document.element.body):
        node = (
            etree.Comment("Not document text")
            if node_kind == "comment"
            else etree.ProcessingInstruction("synthetic", "not document text")
        )
        element.addprevious(node)
    document.save(source)

    result = extract_file(source)

    assert result.status is (
        ExtractionStatus.NEEDS_REVIEW if unsupported_element else ExtractionStatus.SUCCESS
    )
    assert result.error_category is None
    assert [block.kind for block in result.blocks] == ["text", "table", "text"]
    assert [block.order for block in result.blocks] == [0, 1, 2]
    assert result.blocks[0].text == "before −001.50"
    assert result.blocks[2].text == "after"
    extracted = result.blocks[1]
    assert isinstance(extracted, TableBlock)
    assert extracted.table_index == 0
    cells = [cell for row in extracted.rows for cell in row.cells]
    assert [(c.text, c.row_index, c.column_index, c.row_span, c.column_span) for c in cells] == [
        ("merged header", 0, 0, 1, 2),
        ("left", 1, 0, 1, 1),
        ("003 −12.50", 1, 1, 1, 1),
    ]
    assert [(w.code, w.count) for w in result.warnings] == (
        [(WarningCode.UNSUPPORTED_BODY_CONTENT_NOT_EXTRACTED, 1)] if unsupported_element else []
    )

    # Preserve comments/PIs while resolving paths against the actual saved original.
    with zipfile.ZipFile(source) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    body = root.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}body")
    assert body is not None
    body_index = list(root).index(body)
    body_ordinals = (1, 5, 7) if unsupported_element else (1, 3, 5)
    for block, ordinal in zip(result.blocks, body_ordinals, strict=True):
        assert block.source.part_uri == "/word/document.xml"
        assert block.source.xml_path == (body_index, ordinal)
        assert block.source.page_number is None
        node = root
        for index in block.source.xml_path:
            node = node[index]
        assert node is body[ordinal]
        if isinstance(block, TextBlock):
            assert "".join(node.itertext()) == block.text
        else:
            assert node.tag == "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl"
    for cell in cells:
        assert cell.source.part_uri == "/word/document.xml"
        assert cell.source.xml_path[:2] == (body_index, body_ordinals[1])
        node = root
        for index in cell.source.xml_path:
            node = node[index]
        assert node.tag == "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc"
        assert "".join(node.itertext()) == cell.text

    for limits in (
        ExtractionLimits(max_document_blocks=2),
        ExtractionLimits(max_extracted_characters=1),
        ExtractionLimits(max_table_cells=2),
    ):
        bounded = extract_file(source, limits=limits)
        assert bounded.status is ExtractionStatus.FAILED
        assert bounded.error_category is ErrorCategory.LIMIT_EXCEEDED


def test_docx_preserves_header_with_explicit_review_status(tmp_path: Path) -> None:
    source = tmp_path / "header.docx"
    document = Document()
    document.add_paragraph("body")
    document.sections[0].header.paragraphs[0].text = "private header"
    document.save(source)

    result = extract_file(source)

    assert result.status is ExtractionStatus.NEEDS_REVIEW
    assert WarningCode.ANCILLARY_PROVENANCE_UNREVIEWED in {
        warning.code for warning in result.warnings
    }
    assert result.blocks[-1].text == "private header"
    assert result.blocks[-1].source.part_uri == "/word/header1.xml"
    assert "private header" not in json.dumps(
        [warning.code.value for warning in result.warnings]
    )


def test_xlsx_extracts_sparse_cells_formula_cache_and_table_metadata(
    tmp_path: Path,
) -> None:
    source = tmp_path / "sample.xlsx"
    _write_xlsx(source, formula=True, extreme=True)
    source = _set_formula_cache(source, "3")

    result = extract_file(source)

    assert result.status is ExtractionStatus.SUCCESS
    assert [sheet.name for sheet in result.sheets] == ["Metrics", "Second"]
    assert len(result.sheets[0].cells) == 7
    assert result.sheets[0].cells[-1].coordinate == "XFD1048576"
    formula_cell = next(cell for cell in result.sheets[0].cells if cell.coordinate == "C2")
    assert formula_cell.formula == "A2+B2"
    assert formula_cell.cached_value == "3"
    assert result.sheets[0].tables[0].display_name == "MetricsTable"
    assert result.sheets[0].tables[0].reference == "A1:C2"


def test_xlsx_missing_formula_cache_requires_review(tmp_path: Path) -> None:
    source = tmp_path / "formula.xlsx"
    _write_xlsx(source, formula=True)

    result = extract_file(source)

    assert result.status is ExtractionStatus.NEEDS_REVIEW
    formula_cell = next(cell for cell in result.sheets[0].cells if cell.coordinate == "C2")
    assert formula_cell.formula == "A2+B2"
    assert formula_cell.cached_value is None
    assert WarningCode.FORMULA_CACHED_VALUE_MISSING in {
        warning.code for warning in result.warnings
    }


def test_xlsx_stored_cell_limit_is_enforced_without_rectangular_iteration(
    tmp_path: Path,
) -> None:
    source = tmp_path / "sparse.xlsx"
    _write_xlsx(source, extreme=True)

    result = extract_file(source, limits=ExtractionLimits(max_xlsx_stored_cells=2))

    assert result.status is ExtractionStatus.FAILED
    assert result.error_category is ErrorCategory.LIMIT_EXCEEDED


@pytest.mark.parametrize("suffix", ["pdf", "docx", "xlsx"])
def test_corrupt_documents_return_sanitized_failures(tmp_path: Path, suffix: str) -> None:
    source = tmp_path / f"sensitive-name.{suffix}"
    source.write_bytes(b"not a valid document PRIVATE_DETAIL")

    result = extract_file(source)

    assert result.status is ExtractionStatus.FAILED
    assert result.error_category is ErrorCategory.CORRUPT_DOCUMENT
    serialized = json.dumps(result.to_dict(), sort_keys=True)
    assert "sensitive-name" not in serialized
    assert "PRIVATE_DETAIL" not in serialized


def test_extraction_is_deterministic(tmp_path: Path) -> None:
    source = tmp_path / "stable.docx"
    document = Document()
    document.add_paragraph("stable content")
    document.save(source)

    first = extract_file(source).to_dict()
    second = extract_file(source).to_dict()

    assert first == second


def test_corpus_runner_isolates_failures_and_persists_no_content(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    valid = corpus / "identifying-name.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "TOP_SECRET_PAYLOAD"
    workbook.save(valid)
    (corpus / "broken.pdf").write_bytes(b"broken SECRET_ERROR_DETAIL")

    first_output = tmp_path / "output-one"
    second_output = tmp_path / "output-two"
    first = extract_corpus(corpus, first_output, measure_latency=False)
    second = extract_corpus(corpus, second_output, measure_latency=False)

    assert first == second
    assert first["aggregate"]["document_count"] == 2
    assert first["aggregate"]["status_counts"] == {"failed": 1, "success": 1}
    persisted = "\n".join(
        path.read_text(encoding="utf-8") for path in first_output.iterdir()
    )
    assert "TOP_SECRET_PAYLOAD" not in persisted
    assert "SECRET_ERROR_DETAIL" not in persisted
    assert "identifying-name" not in persisted
    assert "broken.pdf" not in persisted


def test_unsupported_format_is_an_isolated_failure(tmp_path: Path) -> None:
    source = tmp_path / "sample.txt"
    source.write_text("content", encoding="utf-8")

    result = extract_file(source)

    assert result.status is ExtractionStatus.FAILED
    assert result.error_category is ErrorCategory.UNSUPPORTED_FORMAT


def test_corpus_file_limit_is_enforced_before_parsing(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    (corpus / "one.pdf").write_bytes(b"one")
    (corpus / "two.pdf").write_bytes(b"two")

    with pytest.raises(ExtractionConfigurationError, match="file limit"):
        extract_corpus(
            corpus,
            tmp_path / "output",
            limits=ExtractionLimits(max_files=1),
        )

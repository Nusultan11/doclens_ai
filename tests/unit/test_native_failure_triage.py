from __future__ import annotations

import json
import subprocess
import sys
import zipfile
from copy import deepcopy

import pytest

from doclens_ai.eda import corpus_v2 as corpus
from doclens_ai.extraction.fidelity_gold import fingerprint
from tools.extraction_quality import native_failure_triage as triage
from tools.extraction_quality.readiness_io import atomic


@pytest.fixture
def paired():
    registered, rows = [], []
    for index, (fmt, category) in enumerate(
        [("xlsx", "unsupported_format")] * 5
        + [("xlsx", "limit_exceeded")] * 2
        + [("docx", "extraction_failed")]
    ):
        metadata = {
            "id": f"nd-{index}",
            "sha256": str(index) * 64,
            "format": fmt,
            "size_bytes": 100 + index,
            "eligible": True,
            "split": "dev",
        }
        registered.append(metadata)
        rows.append(
            {
                **metadata,
                "status": "failed",
                "failure": category,
                "attempted": True,
                "capture": {
                    "status": "failed",
                    "error_category": category,
                    "content_sha256": "a" * 64,
                },
                "runtime": {"status": "completed", "seconds": 1.0},
            }
        )
    left = {"rows": rows}
    right = deepcopy(left)
    for row in right["rows"]:
        row["runtime"]["seconds"] = 2.0
    return left, right, registered


def test_all_eight_paired_failures_reconcile_independent_of_runtime(paired):
    assert len(triage.reconcile(*paired)) == 8


@pytest.mark.parametrize(
    "fault",
    [
        "missing",
        "duplicate",
        "category",
        "test",
        "excluded",
        "sha",
        "capture",
        "operational",
        "reproducibility",
    ],
)
def test_invalid_failed_populations_fail_closed(paired, fault):
    left, right, registered = paired
    if fault == "missing":
        right["rows"].pop()
    elif fault == "duplicate":
        right["rows"].append(deepcopy(right["rows"][0]))
    elif fault == "category":
        left["rows"][0]["failure"] = "limit_exceeded"
    elif fault == "test":
        registered[0]["split"] = "test"
    elif fault == "excluded":
        registered[0]["eligible"] = False
    elif fault == "sha":
        left["rows"][0]["sha256"] = "f" * 64
    elif fault == "capture":
        left["rows"][0]["capture"] = None
    elif fault == "operational":
        left["rows"][0]["runtime"]["status"] = "timeout"
    else:
        right["rows"][0]["capture"]["content_sha256"] = "b" * 64
    with pytest.raises(ValueError):
        triage.reconcile(left, right, registered)


@pytest.mark.parametrize("kind", ["chartsheet", "dialogsheet"])
def test_source_sheet_bindings_confirm_only_the_matching_branch(tmp_path, kind):
    source = tmp_path / "public.xlsx"
    office = triage.parsers._OFFICE_REL_NS
    ns = triage.parsers._SPREADSHEET_NS
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr(
            "xl/workbook.xml",
            f'<workbook xmlns="{ns}" xmlns:r="{office}">'
            '<sheets><sheet name="Synthetic" sheetId="1" r:id="r1"/>'
            "</sheets></workbook>",
        )
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            f'<Relationships xmlns="{triage.parsers._PACKAGE_REL_NS}">'
            f'<Relationship Id="r1" Type="{office}/{kind}" '
            'Target="sheets/public.xml"/></Relationships>',
        )
        archive.writestr("xl/sheets/public.xml", f'<{kind} xmlns="{ns}"/>')
    package = triage.package_evidence(source, "xlsx")
    row = {"failure": "unsupported_format"}
    diagnosis = triage.diagnosis(row, package, None)
    assert diagnosis["classification"] == "CONFIRMED"
    assert diagnosis["mechanism"].endswith(kind)
    package["sheets"][0]["root_tag"] = f"{{{ns}}}worksheet"
    assert triage.diagnosis(row, package, None)["classification"] == "UNRESOLVED"


def test_limit_trigger_requires_matching_capture_and_measured_counter():
    row = {"failure": "limit_exceeded", "capture": {"content_sha256": "same"}}
    event = {
        "function": "_extract_xlsx",
        "exception_type": "_ExtractionFailure",
        "numeric_locals": {"character_count": 5_000_001},
        "line": 859,
    }
    replay = {"capture_fingerprint": "same", "exception_events": [event]}
    result = triage.diagnosis(row, {}, replay)
    assert result["classification"] == "CONFIRMED"
    event["numeric_locals"]["character_count"] = 5_000_000
    assert triage.diagnosis(row, {}, replay)["classification"] == "UNRESOLVED"
    replay["capture_fingerprint"] = "changed"
    with pytest.raises(ValueError, match="does not reproduce"):
        triage.diagnosis(row, {}, replay)


def test_docx_exception_confirmation_does_not_invent_underlying_source_cause():
    row = {"failure": "extraction_failed", "capture": {"content_sha256": "same"}}
    replay = {
        "capture_fingerprint": "same",
        "exception_events": [
            {
                "function": "_docx_table_block",
                "exception_type": "IndexError",
                "line": 472,
                "numeric_locals": {},
                "traceback": [],
            }
        ],
    }
    result = triage.diagnosis(row, {}, replay)
    assert result["classification"] == "CONFIRMED"
    assert result["mechanism"] == "docx_IndexError__docx_table_block"
    assert (
        "deeper source-causal interpretation requires package evidence"
        in result["recommended_minimal_fix"]
    )
    replay["exception_events"][0]["function"] = "_extract_xlsx"
    assert triage.diagnosis(row, {}, replay)["classification"] == "UNRESOLVED"


def test_original_body_xml_comment_explains_current_docx_failure(tmp_path, monkeypatch):
    from docx import Document
    from lxml.etree import Comment

    monkeypatch.setattr(triage, "ROOT", tmp_path)
    source = tmp_path / "comment-before-paragraph.docx"
    document = Document()
    document.add_paragraph("Synthetic regression content")
    document.element.body.insert(0, Comment("Synthetic XML comment"))
    document.save(source)
    metadata = {
        "sha256": corpus.sha256(source),
        "size_bytes": source.stat().st_size,
        "format": "docx",
    }
    package = triage.package_evidence(source, "docx")
    replay = triage.trace_extraction(source, metadata)
    assert replay["capture"]["error_category"] == "extraction_failed"
    result = triage.diagnosis(
        {
            "failure": "extraction_failed",
            "capture": {"content_sha256": replay["capture_fingerprint"]},
        },
        package,
        replay,
    )
    assert result["classification"] == "CONFIRMED"
    assert result["mechanism"] == "docx_xml_comment_body_traversal"
    assert package["docx_structure"]["non_element_body_nodes"][0]["child_ordinal"] == 0


def test_diagnostic_artifact_verification_rejects_changed_bytes(tmp_path, monkeypatch):
    monkeypatch.setattr(triage, "ROOT", tmp_path)
    output = tmp_path / "evaluation-output" / "triage"
    atomic(output / "ledger.json", {"rows": []})
    marker = {"artifact_sha256": triage.baseline.artifact_hashes(output)}
    marker["fingerprint"] = fingerprint(marker)
    atomic(output / "COMPLETE.json", marker)
    triage.verify_artifacts(output)
    atomic(output / "ledger.json", {"rows": ["tampered"]})
    with pytest.raises(ValueError, match="evidence drift"):
        triage.verify_artifacts(output)


def test_worker_cli_traces_unchanged_native_result_and_checks_source(tmp_path, monkeypatch):
    # An original unsupported chartsheet exercises the native exception path without a fix.
    from openpyxl import Workbook

    monkeypatch.setattr(triage, "ROOT", tmp_path)
    source = tmp_path / "public.xlsx"
    workbook = Workbook()
    workbook.create_chartsheet("SyntheticChart")
    workbook.save(source)
    metadata = {
        "sha256": corpus.sha256(source),
        "size_bytes": source.stat().st_size,
        "format": "xlsx",
    }
    result = triage.trace_extraction(source, metadata)
    assert result["capture"]["error_category"] == "unsupported_format"
    assert result["capture_fingerprint"] == fingerprint(result["capture"])
    assert any(e["function"] == "_extract_xlsx" for e in result["exception_events"])
    source.write_bytes(source.read_bytes() + b"drift")
    with pytest.raises(ValueError, match="bytes mismatch"):
        triage.trace_extraction(source, metadata)


def test_actual_worker_module_invocation_and_private_output(tmp_path):
    # The worker branch must be invocable as a module, not '-m __main__'.
    output = (
        triage.ROOT / "evaluation-output" / "triage-unit" / tmp_path.parent.name / tmp_path.name
    )
    source = (
        triage.ROOT / "evaluation-output" / "triage-unit" / tmp_path.parent.name / "public.docx"
    )
    source.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("word/document.xml", "<document/>")
    metadata = {
        "sha256": corpus.sha256(source),
        "size_bytes": source.stat().st_size,
        "format": "docx",
    }
    atomic(output / "request.json", {"source": str(source), "metadata": metadata})
    completed = subprocess.run(
        [sys.executable, "-B", "-m", triage.MODULE, "worker", "--output", str(output)],
        capture_output=True,
        timeout=30,
        check=False,
    )
    assert completed.returncode == 0
    result = json.loads((output / "result.json").read_text(encoding="utf-8"))
    assert result["capture"]["status"] == "failed"
    assert result["capture"]["error_category"] == "corrupt_document"

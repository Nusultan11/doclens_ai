"""Synthetic contract tests only; never substitutes for the actual private audit."""

from __future__ import annotations

import csv
import json
import logging
from dataclasses import replace

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from doclens_ai.eda import data_quality_gate as gate
from doclens_ai.evaluation.gold import load_gold_jsonl
from doclens_ai.extraction.fidelity_gold import fingerprint, load_gold
from doclens_ai.extraction.models import ExtractionStatus

PRIVATE = "PRIVATE_CONTENT_SENTINEL"


def write_jsonl(path, records):
    path.write_text("".join(json.dumps(r) + "\n" for r in records), encoding="utf-8")


@pytest.fixture
def bundle(tmp_path):
    corpus = tmp_path / "private-corpus"
    corpus.mkdir()
    pdf = PdfWriter()
    pdf.add_blank_page(width=100, height=100)
    pdf.write(corpus / "private.pdf")
    pdf.close()
    doc = Document()
    doc.add_paragraph(PRIVATE)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    doc.save(corpus / "private.docx")
    workbook = Workbook()
    workbook.active.title = "PRIVATE_SHEET_NAME"
    workbook.active["A1"] = PRIVATE
    workbook.active["B1"] = 17
    workbook.save(corpus / "private.xlsx")
    workbook.close()
    manifest, fidelity, retrieval = (
        tmp_path / n for n in ("manifest.csv", "private-fidelity.jsonl", "private-retrieval.jsonl")
    )
    rows, records, queries = [], [], []
    for i, fmt in enumerate(gate.FORMATS, 1):
        path = corpus / ("private." + fmt)
        checksum, size = gate.digest(path), path.stat().st_size
        rows.append(
            {"document_id": f"c1-{i:04d}", "format": fmt, "sha256": checksum, "size_bytes": size}
        )
        locator = (
            {"page_number": 1}
            if fmt == "pdf"
            else {"block_index": 0}
            if fmt == "docx"
            else {"sheet_index": 0, "cell_coordinate": "A1"}
        )
        records.append(
            {
                "schema_version": "1.0",
                "dataset_revision": "corpus-v1-fidelity-v1",
                "dataset_kind": "synthetic",
                "review_basis": "synthetic_fixture",
                "document_id": f"c1-{i:04d}",
                "document_format": fmt,
                "source_sha256": checksum,
                "source_size_bytes": size,
                "review_status": "reviewed",
                "reviewer_count": 1,
                "reviewed_on": "2026-01-01",
                "anchors": [
                    {"anchor_id": "a1", "text": PRIVATE, "source": {"region": "body", **locator}}
                ],
                "order_pairs": [],
                "table_cells": [],
                "table_relationships": [],
                "cells": [],
            }
        )
        queries.append(
            {
                "schema_version": "1.0",
                "dataset_revision": "synthetic-retrieval-gold-v1",
                "synthetic_fixture": True,
                "question_id": f"fixture-q-{i:03d}",
                "document_family_id": f"fixture-family-{i:03d}",
                "language": "en",
                "question": PRIVATE + "?",
                "category": "fact",
                "answerable": True,
                "source_document_sha256": checksum,
                "split": "dev" if i < 3 else "test",
                "review_status": "reviewed",
                "reviewer_count": 1,
                "adjudication_state": "not_required",
                "evidence": [
                    {
                        "evidence_id": f"fixture-e-{i:03d}",
                        "document_sha256": checksum,
                        "document_format": fmt,
                        "locator": locator,
                    }
                ],
            }
        )
    with manifest.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    write_jsonl(fidelity, records)
    write_jsonl(retrieval, queries)
    sources = gate.fidelity_eval._read_sources(manifest)
    normalized = load_gold(fidelity, sources, allow_synthetic=True)
    contract = gate.FrozenContract(
        "synthetic-unit-fixture-NOT-real-corpus",
        1,
        gate.digest(manifest),
        fingerprint([r.model_dump(mode="json") for r in normalized]),
        load_gold_jsonl(retrieval).fingerprint,
        synthetic_fixture_only=True,
    )
    return corpus, manifest, fidelity, retrieval, contract


def run(bundle):
    return gate.audit(*bundle[:4], contract=bundle[4])


def edit_gold(bundle, which, edit):
    path = bundle[which]
    records = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]
    edit(records)
    write_jsonl(path, records)


def test_deterministic_private_and_read_only(bundle, tmp_path, monkeypatch):
    inputs = [p for p in bundle[0].iterdir()] + list(bundle[1:4])
    before = {p: gate.digest(p) for p in inputs}
    first, second = run(bundle), run(bundle)
    assert first == second and first["decision"] == "PASS"
    assert first["contract"]["synthetic_fixture_only"] is True
    assert first["corpus"]["inclusion_denominator"] == 3
    assert first["extraction"]["audited"] == 3
    assert first["extraction"]["needs_review"] == 1  # Blank PDF remains in denominator.
    assert first["extraction_gold"]["baseline_parser_evidence"]["missing_required_anchors"] == 1
    assert first["retrieval_gold"]["source_split_crossings"] == 0
    assert first["retrieval_gold"]["source_units"]["verified"] == 2
    root = tmp_path / "evaluation-output"
    monkeypatch.setattr(gate, "OUTPUT_ROOT", root)
    for name, report in (("run1", first), ("run2", second)):
        gate.write_reports(report, root / name)
    for filename in ("metrics.json", "aggregate.md"):
        a = (root / "run1" / filename).read_bytes()
        assert a == (root / "run2" / filename).read_bytes()
        assert all(
            token not in a.decode()
            for token in (
                PRIVATE,
                "PRIVATE_SHEET_NAME",
                str(tmp_path),
                "private.pdf",
                "private-fidelity",
            )
        )
    assert {p: gate.digest(p) for p in inputs} == before
    with pytest.raises(ValueError):
        gate.write_reports(first, root / "run1")
    with pytest.raises(ValueError):
        gate.write_reports(first, tmp_path / "outside")


@pytest.mark.parametrize("kind", ["missing", "changed", "same-size-change", "duplicate", "suffix"])
def test_corrupt_or_missing_sources_fail_closed(bundle, kind):
    path = bundle[0] / "private.pdf"
    if kind == "missing":
        path.unlink()
    elif kind == "changed":
        path.write_bytes(path.read_bytes() + b"changed")
    elif kind == "same-size-change":
        data = path.read_bytes()
        path.write_bytes(bytes([data[0] ^ 1]) + data[1:])
    elif kind == "duplicate":
        (bundle[0] / "copy.pdf").write_bytes(path.read_bytes())
    else:
        path.rename(path.with_suffix(".docx.bad"))
    result = run(bundle)
    assert result["decision"] == "FIX"
    assert result["findings"]["corpus_integrity_failed"] == 1
    assert result["extraction"]["not_evaluated"] == 3
    assert sum(sum(c.values()) for c in result["extraction"]["by_format"].values()) == 3
    assert result["extraction_gold"]["source_units"]["not_evaluated"] == 3
    if kind == "duplicate":
        assert result["corpus"]["inventory"]["duplicate_file_hashes"] == 1
    else:
        assert result["corpus"]["inventory"]["missing_or_changed_source_identities"] == 1


@pytest.mark.parametrize("which", [1, 2, 3])
def test_missing_required_inputs_are_blocked(bundle, which):
    bundle[which].unlink()
    assert run(bundle)["decision"] == "BLOCKED"


@pytest.mark.parametrize("field", ["document_id", "sha256"])
def test_duplicate_manifest_identity_or_hash(bundle, field):
    path = bundle[1]
    with path.open(encoding="utf-8", newline="") as stream:
        rows = list(csv.DictReader(stream))
    rows[1][field] = rows[0][field]
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    result = run(bundle)
    assert result["decision"] == "FIX" and result["findings"]["manifest_invalid"] == 1
    key = "duplicate_identities" if field == "document_id" else "duplicate_manifest_hashes"
    assert result["corpus"][key] == 1


@pytest.mark.parametrize("which", [2, 3])
@pytest.mark.parametrize("kind", ["malformed", "draft", "duplicate-json-key", "source-hash"])
def test_malformed_unapproved_gold_is_rejected_without_echo(bundle, which, kind, capsys):
    path = bundle[which]
    if kind == "malformed":
        path.write_text("{PRIVATE_CONTENT_SENTINEL", encoding="utf-8")
    elif kind == "duplicate-json-key":
        path.write_text(
            path.read_text(encoding="utf-8").replace(
                '"schema_version": "1.0"', '"schema_version": "bad", "schema_version": "1.0"'
            ),
            encoding="utf-8",
        )
    elif kind == "source-hash":
        field = "source_sha256" if which == 2 else "source_document_sha256"
        edit_gold(bundle, which, lambda rows: rows[0].update({field: "a" * 64}))
    else:
        edit_gold(bundle, which, lambda rows: rows[0].update(review_status="draft"))
    result = run(bundle)
    assert result["decision"] == "FIX"
    assert (
        result["findings"]["extraction_gold_invalid" if which == 2 else "retrieval_gold_invalid"]
        == 1
    )
    assert PRIVATE not in json.dumps(result)
    assert not capsys.readouterr().out


@pytest.mark.parametrize("which", [2, 3])
def test_out_of_bounds_locator_is_not_schema_only_pass(bundle, which):
    def edit(rows):
        loc = rows[0]["anchors"][0]["source"] if which == 2 else rows[0]["evidence"][0]["locator"]
        loc["page_number"] = 999

    edit_gold(bundle, which, edit)
    result = run(bundle)
    assert result["decision"] == "FIX"
    section = "extraction_gold" if which == 2 else "retrieval_gold"
    assert result[section]["source_units"]["invalid"] == 1


@pytest.mark.parametrize("kind", ["same-family", "same-source-new-family"])
def test_family_and_source_split_leakage_rejected(bundle, kind):
    def edit(rows):
        copy = json.loads(json.dumps(rows[1]))
        copy.update(question_id="another-question", split="test")
        if kind == "same-source-new-family":
            copy["document_family_id"] = "another-family"
        rows.append(copy)

    edit_gold(bundle, 3, edit)
    assert run(bundle)["findings"]["retrieval_gold_invalid"] == 1


def test_retrieval_offsets_use_existing_prefixed_units(bundle):
    def edit(rows):
        anchor = rows[2]["evidence"][0]
        result = gate.extract_file(bundle[0] / "private.xlsx")
        unit = next(u for u in gate._build_source_units(result) if u.source.cell_coordinate == "A1")
        assert len(unit.text) > len(PRIVATE)
        anchor.update(source_start=0, source_end=len(unit.text))

    edit_gold(bundle, 3, edit)
    result = run(bundle)
    # Snapshot changed, but the approved schema/source-unit boundary is correct.
    assert result["retrieval_gold"]["source_units"]["invalid"] == 0
    edit_gold(bundle, 3, lambda rows: rows[2]["evidence"][0].update(source_end=999999))
    assert run(bundle)["retrieval_gold"]["source_units"]["invalid"] == 1


def test_failure_isolation_and_denominator(bundle, monkeypatch):
    original = gate.extract_file

    def extract(path):
        result = original(path)
        return replace(result, status=ExtractionStatus.FAILED) if path.suffix == ".pdf" else result

    monkeypatch.setattr(gate, "extract_file", extract)
    result = run(bundle)
    assert result["extraction"]["audited"] == 3
    assert result["extraction"]["failures"] == 1
    assert result["extraction"]["eligible_by_status"] == 2
    assert result["corpus"]["inclusion_denominator"] == 3
    assert result["decision"] == "FIX"


def test_source_changes_during_audit_are_detected(bundle, monkeypatch):
    original = gate.extract_file

    def extract(path):
        result = original(path)
        if path.suffix == ".pdf":
            path.write_bytes(path.read_bytes() + b"change-after-extraction")
        return result

    monkeypatch.setattr(gate, "extract_file", extract)
    assert run(bundle)["findings"]["corpus_changed_during_audit"] == 1


def test_cli_and_diagnostics_are_content_free(bundle, tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(gate, "CORPUS_V1", bundle[4])
    root = tmp_path / "evaluation-output"
    monkeypatch.setattr(gate, "OUTPUT_ROOT", root)
    args = [
        "--corpus",
        str(bundle[0]),
        "--manifest",
        str(bundle[1]),
        "--extraction-gold",
        str(bundle[2]),
        "--retrieval-gold",
        str(bundle[3]),
        "--output-dir",
        str(root / "run"),
    ]
    previous = logging.root.manager.disable
    assert gate.main(args) == 0
    assert logging.root.manager.disable == previous
    assert json.loads(capsys.readouterr().out)["decision"] == "PASS"
    assert gate.main(args) == 2  # Never overwrite an existing output.
    assert PRIVATE not in capsys.readouterr().out
    with pytest.raises(SystemExit):
        gate.main(["--PRIVATE_CONTENT_SENTINEL"])
    assert PRIVATE not in capsys.readouterr().err


def test_cli_cannot_use_synthetic_gold_as_real_corpus(bundle):
    result = gate.audit(*bundle[:4])
    assert result["decision"] != "PASS"
    assert result["findings"]["manifest_snapshot_changed"] == 1
    assert result["findings"]["extraction_gold_invalid"] == 1
    assert result["findings"]["retrieval_gold_invalid"] == 1


def test_absent_corpus_is_blocked_and_never_an_empty_pass(bundle):
    result = gate.audit(bundle[0].parent / "missing", *bundle[1:4], contract=bundle[4])
    assert result["decision"] == "BLOCKED"
    assert result["blockers"]["corpus_unavailable"] == 1
    assert result["extraction"]["not_evaluated"] == 3


def test_supported_relationships_and_invalid_table_locator(bundle):
    def edit(rows):
        record = rows[1]
        record["anchors"] = [
            {
                "anchor_id": "a" + str(i + 1),
                "text": text,
                "source": {
                    "region": "body",
                    "block_index": 1,
                    "table_index": 0,
                    "row_index": 0,
                    "column_index": i,
                },
            }
            for i, text in enumerate(("left", "right"))
        ]
        record["table_relationships"] = [{"left": "a1", "right": "a2", "relationship": "same_row"}]

    edit_gold(bundle, 2, edit)
    report = run(bundle)
    assert report["extraction_gold"]["relationships"]["verified"] == 1
    edit_gold(bundle, 2, lambda rows: rows[1]["anchors"][1]["source"].update(column_index=999))
    report = run(bundle)
    assert report["extraction_gold"]["relationships"]["invalid"] == 1


def test_invalid_xlsx_coordinates_and_sheet_identity(bundle):
    edit_gold(
        bundle,
        3,
        lambda rows: rows[2]["evidence"][0]["locator"].update(cell_coordinate="XFE1048577"),
    )
    assert run(bundle)["retrieval_gold"]["source_units"]["invalid"] == 1
    edit_gold(
        bundle,
        3,
        lambda rows: rows[2]["evidence"][0]["locator"].update(
            cell_coordinate="A1", sheet_name="another-private-sheet"
        ),
    )
    assert run(bundle)["retrieval_gold"]["source_units"]["invalid"] == 1


def test_no_cli_synthetic_bypass_option(bundle, capsys):
    with pytest.raises(SystemExit):
        gate.main(["--allow-synthetic"])
    assert "synthetic" not in capsys.readouterr().err


def test_output_cannot_be_inside_inputs(bundle, monkeypatch):
    # Simulate an input corpus below the allowed output root: reject nesting too.
    monkeypatch.setattr(gate, "OUTPUT_ROOT", bundle[0].parent)
    with pytest.raises(ValueError):
        gate.write_reports(run(bundle), bundle[0] / "output", protected=(bundle[0],))


def test_output_symlink_escape_rejected(bundle, tmp_path, monkeypatch):
    root = tmp_path / "evaluation-output"
    root.mkdir()
    outside = tmp_path / "outside"
    outside.mkdir()
    try:
        (root / "link").symlink_to(outside, target_is_directory=True)
    except OSError:
        pytest.skip("Windows symlink privilege unavailable")
    monkeypatch.setattr(gate, "OUTPUT_ROOT", root)
    with pytest.raises(ValueError):
        gate.write_reports(run(bundle), root / "link/escape")


def test_unsupported_pdf_drawing_still_requires_real_page(bundle):
    edit_gold(
        bundle,
        2,
        lambda rows: rows[0]["anchors"][0]["source"].update(region="drawing", page_number=999),
    )
    assert run(bundle)["extraction_gold"]["source_units"]["invalid"] == 1


def test_invalid_snapshot_is_never_used_for_evidence_metrics(bundle):
    contract = replace(bundle[4], synthetic_fixture_only=False)
    result = gate.audit(*bundle[:4], contract=contract)
    for key in ("extraction_gold", "retrieval_gold"):
        assert result[key]["status"] == "INVALID"
        assert sum(result[key]["source_units"].values()) == 0
    assert "baseline_parser_evidence" not in result["extraction_gold"]

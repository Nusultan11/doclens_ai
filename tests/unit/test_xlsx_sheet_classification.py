"""Unsupported sheet kinds must not be confused with damaged source parts."""

from __future__ import annotations

import zipfile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook

from doclens_ai.extraction import ErrorCategory, ExtractionStatus, extract_file
from tests.unit.test_evidence_regressions import rewrite

S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"


@pytest.mark.parametrize("kind", ["dialogsheet", "chartsheet"])
@pytest.mark.parametrize("valid_root", [True, False])
def test_unsupported_sheet_parts_are_classified_without_partial_evidence(
    tmp_path, kind, valid_root
):
    path = tmp_path / "sheets.xlsx"
    wb = Workbook()
    wb.active["A1"] = "001.50 kg"
    wb.create_sheet("Unsupported")
    wb.save(path)
    with zipfile.ZipFile(path) as archive:
        relationships = archive.read("xl/_rels/workbook.xml.rels").decode()
    root = ET.fromstring(relationships)
    for node in root:
        if node.get("Target") == "/xl/worksheets/sheet2.xml":
            node.set("Type", R + kind)
    rewrite(
        path,
        {
            "xl/_rels/workbook.xml.rels": ET.tostring(root),
            "xl/worksheets/sheet2.xml": f'<{kind if valid_root else "worksheet"} xmlns="{S}"/>',
        },
    )
    result = extract_file(path)
    assert result.status is ExtractionStatus.FAILED
    assert result.error_category is (
        ErrorCategory.UNSUPPORTED_FORMAT if valid_root else ErrorCategory.CORRUPT_DOCUMENT
    )
    assert not result.sheets


@pytest.mark.parametrize(
    "relationship_type,root_tag",
    [(R + "worksheet", "chartsheet"), ("https://example.invalid/worksheet", "worksheet")],
)
def test_mismatched_worksheet_identity_is_corrupt(tmp_path, relationship_type, root_tag):
    path = tmp_path / "wrong-root.xlsx"
    Workbook().save(path)
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    for node in root:
        if node.get("Type") == R + "worksheet":
            node.set("Type", relationship_type)
    rewrite(
        path,
        {
            "xl/_rels/workbook.xml.rels": ET.tostring(root),
            "xl/worksheets/sheet1.xml": f'<{root_tag} xmlns="{S}"/>',
        },
    )
    result = extract_file(path)
    assert result.error_category is ErrorCategory.CORRUPT_DOCUMENT
    assert not result.sheets

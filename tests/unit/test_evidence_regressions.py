"""Public synthetic regressions for source fidelity and publication eligibility."""

from __future__ import annotations

import asyncio
import hashlib
import io
import json
import tomllib
import zipfile
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.worksheet.table import Table as XlsxTable

from doclens_ai.chunking import ChunkingConfig, ChunkingStatus, ChunkingStrategy, chunk_document
from doclens_ai.extraction.models import (
    DocumentFormat,
    ErrorCategory,
    ExtractionResult,
    ExtractionStatus,
    SourceLocation,
    TableBlock,
    TableCell,
    TableRow,
    TextBlock,
    WarningCode,
)
from doclens_ai.extraction.parsers import _ExtractionFailure, _parse_xml, extract_file
from doclens_ai.extraction.runtime import PARSER_VERSIONS, runtime_versions_match
from doclens_ai.extraction.validation import evidence_is_indexable
from doclens_ai.services.ingestion import (
    IngestionFailure,
    IngestionFailureCategory,
    IngestionService,
    _SourceIdentity,
    canonicalize_chunk,
)
from doclens_ai.services.processing import (
    PROCESSING_CONFIGURATION_VERSION,
    default_processing_configuration_fingerprint,
)
from doclens_ai.services.retrieval import source_locator
from tests.unit.test_docx_footnote_candidate import W, note, package
from tests.unit.test_extraction import _write_pdf


def rewrite(path: Path, updates: dict[str, str | bytes]) -> None:
    with zipfile.ZipFile(path) as archive:
        contents = {item.filename: archive.read(item) for item in archive.infolist()}
    contents.update(updates)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, value in contents.items():
            archive.writestr(name, value)


def chunks(result: ExtractionResult, maximum: int = 400):
    return chunk_document(
        result,
        document_id="synthetic",
        config=ChunkingConfig(
            strategy=ChunkingStrategy.STRUCTURE_AWARE,
            max_characters=maximum,
            overlap_characters=0,
            allow_needs_review=True,
        ),
    )


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16", "utf-16-le", "utf-16-be"])
def test_xml_declarations_rejected_before_expansion_in_any_encoding(encoding):
    xml = '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY e "WRONG_FACT">]><x>&e;</x>'
    with pytest.raises(_ExtractionFailure) as caught:
        _parse_xml(xml.encode(encoding))
    assert caught.value.category is ErrorCategory.CORRUPT_DOCUMENT
    safe = "<x>Қазақ Русский −001.50 %</x>"
    assert _parse_xml(safe.encode(encoding)).text == "Қазақ Русский −001.50 %"


@pytest.mark.parametrize("part", ["word/document.xml", "word/styles.xml", "word/footnotes.xml"])
def test_docx_validates_all_package_xml_before_library_parsing(tmp_path, part):
    path = tmp_path / "input.docx"
    doc = Document()
    doc.add_paragraph("safe")
    doc.save(path)
    rewrite(path, {part: '<!DOCTYPE x [<!ENTITY e "bad">]><x>&e;</x>'.encode("utf-16")})
    result = extract_file(path)
    assert result.status is ExtractionStatus.FAILED
    assert result.error_category is ErrorCategory.CORRUPT_DOCUMENT
    assert not result.blocks


def test_docx_omitted_grid_columns_merges_and_exact_ooxml_addresses(tmp_path):
    doc = Document()
    table = doc.add_table(rows=3, cols=4)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.text = f"column-{i}"
    table.cell(1, 1).merge(table.cell(2, 1)).text = "−12 kg"
    table.cell(1, 2).merge(table.cell(1, 3)).text = "2026 amount"
    first = table.rows[0]._tr
    first.remove(first.tc_lst[0])
    first.remove(first.tc_lst[-1])
    for tag in ("gridBefore", "gridAfter"):
        element = OxmlElement(f"w:{tag}")
        element.set(qn("w:val"), "1")
        first.get_or_add_trPr().append(element)
    path = tmp_path / "grid.docx"
    doc.save(path)
    result = extract_file(path)
    assert result.status is ExtractionStatus.SUCCESS
    assert evidence_is_indexable(result)
    extracted = result.blocks[0]
    assert [(c.text, c.column_index) for c in extracted.rows[0].cells] == [
        ("column-1", 1),
        ("column-2", 2),
    ]
    cells = [c for row in extracted.rows for c in row.cells]
    vertical = next(c for c in cells if c.text == "−12 kg")
    assert (vertical.row_index, vertical.column_index, vertical.row_span) == (1, 1, 2)
    horizontal = next(c for c in cells if c.text == "2026 amount")
    assert horizontal.column_span == 2
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
        for cell in cells:
            node = root
            for index in cell.source.xml_path:
                node = node[index]
            assert node.tag == f"{{{W}}}tc"
            assert cell.text == "".join(node.itertext())


def test_docx_invalid_grid_cannot_publish_partial_cells(tmp_path):
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    before = OxmlElement("w:gridBefore")
    before.set(qn("w:val"), "1")
    table.rows[0]._tr.get_or_add_trPr().append(before)
    path = tmp_path / "invalid.docx"
    doc.save(path)
    result = extract_file(path)
    assert result.status is ExtractionStatus.FAILED and not result.blocks


def test_docx_orphan_vertical_continuation_does_not_discard_facts(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "10 kg"
    table.cell(1, 0).text = "999 kg"
    continuation = OxmlElement("w:vMerge")
    table.cell(1, 0)._tc.get_or_add_tcPr().append(continuation)
    path = tmp_path / "orphan-merge.docx"
    doc.save(path)
    result = extract_file(path)
    assert result.status is ExtractionStatus.FAILED
    assert result.error_category is ErrorCategory.CORRUPT_DOCUMENT
    assert not result.blocks


@pytest.mark.parametrize("part", ["body", "header", "footer"])
def test_docx_grid_span_budget_checked_before_library_expansion(tmp_path, monkeypatch, part):
    from docx.shared import Inches
    from docx.table import _Row

    from doclens_ai.extraction.parsers import ExtractionLimits

    doc = Document()
    table = (
        doc.add_table(rows=1, cols=1)
        if part == "body"
        else getattr(doc.sections[0], part).add_table(rows=1, cols=1, width=Inches(2))
    )
    table.cell(0, 0).text = "bounded evidence"
    span = OxmlElement("w:gridSpan")
    span.set(qn("w:val"), "1000000")
    table.cell(0, 0)._tc.get_or_add_tcPr().append(span)
    path = tmp_path / "large-span.docx"
    doc.save(path)
    monkeypatch.setattr(
        _Row, "cells", property(lambda _: pytest.fail("unbounded cell expansion invoked"))
    )
    result = extract_file(path, limits=ExtractionLimits(max_table_cells=10))
    assert result.status is ExtractionStatus.FAILED
    assert result.error_category is ErrorCategory.LIMIT_EXCEEDED
    assert not result.blocks


def test_docx_unicode_controls_and_unsupported_inline_content(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph("Қазақ −001.50")
    run = paragraph.add_run()._r
    for tag in ("noBreakHyphen", "softHyphen", "tab"):
        run.append(OxmlElement(f"w:{tag}"))
    path = tmp_path / "unicode.docx"
    doc.save(path)
    result = extract_file(path)
    assert result.blocks[0].text == "Қазақ −001.50\u2011\u00ad\t"
    assert evidence_is_indexable(result)
    unsupported = OxmlElement("w:sdt")
    paragraph._p.append(unsupported)
    doc.save(path)
    assert not evidence_is_indexable(extract_file(path))


@pytest.mark.parametrize("moved", [False, True])
def test_footnotes_references_deduplication_and_real_part_paths(tmp_path, moved):
    path = package(
        tmp_path, notes=note("7", text="Қазақ −007.50 kg"), references=("7", "7"), moved=moved
    )
    main_part = "/parts/body.xml" if moved else "/word/document.xml"
    note_part = "/notes/actual.xml" if moved else "/word/footnotes.xml"
    result = extract_file(path)
    assert result.status is ExtractionStatus.NEEDS_REVIEW
    assert [b.text for b in result.blocks] == ["Synthetic body 42", "Қазақ −007.50 kg"]
    assert len(result.note_references) == 2
    assert len({ref.source.xml_path for ref in result.note_references}) == 2
    assert all(ref.target_part_uri == note_part for ref in result.note_references)
    assert result.blocks[-1].source.part_uri == note_part
    assert result.blocks[-1].source.xml_path == (0, 0)
    assert result.blocks[-1].source.note_id == "7"
    assert not evidence_is_indexable(result)
    outcome = chunks(result)
    assert outcome.status is ChunkingStatus.CHUNKED
    canonical = canonicalize_chunk(uuid4(), outcome.chunks[-1])
    locator = source_locator(
        SimpleNamespace(
            document_format="docx",
            source_provenance=canonical.source_provenance,
        )
    )
    assert any(span.source.part_uri == note_part for span in locator.spans)
    assert any(span.source.part_uri == main_part for span in locator.spans)
    assert "footnote_reference=7" in canonical.chunk_text
    assert locator.spans[0].offset_basis == "rendered_source_unit"


def test_dangling_or_unsupported_footnote_cannot_be_trusted(tmp_path):
    result = extract_file(package(tmp_path, references=("99",)))
    assert result.status is ExtractionStatus.NEEDS_REVIEW
    assert not evidence_is_indexable(result)


def test_shared_docx_header_footer_and_part_local_tables_keep_source_identity(tmp_path):
    from docx.enum.section import WD_SECTION
    from docx.shared import Inches

    doc = Document()
    doc.add_paragraph("body first")
    header = doc.sections[0].header
    header.paragraphs[0].text = "Shared header −10 kg"
    header.add_table(rows=1, cols=1, width=Inches(2)).cell(0, 0).text = "Header value 2026"
    footer = doc.sections[0].footer
    footer.paragraphs[0].text = "Footer 001"
    footer.add_table(rows=1, cols=1, width=Inches(2)).cell(0, 0).text = "Footer value 42"
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("body last")
    path = tmp_path / "parts.docx"
    doc.save(path)
    result = extract_file(path)
    assert result.status is ExtractionStatus.NEEDS_REVIEW
    assert result.blocks[0].text == "body first"
    texts = [b.text for b in result.blocks if isinstance(b, TextBlock)]
    assert texts.count("Shared header −10 kg") == 1
    assert texts.count("Footer 001") == 1
    tables = [b for b in result.blocks if isinstance(b, TableBlock)]
    assert [b.table_index for b in tables] == [0, 0]
    assert {b.source.part_uri for b in tables} == {"/word/header1.xml", "/word/footer1.xml"}
    with zipfile.ZipFile(path) as archive:
        for block in tables:
            cell = block.rows[0].cells[0]
            root = ET.fromstring(archive.read(cell.source.part_uri.lstrip("/")))
            for index in cell.source.xml_path:
                root = root[index]
            assert root.tag == f"{{{W}}}tc" and cell.text == "".join(root.itertext())
    outcome = chunks(result)
    assert outcome.status is ChunkingStatus.CHUNKED
    assert outcome.covered_source_characters == outcome.normalized_source_characters
    assert not evidence_is_indexable(result)


@pytest.mark.parametrize(
    "kind,value",
    [
        ("s", "-1"),
        ("s", "1"),
        ("s", "0.0"),
        ("s", "bad"),
        ("b", "invalid"),
        ("b", "2"),
        ("n", "NaN"),
        ("unknown", "1"),
    ],
)
def test_invalid_xlsx_values_fail_without_plausible_substitutes(tmp_path, kind, value):
    path = tmp_path / "invalid.xlsx"
    wb = Workbook()
    wb.active["A1"] = "original"
    wb.save(path)
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rewrite(
        path,
        {
            "xl/worksheets/sheet1.xml": (
                f'<worksheet xmlns="{ns}"><sheetData><row r="1">'
                f'<c r="A1" t="{kind}"><v>{value}</v></c></row></sheetData></worksheet>'
            ),
            "xl/sharedStrings.xml": f'<sst xmlns="{ns}"><si><t>LAST</t></si></sst>',
        },
    )
    result = extract_file(path)
    assert result.status is ExtractionStatus.FAILED
    assert result.error_category is ErrorCategory.CORRUPT_DOCUMENT
    assert not result.sheets


@pytest.mark.parametrize(
    "raw,expected", [("0", "false"), ("1", "true"), ("true", "true"), ("false", "false")]
)
def test_xlsx_boolean_lexical_forms_preserve_raw(tmp_path, raw, expected):
    path = tmp_path / "bool.xlsx"
    wb = Workbook()
    wb.active["A1"] = True
    wb.save(path)
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode()
    rewrite(path, {"xl/worksheets/sheet1.xml": xml.replace("<v>1</v>", f"<v>{raw}</v>")})
    cell = extract_file(path).sheets[0].cells[0]
    assert (cell.value, cell.raw_value, cell.data_type) == (expected, raw, "b")


@pytest.mark.parametrize("date_system", ["1900", "1904"])
def test_xlsx_formats_dates_raw_precision_formula_and_missing_shared_cache(tmp_path, date_system):
    from openpyxl.utils.datetime import CALENDAR_MAC_1904

    path = tmp_path / "formats.xlsx"
    wb = Workbook()
    if date_system == "1904":
        wb.epoch = CALENDAR_MAC_1904
    sheet = wb.active
    sheet["A1"] = 0.25
    sheet["A1"].number_format = "0.00%"
    sheet["B1"] = 45000
    sheet["B1"].number_format = "yyyy-mm-dd"
    sheet["C1"] = "=A1*2"
    sheet["D1"] = "=A1*2"
    sheet.merge_cells("A2:B2")
    sheet["A2"] = "unit kg"
    wb.save(path)
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode()
    xml = xml.replace("<f>A1*2</f><v></v>", "<f>A1*2</f><v>0.50</v>", 1)
    xml = xml.replace("<f>A1*2</f><v></v>", '<f t="shared" si="0"/><v/>')
    rewrite(path, {"xl/worksheets/sheet1.xml": xml})
    result = extract_file(path)
    assert result.sheets[0].date_system == date_system
    cells = {c.coordinate: c for c in result.sheets[0].cells}
    assert (cells["A1"].value, cells["A1"].number_format) == ("0.25", "0.00%")
    assert (cells["B1"].value, cells["B1"].number_format) == ("45000", "yyyy-mm-dd")
    assert cells["C1"].formula == "A1*2" and cells["C1"].cached_value == "0.50"
    assert cells["D1"].has_formula and cells["D1"].formula is None
    assert cells["D1"].cached_value is None
    assert dict(cells["D1"].formula_attributes) == {"si": "0", "t": "shared"}
    assert "B2" not in cells and result.sheets[0].merged_ranges == ("A2:B2",)
    assert result.status is ExtractionStatus.NEEDS_REVIEW and not evidence_is_indexable(result)
    rendered = "\n".join(c.text for c in chunks(result).chunks)
    assert "number_format=0.00%" in rendered and "value=0.25" in rendered
    assert "25.00%" not in rendered
    assert "formula=<unavailable>\tcache=<missing>" in rendered


def test_note_only_paragraph_keeps_its_reference_through_chunking(tmp_path):
    path = package(tmp_path)
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode()
    rewrite(path, {"word/document.xml": xml.replace("Synthetic body 42", "")})
    result = extract_file(path)
    assert result.note_references and result.blocks[0].text == ""
    outcome = chunks(result)
    assert outcome.status is ChunkingStatus.CHUNKED
    assert any("Synthetic note 007.50" in chunk.text for chunk in outcome.chunks)
    assert any("footnote_reference=7" in chunk.text for chunk in outcome.chunks)


@pytest.mark.parametrize("fmt", [DocumentFormat.PDF, DocumentFormat.DOCX])
def test_generic_table_identity_never_collides_across_pages_or_parts(fmt):
    blocks = []
    for index in range(2):
        source = SourceLocation(
            page_number=index + 1 if fmt is DocumentFormat.PDF else None,
            part_uri=f"/part-{index}.xml" if fmt is DocumentFormat.DOCX else None,
            table_index=0,
            block_index=0,
        )
        cell = TableCell(0, 0, f"value-{index}", replace(source, row_index=0, column_index=0))
        blocks.append(TableBlock(index, 0, (TableRow(0, (cell,)),), source))
    result = ExtractionResult("1.1", fmt, ExtractionStatus.SUCCESS, "a" * 64, 42, tuple(blocks))
    outcome = chunks(result)
    assert outcome.status is ChunkingStatus.CHUNKED and len(outcome.chunks) == 2
    assert outcome.covered_source_characters == outcome.normalized_source_characters
    ids = [c.provenance[0].unit_id for c in outcome.chunks]
    assert len(set(ids)) == 2
    assert all(fmt.value in identity for identity in ids)


def test_pdf_missing_page_locator_is_not_invented():
    result = ExtractionResult(
        "1.1",
        DocumentFormat.PDF,
        ExtractionStatus.SUCCESS,
        "a" * 64,
        42,
        (TextBlock(12, "unlocated", SourceLocation()),),
    )
    assert chunks(result).status is ChunkingStatus.FAILED


def test_declared_xlsx_headers_stay_with_values_across_chunks(tmp_path):
    path = tmp_path / "headers.xlsx"
    wb = Workbook()
    sheet = wb.active
    sheet.append(["Period 2026", "Amount kg"])
    for index in range(1, 12):
        sheet.append([f"entity-{index}", -index])
    sheet.add_table(XlsxTable(displayName="Facts", ref="A1:B12"))
    wb.save(path)
    result = extract_file(path)
    assert evidence_is_indexable(result)
    outcome = chunks(result, 450)
    assert len(outcome.chunks) > 1
    assert not any(c.split_oversized_element for c in outcome.chunks)
    for chunk in outcome.chunks:
        if "entity-" in chunk.text:
            assert "Period 2026" in chunk.text and "Amount kg" in chunk.text
            coordinates = {span.source.cell_coordinate for span in chunk.provenance}
            assert {"A1", "B1"} <= coordinates
            assert "table=Facts" in chunk.text
    assert outcome.covered_source_characters == outcome.normalized_source_characters


def test_declared_docx_headers_repeat_with_real_provenance(tmp_path):
    doc = Document()
    table = doc.add_table(rows=9, cols=2)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    table.cell(0, 0).text = "Period 2026"
    table.cell(0, 1).text = "Amount kg"
    for index in range(1, 9):
        table.cell(index, 0).text = f"entity-{index}"
        table.cell(index, 1).text = f"−{index}"
    path = tmp_path / "headers.docx"
    doc.save(path)
    result = extract_file(path)
    assert result.blocks[0].header_row_indices == (0,)
    outcome = chunks(result, 130)
    assert len(outcome.chunks) > 1
    assert not any(c.split_oversized_element for c in outcome.chunks)
    for chunk in outcome.chunks:
        if "entity-" in chunk.text:
            assert "Period 2026" in chunk.text and "Amount kg" in chunk.text
            assert {span.source.row_index for span in chunk.provenance} >= {0}


@pytest.mark.parametrize("text", ["corrupt \ufffd 10 kg", "bad \ud800"])
def test_corrupted_unicode_and_locator_mismatches_cannot_index(text):
    location = SourceLocation(block_index=0, part_uri="/word/document.xml", xml_path=(0,))
    result = ExtractionResult(
        "1.1",
        DocumentFormat.DOCX,
        ExtractionStatus.SUCCESS,
        "a" * 64,
        42,
        (TextBlock(0, text, location),),
    )
    assert not evidence_is_indexable(result)
    cell = TableCell(0, 0, "10 kg", replace(location, table_index=99, row_index=0, column_index=0))
    result = replace(
        result,
        blocks=(TableBlock(0, 0, (TableRow(0, (cell,)),), replace(location, table_index=0)),),
    )
    assert not evidence_is_indexable(result)


def test_pinned_parser_stack_matches_dependency_contract():
    config = tomllib.loads(Path("pyproject.toml").read_text())
    assert set(config["project"]["optional-dependencies"]["eda"]) == {
        f"{name}=={version}" for name, version in PARSER_VERSIONS.items()
    }
    assert runtime_versions_match()


def test_ingestion_rejects_pdf_success_and_partial_table_evidence(tmp_path):
    from tests.integration.test_canonical_chunks import MemoryStorage

    path = tmp_path / "table.pdf"
    _write_pdf(path, ["Period 2026 Amount kg entity A -10, source text preserved."])
    result = extract_file(path)
    assert result.status is ExtractionStatus.SUCCESS
    assert WarningCode.PDF_TABLE_STRUCTURE_NOT_PRESERVED in {w.code for w in result.warnings}
    assert not evidence_is_indexable(result)
    payload = path.read_bytes()
    storage = MemoryStorage()
    content_type = "application/pdf"
    checksum = hashlib.sha256(payload).hexdigest()
    source = _SourceIdentity(
        uuid4(),
        uuid4(),
        uuid4(),
        PROCESSING_CONFIGURATION_VERSION,
        default_processing_configuration_fingerprint(),
        "test-key",
        content_type,
        len(payload),
        checksum,
    )

    async def exercise():
        await storage.put_stream(
            "test-key", io.BytesIO(payload), content_type=content_type, sha256_checksum=checksum
        )
        service = IngestionService(object(), storage)
        with pytest.raises(IngestionFailure) as failure:
            await service._extract_and_chunk(source, suffix=".pdf")
        assert failure.value.category is IngestionFailureCategory.EXTRACTION_NEEDS_REVIEW
        assert not failure.value.retryable

    asyncio.run(exercise())
    # Even a structurally valid extraction cannot silently publish a split table row.
    block = TableBlock(
        0,
        0,
        (
            TableRow(
                0,
                (
                    TableCell(
                        0,
                        0,
                        "long evidence " * 200,
                        SourceLocation(block_index=0, table_index=0, row_index=0, column_index=0),
                    ),
                ),
            ),
        ),
        SourceLocation(block_index=0, table_index=0),
    )
    extracted = ExtractionResult(
        "1.1",
        DocumentFormat.DOCX,
        ExtractionStatus.SUCCESS,
        checksum,
        len(payload),
        (block,),
    )
    outcome = chunk_document(
        extracted,
        document_id=str(source.document_version_id),
        config=ChunkingConfig(strategy=ChunkingStrategy.STRUCTURE_AWARE),
    )
    assert outcome.status is ChunkingStatus.CHUNKED
    with pytest.raises(IngestionFailure) as failure:
        IngestionService(object(), storage)._validate_chunks(
            replace(
                source,
                content_type="application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document",
            ),
            outcome.chunks,
        )
    assert failure.value.category is IngestionFailureCategory.CHUNKING_FAILED


def test_no_private_text_in_warning_serialization(tmp_path):
    result = extract_file(package(tmp_path, notes=note(text="PRIVATE_SECRET")))
    assert "PRIVATE_SECRET" not in json.dumps([w.code for w in result.warnings])


@pytest.mark.parametrize("suffix", [".docx", ".xlsx"])
def test_supported_ingestion_and_canonical_provenance_round_trip(suffix):
    from tests.integration.test_canonical_chunks import MemoryStorage, _docx_bytes, _xlsx_bytes

    payload = _docx_bytes() if suffix == ".docx" else _xlsx_bytes()
    content_type = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        if suffix == ".docx"
        else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    source = _SourceIdentity(
        uuid4(),
        uuid4(),
        uuid4(),
        PROCESSING_CONFIGURATION_VERSION,
        default_processing_configuration_fingerprint(),
        "test-key",
        content_type,
        len(payload),
        hashlib.sha256(payload).hexdigest(),
    )

    async def exercise():
        storage = MemoryStorage()
        await storage.put_stream(
            "test-key",
            io.BytesIO(payload),
            content_type=content_type,
            sha256_checksum=source.sha256_checksum,
        )
        service = IngestionService(object(), storage)
        service._validate_processing_configuration(source)
        first = await service._extract_and_chunk(source, suffix=suffix)
        assert first == await service._extract_and_chunk(source, suffix=suffix)
        canonical = tuple(canonicalize_chunk(source.processing_run_id, chunk) for chunk in first)
        for chunk in canonical:
            assert json.loads(json.dumps(chunk.source_provenance)) == chunk.source_provenance
            locator = source_locator(
                SimpleNamespace(
                    document_format=chunk.document_format,
                    source_provenance=chunk.source_provenance,
                )
            )
            assert all(span.offset_basis == "rendered_source_unit" for span in locator.spans)

    asyncio.run(exercise())


def test_footnote_source_hash_mismatch_discards_all_output(tmp_path, monkeypatch):
    from doclens_ai.extraction import docx_footnote_candidate

    path = package(tmp_path)
    result = docx_footnote_candidate.extract_footnotes(path)
    monkeypatch.setattr(
        docx_footnote_candidate,
        "extract_footnotes",
        lambda _: replace(result, source_sha256="0" * 64),
    )
    extracted = extract_file(path)
    assert extracted.status is ExtractionStatus.FAILED
    assert not extracted.blocks and not extracted.note_references


def test_unknown_number_format_preserves_raw_number_but_requires_review(tmp_path):
    path = tmp_path / "unknown-format.xlsx"
    wb = Workbook()
    wb.active["A1"] = -1.25
    wb.save(path)
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/styles.xml").decode()
    rewrite(path, {"xl/styles.xml": xml.replace('numFmtId="0"', 'numFmtId="23"')})
    result = extract_file(path)
    cell = result.sheets[0].cells[0]
    assert (cell.raw_value, cell.value) == ("-1.25", "-1.25")
    assert cell.number_format is None and cell.number_format_id == 23
    assert result.status is ExtractionStatus.NEEDS_REVIEW
    assert not evidence_is_indexable(result)


def test_ambiguous_xlsx_headers_and_merged_values_are_held_for_review(tmp_path):
    path = tmp_path / "stale-header.xlsx"
    wb = Workbook()
    sheet = wb.active
    sheet.append(["Header", "Value kg"])
    sheet.append(["entity", 42])
    sheet.add_table(XlsxTable(displayName="Facts", ref="A1:B2"))
    wb.save(path)
    with zipfile.ZipFile(path) as archive:
        table = archive.read("xl/tables/table1.xml").decode()
    rewrite(path, {"xl/tables/table1.xml": table.replace('name="Header"', 'name="Other"')})
    result = extract_file(path)
    assert result.status is ExtractionStatus.NEEDS_REVIEW
    assert not evidence_is_indexable(result)
    assert result.sheets[0].cells[-1].raw_value == "42"
    wb.save(path)
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode()
    rewrite(
        path,
        {
            "xl/worksheets/sheet1.xml": xml.replace(
                "</sheetData>", '</sheetData><mergeCells><mergeCell ref="A2:B2"/></mergeCells>'
            )
        },
    )
    result = extract_file(path)
    assert result.status is ExtractionStatus.NEEDS_REVIEW
    assert not evidence_is_indexable(result)
    assert result.sheets[0].cells[-1].raw_value == "42"


@pytest.mark.parametrize("index", [0, 1])
def test_valid_shared_string_boundaries_and_phonetic_text(index):
    from doclens_ai.extraction.parsers import _xlsx_cell_value, _xlsx_string

    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    cell = ET.fromstring(f'<c xmlns="{ns}" t="s"><v>{index}</v></c>')
    assert _xlsx_cell_value(cell, cell[0], ("FIRST", "LAST")) == ("FIRST", "LAST")[index]
    rich = ET.fromstring(
        f'<si xmlns="{ns}"><r><t>Қазақ −001.50</t></r>'
        '<rPh sb="0" eb="1"><t>phonetic annotation</t></rPh><r><t> kg</t></r></si>'
    )
    assert _xlsx_string(rich) == "Қазақ −001.50 kg"


def test_office_limits_apply_before_xml_is_read(tmp_path, monkeypatch):
    from doclens_ai.extraction import parsers

    path = tmp_path / "bounded.docx"
    doc = Document()
    doc.add_paragraph("safe synthetic source")
    doc.save(path)

    def forbidden(_):
        raise AssertionError("No XML read is permitted before archive entry limits")

    monkeypatch.setattr(parsers, "_parse_xml", forbidden)
    result = extract_file(path, limits=parsers.ExtractionLimits(max_office_entry_bytes=1))
    assert result.status is ExtractionStatus.FAILED
    assert result.error_category is ErrorCategory.LIMIT_EXCEEDED


def test_duplicate_external_internal_relationship_identity_is_rejected(tmp_path):
    path = tmp_path / "relationships.xlsx"
    wb = Workbook()
    wb.active["A1"] = 42
    wb.save(path)
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/_rels/workbook.xml.rels").decode()
    xml = xml.replace(
        "</Relationships>",
        '<Relationship Id="rId1" TargetMode="External" Target="https://example.invalid" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"/>'
        "</Relationships>",
    )
    rewrite(path, {"xl/_rels/workbook.xml.rels": xml})
    result = extract_file(path)
    assert result.status is ExtractionStatus.FAILED
    assert result.error_category is ErrorCategory.CORRUPT_DOCUMENT
    assert not result.sheets


@pytest.mark.parametrize("tag", ["endnoteReference", "commentReference"])
def test_unsupported_ancillary_reference_without_default_part_is_not_trusted(tmp_path, tag):
    doc = Document()
    paragraph = doc.add_paragraph("Body text must remain available.")
    reference = OxmlElement(f"w:{tag}")
    reference.set(qn("w:id"), "7")
    paragraph.add_run()._r.append(reference)
    path = tmp_path / "unsupported-note.docx"
    doc.save(path)
    result = extract_file(path)
    assert result.blocks[0].text == "Body text must remain available."
    assert result.status is ExtractionStatus.NEEDS_REVIEW
    assert not evidence_is_indexable(result)


def test_number_formats_count_towards_extracted_character_limit(tmp_path):
    from doclens_ai.extraction.parsers import ExtractionLimits

    path = tmp_path / "format-limit.xlsx"
    wb = Workbook()
    wb.active["A1"] = 1
    wb.active["A1"].number_format = "0" * 100
    wb.save(path)
    result = extract_file(path, limits=ExtractionLimits(max_extracted_characters=30))
    assert result.status is ExtractionStatus.FAILED
    assert result.error_category is ErrorCategory.LIMIT_EXCEEDED

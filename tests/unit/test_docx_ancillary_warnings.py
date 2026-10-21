"""Unsupported header/footer content must never silently become indexable."""

import pytest
from docx import Document
from docx.oxml import OxmlElement

from doclens_ai.extraction import ExtractionStatus, WarningCode, extract_file
from doclens_ai.extraction.validation import evidence_is_indexable


@pytest.mark.parametrize("part", ["header", "footer"])
@pytest.mark.parametrize(
    "tag,warning",
    [
        ("drawing", WarningCode.DRAWINGS_NOT_EXTRACTED),
        ("fldSimple", WarningCode.UNSUPPORTED_BODY_CONTENT_NOT_EXTRACTED),
        ("endnoteReference", WarningCode.ENDNOTES_NOT_EXTRACTED),
    ],
)
def test_unsupported_content_in_textless_ancillary_parts_stays_untrusted(
    tmp_path, part, tag, warning
):
    doc = Document()
    doc.add_paragraph("Body facts 001.50 kg remain unchanged.")
    container = getattr(doc.sections[0], part)
    container.paragraphs[0].add_run()._r.append(OxmlElement(f"w:{tag}"))
    path = tmp_path / "ancillary.docx"
    doc.save(path)

    result = extract_file(path)

    assert result.status is ExtractionStatus.NEEDS_REVIEW
    assert warning in {w.code for w in result.warnings}
    assert not evidence_is_indexable(result)
    assert result.blocks[0].text == "Body facts 001.50 kg remain unchanged."

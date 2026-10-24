"""Lexical values, semantic versus ordering links and genuine offset provenance."""

import hashlib
import json
from dataclasses import asdict, replace
from uuid import uuid4

import pytest

from doclens_ai.chunking.core import chunk_document
from doclens_ai.chunking.ir import _render, chunk_verified_ir
from doclens_ai.chunking.models import ChunkingConfig, ChunkingStrategy
from doclens_ai.extraction.engines import default_registry, default_route
from doclens_ai.extraction.ir import CellElement, DocumentIR, FormulaElement, Relation
from doclens_ai.extraction.models import (
    DocumentFormat,
    ExtractionResult,
    ExtractionStatus,
    SourceLocation,
    TextBlock,
)
from doclens_ai.services.ingestion import canonicalize_chunk
from tests.unit.test_document_ir import provenance
from tests.unit.test_extraction import _write_pdf
from tests.unit.test_ir_evidence import document, linked_table, policy, review


@pytest.mark.parametrize(
    "expression,cached,expected",
    [
        (None, None, ""),
        ("1-1", None, "formula=1-1"),
        (None, "0", "cached_value=0"),
        ("1-1", "0", "formula=1-1\ncached_value=0"),
        ("A1", "", "formula=A1\ncached_value="),
    ],
)
def test_formula_rendering_never_invents_literal_none(expression, cached, expected):
    value = FormulaElement(
        id="formula",
        source=SourceLocation(page_number=1),
        provenance=provenance(),
        expression=expression,
        cached_value=cached,
    )
    assert _render(value) == expected


def ordered_document():
    ir = document()
    first = ir.elements[0]
    elements = tuple(
        first.model_copy(
            update={
                "id": str(i),
                "source": replace(first.source, page_number=i + 1),
                "text": f"text-{i}",
            }
        )
        for i in range(3)
    )
    relations = tuple(
        Relation(
            id=f"order-{i}",
            kind="reading_order",
            source_id=str(i),
            target_id=str(i + 1),
            provenance=first.provenance,
        )
        for i in range(2)
    )
    return ir.model_copy(update={"elements": elements, "relations": relations})


def test_reading_order_does_not_force_an_oversized_atomic_group():
    ir = ordered_document()
    out = chunk_verified_ir(
        ir,
        expected=ir.source,
        review=review(ir),
        policy=policy(ir, allow_verified_pdf=True),
        max_characters=6,
    )
    assert out.status == "complete" and not out.held
    assert [c.text for c in out.chunks] == ["text-0", "text-1", "text-2"]
    assert [r for unit in out.units for r in unit.relations] == list(ir.relations)
    assert [c.ordinal for c in out.chunks] == [0, 1, 2]
    assert out == chunk_verified_ir(
        ir,
        expected=ir.source,
        review=review(ir),
        policy=policy(ir, allow_verified_pdf=True),
        max_characters=6,
    )


def test_bad_order_neighbor_does_not_invalidate_unrelated_semantic_evidence():
    ir = ordered_document()
    out = chunk_verified_ir(
        ir,
        expected=ir.source,
        review=review(ir, failed=("2",)),
        policy=policy(ir, allow_verified_pdf=True),
        max_characters=6,
    )
    assert [c.text for c in out.chunks] == ["text-0", "text-1"]
    assert out.status == "partial_review_required"


def test_semantic_relations_survive_wire_roundtrip_and_stay_atomic():
    ir = linked_table()
    restored = DocumentIR.model_validate_json(ir.model_dump_json())
    assert restored == ir
    out = chunk_verified_ir(
        restored,
        expected=ir.source,
        review=review(restored),
        policy=policy(ir, allow_verified_pdf=True),
    )
    assert len(out.chunks) == 1
    assert out.units[0].elements == ir.elements and out.units[0].relations == ir.relations
    assert all(span.text_block_map is None for span in out.chunks[0].provenance)
    held = chunk_verified_ir(
        ir,
        expected=ir.source,
        review=review(ir),
        policy=policy(ir, allow_verified_pdf=True),
        max_characters=6,
    )
    assert held.status == "blocked" and held.held[0][1] == "oversized_evidence"


@pytest.mark.parametrize("lexical", ["0", "-0", "-001.5000", "00123", "", None])
def test_cell_lexical_distinctions_survive_wire_and_rendering(lexical):
    value = CellElement(
        id="cell",
        source=SourceLocation(page_number=1, row_index=0, column_index=0),
        provenance=provenance(),
        row_index=0,
        column_index=0,
        text=lexical,
        raw_value=lexical,
    )
    ir = document((value,))
    restored = DocumentIR.model_validate_json(ir.model_dump_json())
    assert restored.elements[0].raw_value == lexical
    assert _render(restored.elements[0]) == (
        "" if lexical is None else f"text={lexical}\nraw_value={lexical}"
    )


def test_canonicalization_keeps_native_text_map_separate_from_generated_offsets():
    original = "  -001.5000 kg\r\n00123  "
    result = ExtractionResult(
        "1.1",
        DocumentFormat.DOCX,
        ExtractionStatus.SUCCESS,
        "a" * 64,
        100,
        blocks=(TextBlock(0, original, SourceLocation(block_index=0)),),
    )
    out = chunk_document(
        result,
        document_id=str(uuid4()),
        config=ChunkingConfig(
            strategy=ChunkingStrategy.STRUCTURE_AWARE,
        ),
    )
    chunk = out.chunks[0]
    span = chunk.provenance[0]
    assert span.text_block_map is not None
    saved = canonicalize_chunk(uuid4(), chunk).source_provenance[0]
    assert saved["offset_basis"] == "rendered_source_unit"
    assert saved["text_block_map"] == json.loads(json.dumps(asdict(span.text_block_map)))
    assert json.loads(json.dumps(saved)) == saved
    assert saved["text_block_map"]["offset_basis"] == "text_block_code_points"


@pytest.mark.parametrize("fmt", ["pdf", "docx", "xlsx"])
def test_real_source_to_ir_to_wire_retains_native_values_locators_and_failures(tmp_path, fmt):
    """Actual small source bytes, never a human Gold or model-quality assertion."""
    path = tmp_path / ("source." + fmt)
    literal = "-001.5000 kg; identifier 00123; original lexical evidence."
    if fmt == "pdf":
        _write_pdf(path, [literal])
    elif fmt == "docx":
        from docx import Document

        doc = Document()
        doc.add_paragraph(literal)
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Net kg 2026"
        table.cell(1, 0).text = "00123"
        table.cell(1, 1).text = "-001.5000"
        doc.save(path)
    else:
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.append(["00123", 0, -12.125, "=B1-B1", ""])
        book.save(path)
    payload = path.read_bytes()
    (observed,) = default_registry().extract(path, route=default_route())
    restored = DocumentIR.model_validate_json(observed.model_dump_json())
    assert restored == observed
    assert restored.native.to_dict() == observed.native.to_dict()
    assert restored.source.sha256 == hashlib.sha256(payload).hexdigest()
    assert restored.source.size_bytes == len(payload)
    assert all(e.verification.state == "unknown" for e in restored.elements)
    if fmt == "xlsx":
        cells = {c.coordinate: c for c in restored.native.sheets[0].cells}
        assert cells["A1"].value == "00123"
        assert cells["B1"].raw_value == "0"
        assert cells["C1"].raw_value == "-12.125"
        assert cells["D1"].formula == "B1-B1"
    else:
        assert restored.elements[0].text == literal

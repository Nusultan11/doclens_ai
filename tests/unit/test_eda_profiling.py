from __future__ import annotations

import csv
import json
import shutil
import sys
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from doclens_ai.eda import ProfileLimits, ProfilingError, profile_corpus, profiling
from doclens_ai.eda import __main__ as eda_main


def _write_pdf(path: Path, *, encrypted: bool = False) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    if encrypted:
        writer.encrypt("test-password")
    with path.open("wb") as destination:
        writer.write(destination)


def _write_docx(path: Path, text: str = "Document text") -> None:
    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    document.save(path)


def _write_xlsx(path: Path, secret_value: str = "value") -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet.append(["header", secret_value])
    worksheet.append(["row", 42])
    workbook.save(path)
    workbook.close()


def _by_format(result: dict[str, object], file_format: str) -> dict[str, object]:
    documents = result["documents"]
    assert isinstance(documents, list)
    return next(document for document in documents if document["format"] == file_format)


def test_profiles_valid_pdf_docx_and_xlsx(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    _write_pdf(corpus / "sample.pdf")
    _write_docx(corpus / "sample.docx")
    _write_xlsx(corpus / "sample.xlsx")

    result = profile_corpus(corpus, tmp_path / "output", measure_latency=False)

    assert result["aggregate"]["successful_extraction_count"] == 3
    pdf = _by_format(result, "pdf")
    assert pdf["pdf_page_count"] == 1
    assert pdf["possible_scanned_pdf"] is True
    docx = _by_format(result, "docx")
    assert docx["docx_paragraph_count"] == 1
    assert docx["docx_table_count"] == 1
    assert docx["docx_populated_table_cell_count"] == 2
    xlsx = _by_format(result, "xlsx")
    assert xlsx["xlsx_sheet_count"] == 1
    assert xlsx["xlsx_populated_cell_count"] == 4
    assert xlsx["xlsx_rows_with_data"] == 2


def test_records_empty_corrupt_encrypted_and_unsupported_files(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    (corpus / "empty.pdf").write_bytes(b"")
    (corpus / "corrupt.docx").write_bytes(b"not-a-zip")
    _write_pdf(corpus / "encrypted.pdf", encrypted=True)
    (corpus / "unsupported.txt").write_text("ignored content", encoding="utf-8")

    result = profile_corpus(corpus, tmp_path / "output", measure_latency=False)

    assert result["aggregate"]["error_category_counts"] == {
        "corrupt_document": 1,
        "empty_file": 1,
        "encrypted_document": 1,
        "unsupported_format": 1,
    }
    assert result["aggregate"]["successful_extraction_count"] == 0


def test_detects_duplicate_bytes_without_exposing_names(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    first = corpus / "private-name.docx"
    second = corpus / "other-private-name.docx"
    _write_docx(first, "confidential duplicate content")
    shutil.copyfile(first, second)

    result = profile_corpus(corpus, tmp_path / "output", measure_latency=False)
    documents = result["documents"]

    assert documents[0]["sha256"] == documents[1]["sha256"]
    assert documents[0]["duplicate_of"] is None
    assert documents[1]["duplicate_of"] == documents[0]["document_id"]
    assert result["aggregate"]["duplicate_copy_rate"] == {
        "numerator": 1,
        "denominator": 2,
        "percentage": 50.0,
    }
    serialized = json.dumps(result)
    assert "private-name" not in serialized
    assert "confidential duplicate content" not in serialized


def test_one_extractor_failure_does_not_stop_other_documents(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    _write_docx(corpus / "a.docx")
    _write_docx(corpus / "b.docx")
    original = profiling.EXTRACTORS["docx"]

    def fail_one(path: Path, limits: ProfileLimits) -> dict[str, object]:
        if path.name == "a.docx":
            raise RuntimeError("sensitive parser detail")
        return original(path, limits)

    monkeypatch.setitem(profiling.EXTRACTORS, "docx", fail_one)
    result = profile_corpus(corpus, tmp_path / "output", measure_latency=False)

    assert [document["extraction_success"] for document in result["documents"]] == [False, True]
    assert result["documents"][0]["error_category"] == "extraction_failed"
    assert "sensitive parser detail" not in json.dumps(result)


def test_enforces_file_and_spreadsheet_scan_limits(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    (corpus / "large.pdf").write_bytes(b"x" * 128)
    _write_xlsx(corpus / "cells.xlsx")

    result = profile_corpus(
        corpus,
        tmp_path / "output",
        limits=ProfileLimits(max_file_size_bytes=64, max_xlsx_cells_scanned=2),
        measure_latency=False,
    )

    assert result["aggregate"]["error_category_counts"] == {"limit_exceeded": 2}
    large = next(document for document in result["documents"] if document["file_size_bytes"] == 128)
    assert large["sha256"] is None


def test_repeated_content_reports_are_byte_for_byte_deterministic_and_private(
    tmp_path: Path,
) -> None:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    secret_name = "customer-alice-secret.docx"
    secret_content = "PERSONAL IDENTIFIER 12345"
    spreadsheet_secret = "SPREADSHEET SECRET 67890"
    _write_docx(corpus / secret_name, secret_content)
    _write_xlsx(corpus / "customer-bob-secret.xlsx", spreadsheet_secret)
    outputs = [tmp_path / "output-a", tmp_path / "output-b"]

    for output in outputs:
        profile_corpus(corpus, output, measure_latency=False)

    for report_name in ("metrics.json", "manifest.csv", "eda_report.md"):
        first = (outputs[0] / report_name).read_bytes()
        second = (outputs[1] / report_name).read_bytes()
        assert first == second
        assert secret_name.encode() not in first
        assert secret_content.encode() not in first
        assert spreadsheet_secret.encode() not in first
        assert str(corpus).encode() not in first


def test_manifest_has_fixed_privacy_safe_columns(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    _write_docx(corpus / "name-must-not-leak.docx")
    output = tmp_path / "output"

    profile_corpus(corpus, output, measure_latency=False)

    with (output / "manifest.csv").open(encoding="utf-8", newline="") as source:
        rows = list(csv.DictReader(source))
    assert tuple(rows[0]) == profiling.MANIFEST_FIELDS
    assert all("name" not in column and "path" not in column for column in rows[0])
    assert "name-must-not-leak" not in json.dumps(rows)


def test_aggregate_denominators_exclude_unsupported_from_extraction_rate(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    _write_docx(corpus / "valid.docx")
    (corpus / "corrupt.docx").write_bytes(b"bad")
    (corpus / "empty.xlsx").write_bytes(b"")
    (corpus / "unsupported.bin").write_bytes(b"data")

    result = profile_corpus(corpus, tmp_path / "output", measure_latency=False)

    assert result["aggregate"]["extraction_success_rate"] == {
        "numerator": 1,
        "denominator": 3,
        "percentage": 33.33,
    }
    assert result["aggregate"]["empty_text_rate"]["denominator"] == 1
    assert "unsupported" not in result["aggregate"]["extraction_success_rate_by_format"]
    assert result["aggregate"]["file_size_bytes"]["p50"] is None
    assert (
        result["aggregate"]["file_size_bytes"]["percentiles_unavailable_reason"]
        == "fewer_than_5_samples"
    )


def test_empty_corpus_writes_reports_with_explicit_zero_denominators(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    output = tmp_path / "output"

    result = profile_corpus(corpus, output, measure_latency=False)

    assert result["aggregate"]["document_count"] == 0
    assert result["aggregate"]["extraction_success_rate"] == {
        "numerator": 0,
        "denominator": 0,
        "percentage": None,
    }
    assert {path.name for path in output.iterdir()} == {
        "eda_report.md",
        "manifest.csv",
        "metrics.json",
    }
    report = (output / "eda_report.md").read_text(encoding="utf-8")
    assert "## Format distribution" in report
    assert "## Distributions" in report


@pytest.mark.parametrize("kind", ["missing", "file", "same"])
def test_rejects_invalid_directory_configuration(tmp_path: Path, kind: str) -> None:
    input_path = tmp_path / "corpus"
    output_path = tmp_path / "output"
    if kind == "file":
        input_path.write_text("not a directory", encoding="utf-8")
    elif kind == "same":
        input_path.mkdir()
        output_path = input_path

    with pytest.raises(ProfilingError):
        profile_corpus(input_path, output_path, measure_latency=False)


def test_cli_continues_past_corrupt_and_unsupported_files(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    _write_docx(corpus / "valid.docx")
    _write_pdf(corpus / "valid.pdf")
    _write_xlsx(corpus / "valid.xlsx")
    (corpus / "corrupt.pdf").write_bytes(b"bad")
    (corpus / "unsupported.txt").write_text("data", encoding="utf-8")
    output = tmp_path / "output"
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "doclens_ai.eda",
            "--input-dir",
            str(corpus),
            "--output-dir",
            str(output),
            "--omit-runtime-metrics",
        ],
    )

    assert eda_main.main() == 0
    assert "5 files" in capsys.readouterr().out
    metrics = json.loads((output / "metrics.json").read_text(encoding="utf-8"))
    assert metrics["aggregate"]["successful_extraction_count"] == 3
    assert metrics["aggregate"]["error_category_counts"] == {
        "corrupt_document": 1,
        "unsupported_format": 1,
    }


def test_measured_latency_is_present_but_excluded_from_content_fingerprint(
    tmp_path: Path,
) -> None:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    _write_docx(corpus / "sample.docx")
    ticks = iter([1_000_000, 3_500_000])

    measured = profile_corpus(
        corpus,
        tmp_path / "measured",
        measure_latency=True,
        clock_ns=lambda: next(ticks),
    )
    deterministic = profile_corpus(corpus, tmp_path / "stable", measure_latency=False)

    assert measured["documents"][0]["extraction_latency_ms"] == 2.5
    assert measured["deterministic_content_sha256"] == deterministic["deterministic_content_sha256"]

from __future__ import annotations

import asyncio
import io
from contextlib import asynccontextmanager
from datetime import timedelta
from uuid import uuid4

import pytest
from docx import Document as DocxDocument
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

import doclens_ai.services.processing as processing_module
from doclens_ai.models import (
    CanonicalChunk,
    Document,
    DocumentVersion,
    ProcessingJob,
    ProcessingRun,
    ProcessingStage,
    Workspace,
)
from doclens_ai.services.ingestion import IngestionService
from doclens_ai.services.processing import (
    InvalidProcessingTransition,
    LeaseOwnershipError,
    ProcessingCoordinator,
    ProcessingPolicy,
)
from doclens_ai.services.upload import DOCX_CONTENT_TYPE, PDF_CONTENT_TYPE, UploadService
from doclens_ai.services.worker import IngestionWorker, WorkerSettings
from doclens_ai.storage.s3 import ObjectStorageError, S3ObjectStorage

pytestmark = [pytest.mark.integration, pytest.mark.storage]


@pytest.mark.parametrize("reclaim", [False, True])
def test_heartbeat_failure_prevents_unsafe_handoff(
    migrated_database_url,
    test_storage_settings,
    monkeypatch,
    reclaim,
):
    async def exercise():
        async with _environment(migrated_database_url, test_storage_settings) as (
            factory,
            storage,
            keys,
        ):
            result = await _upload(factory, storage, keys)
            a = _worker(factory, storage)
            persisted, heartbeat_failed, release = asyncio.Event(), asyncio.Event(), asyncio.Event()
            original = a.ingestion.ingest

            async def paused(*args, **kwargs):
                outcome = await original(*args, **kwargs)
                persisted.set()
                await release.wait()
                return outcome

            async def broken_heartbeat(**kwargs):
                heartbeat_failed.set()
                raise RuntimeError("private DB exception details")

            async def expire():
                async with factory() as session, session.begin():
                    await session.execute(
                        text(
                            "UPDATE processing_jobs SET lease_expires_at = clock_timestamp() "
                            "- INTERVAL '1 second' WHERE id = :id"
                        ),
                        {"id": result.processing_job_id},
                    )

            monkeypatch.setattr(a.ingestion, "ingest", paused)
            monkeypatch.setattr(a.coordinator, "heartbeat", broken_heartbeat)
            task = asyncio.create_task(a.run_once())
            try:
                await asyncio.wait_for(persisted.wait(), 5)
                await asyncio.wait_for(heartbeat_failed.wait(), 5)
                async with factory() as session:
                    ids = set(await session.scalars(select(CanonicalChunk.id)))
                b = _worker(factory, storage, identity="worker-b")
                if reclaim:
                    await expire()
                    assert await b.run_once()
                release.set()
                await asyncio.wait_for(task, 5)
                if not reclaim:
                    async with factory() as session:
                        job = await session.get(ProcessingJob, result.processing_job_id)
                        run = await session.get(ProcessingRun, job.processing_run_id)
                        assert job.status == "leased" and job.worker_id == a.worker_id
                        assert run.stage == ProcessingStage.EXTRACTION
                    await expire()
                    assert await b.run_once()
                _, final_ids = await _assert_handoff(factory, result)
                assert set(final_ids) == ids
                assert len(asyncio.all_tasks()) == 1
            finally:
                release.set()
                await task

    asyncio.run(exercise())


def test_handoff_expiry_after_flush_rolls_back_atomically(
    migrated_database_url,
    test_storage_settings,
    monkeypatch,
):
    async def exercise():
        async with _environment(migrated_database_url, test_storage_settings) as (
            factory,
            storage,
            keys,
        ):
            result = await _upload(factory, storage, keys)
            a = _worker(factory, storage)
            claim = await a.coordinator.claim_next(worker_id=a.worker_id)
            await a.ingestion.ingest(
                claim.processing_run_id, worker_id=a.worker_id, lease_token=claim.lease_token
            )
            original = processing_module._processing_now
            calls = 0

            async def expire_at_commit(session, value=None):
                nonlocal calls
                calls += 1
                if calls == 3:
                    # The handoff UPDATEs have flushed; wall-clock expiry must roll them back.
                    job = await session.get(ProcessingJob, claim.job_id)
                    assert job.status == "queued" and job.lease_token is None
                    async with factory() as observer:
                        assert (
                            await observer.scalar(
                                select(ProcessingJob.id)
                                .where(ProcessingJob.id == claim.job_id)
                                .with_for_update(skip_locked=True)
                            )
                            is None
                        )
                        assert (
                            await observer.scalar(
                                select(ProcessingRun.id)
                                .where(ProcessingRun.id == claim.processing_run_id)
                                .with_for_update(skip_locked=True)
                            )
                            is None
                        )
                        visible_run = await observer.get(ProcessingRun, claim.processing_run_id)
                        assert visible_run.stage == ProcessingStage.EXTRACTION
                    await session.scalar(select(func.pg_sleep_until(claim.lease_expires_at)))
                return await original(session, value)

            monkeypatch.setattr(processing_module, "_processing_now", expire_at_commit)
            with pytest.raises(LeaseOwnershipError):
                await a.coordinator.handoff_extraction(
                    job_id=claim.job_id, worker_id=a.worker_id, lease_token=claim.lease_token
                )
            assert calls == 3
            async with factory() as session:
                job = await session.get(ProcessingJob, result.processing_job_id)
                run = await session.get(ProcessingRun, claim.processing_run_id)
                assert job.status == "leased" and job.lease_token == claim.lease_token
                assert job.attempt_count == 1 and run.stage == ProcessingStage.EXTRACTION
                assert run.status == "processing"

    asyncio.run(exercise())


@asynccontextmanager
async def _environment(url, settings):
    engine = create_async_engine(url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    storage = S3ObjectStorage(settings)
    keys = []
    try:
        await storage.probe()
        yield factory, storage, keys
    finally:
        for key in keys:
            await storage.delete(key)
        await asyncio.to_thread(storage._client.close)
        await engine.dispose()


async def _upload(factory, storage, keys, *, corrupt=False):
    async with factory() as session, session.begin():
        document = Document(workspace=Workspace(name="Worker fixture"))
        session.add(document)
        await session.flush()
        document_id, workspace_id = document.id, document.workspace_id
    stream = io.BytesIO()
    if corrupt:
        stream.write(b"%PDF-1.7\nnot a readable document\n%%EOF\n")
    else:
        document = DocxDocument()
        document.add_paragraph("A genuine fixture document with enough extraction content.")
        document.add_table(rows=1, cols=2).cell(0, 0).text = "Table fixture value"
        document.save(stream)
    stream.seek(0)
    result = await UploadService(
        factory, storage, max_upload_size_bytes=1024 * 1024
    ).upload_and_confirm(
        workspace_id=workspace_id,
        document_id=document_id,
        idempotency_key=uuid4().hex,
        original_filename="fixture.pdf" if corrupt else "fixture.docx",
        content_type=PDF_CONTENT_TYPE if corrupt else DOCX_CONTENT_TYPE,
        stream=stream,
    )
    async with factory() as session:
        version = await session.get(DocumentVersion, result.document_version_id)
        keys.append(version.object_storage_key)
    return result


def _worker(factory, storage, *, identity="worker-a"):
    settings = WorkerSettings(
        _env_file=None,
        lease_seconds=2,
        heartbeat_interval_seconds=0.03,
        db_timeout_seconds=0.3,
        poll_interval_seconds=0.02,
    )
    coordinator = ProcessingCoordinator(
        factory,
        policy=ProcessingPolicy(
            lease_duration=timedelta(seconds=settings.lease_seconds),
            retry_base_delay=timedelta(seconds=1),
        ),
    )
    worker = IngestionWorker(
        coordinator, IngestionService(factory, storage), settings=settings, worker_id=identity
    )
    return worker


async def _assert_handoff(factory, result):
    async with factory() as session:
        job = await session.get(ProcessingJob, result.processing_job_id)
        run = await session.get(ProcessingRun, job.processing_run_id)
        assert job.status == "queued" and run.status == "queued"
        assert run.stage == ProcessingStage.EMBEDDING
        assert job.attempt_count == 0 and job.max_attempts == 3
        assert job.worker_id is None and job.lease_token is None and job.lease_expires_at is None
        assert job.finished_at is None and job.next_retry_at is None and job.failure_reason is None
        chunks = tuple(
            await session.scalars(
                select(CanonicalChunk).where(CanonicalChunk.processing_run_id == run.id)
            )
        )
        assert chunks
        assert len({chunk.id for chunk in chunks}) == len(chunks)
        return run.id, tuple(chunk.id for chunk in chunks)


def test_upload_ingestion_handoff_and_embedding_filter(
    migrated_database_url, test_storage_settings
):
    async def exercise():
        async with _environment(migrated_database_url, test_storage_settings) as (
            factory,
            storage,
            keys,
        ):
            result = await _upload(factory, storage, keys)
            async with factory() as session:
                before = (
                    await session.execute(
                        text("SELECT * FROM document_versions WHERE id = :id"),
                        {"id": result.document_version_id},
                    )
                ).one()
            worker = _worker(factory, storage)
            assert await worker.run_once()
            run_id, _ = await _assert_handoff(factory, result)
            assert not await worker.run_once()
            # Unfiltered callers retain their original ability to claim any stage.
            claim = await worker.coordinator.claim_next(worker_id="embedding-future")
            assert claim.processing_run_id == run_id and claim.stage is ProcessingStage.EMBEDDING
            assert claim.attempt_count == 1
            async with factory() as session:
                assert (
                    await session.execute(
                        text("SELECT * FROM document_versions WHERE id = :id"),
                        {"id": result.document_version_id},
                    )
                ).one() == before

    asyncio.run(exercise())


def test_competing_workers_and_independent_progression(
    migrated_database_url, test_storage_settings, monkeypatch
):
    async def exercise():
        async with _environment(migrated_database_url, test_storage_settings) as (
            factory,
            storage,
            keys,
        ):
            first = await _upload(factory, storage, keys)
            a, b = _worker(factory, storage), _worker(factory, storage, identity="worker-b")
            entered, release = asyncio.Event(), asyncio.Event()
            original = a.ingestion.ingest

            async def paused(*args, **kwargs):
                entered.set()
                await release.wait()
                return await original(*args, **kwargs)

            monkeypatch.setattr(a.ingestion, "ingest", paused)
            task = asyncio.create_task(a.run_once())
            try:
                await asyncio.wait_for(entered.wait(), 5)
                assert not await b.run_once()
                second = await _upload(factory, storage, keys)
                assert await b.run_once()
                await _assert_handoff(factory, second)
                release.set()
                assert await asyncio.wait_for(task, 5)
                await _assert_handoff(factory, first)
                assert not await a.run_once() and not await b.run_once()
            finally:
                release.set()
                await task

    asyncio.run(exercise())


def test_crash_cut_after_persistence_reclaim_and_stale_handoff(
    migrated_database_url, test_storage_settings
):
    async def exercise():
        async with _environment(migrated_database_url, test_storage_settings) as (
            factory,
            storage,
            keys,
        ):
            result = await _upload(factory, storage, keys)
            a = _worker(factory, storage)
            claim = await a.coordinator.claim_next(
                worker_id=a.worker_id, stage=ProcessingStage.EXTRACTION
            )
            await a.ingestion.ingest(
                claim.processing_run_id, worker_id=a.worker_id, lease_token=claim.lease_token
            )
            # Crash cut: persisted chunks survive, but no handoff was committed.
            async with factory() as session:
                original_ids = set(await session.scalars(select(CanonicalChunk.id)))
                assert original_ids
            async with factory() as session, session.begin():
                await session.execute(
                    text(
                        "UPDATE processing_jobs SET lease_expires_at = clock_timestamp() "
                        "- INTERVAL '1 second' WHERE id = :id"
                    ),
                    {"id": claim.job_id},
                )
            with pytest.raises(LeaseOwnershipError):
                await a.coordinator.handoff_extraction(
                    job_id=claim.job_id, worker_id=a.worker_id, lease_token=claim.lease_token
                )
            b = _worker(factory, storage, identity="worker-b")
            replacement = await b.coordinator.claim_next(
                worker_id=b.worker_id, stage=ProcessingStage.EXTRACTION
            )
            assert replacement.lease_token != claim.lease_token and replacement.attempt_count == 2
            with pytest.raises(LeaseOwnershipError):
                await a.coordinator.handoff_extraction(
                    job_id=claim.job_id, worker_id=a.worker_id, lease_token=claim.lease_token
                )
            await b._execute(replacement)
            _, ids = await _assert_handoff(factory, result)
            assert set(ids) == original_ids
            with pytest.raises(LeaseOwnershipError):
                await a.coordinator.handoff_extraction(
                    job_id=claim.job_id, worker_id=a.worker_id, lease_token=claim.lease_token
                )

    asyncio.run(exercise())


def test_heartbeat_renews_lease_past_original_expiry(
    migrated_database_url, test_storage_settings, monkeypatch
):
    async def exercise():
        async with _environment(migrated_database_url, test_storage_settings) as (
            factory,
            storage,
            keys,
        ):
            result = await _upload(factory, storage, keys)
            worker = _worker(factory, storage)
            original = worker.ingestion.ingest
            heartbeat = asyncio.Event()
            original_heartbeat = worker.coordinator.heartbeat

            async def observed_heartbeat(**kwargs):
                expiry = await original_heartbeat(**kwargs)
                heartbeat.set()
                return expiry

            async def slow(run_id, **kwargs):
                async with factory() as session:
                    job = await session.get(ProcessingJob, result.processing_job_id)
                    initial_expiry = job.lease_expires_at
                await asyncio.wait_for(heartbeat.wait(), 5)
                async with factory() as session:
                    await session.scalar(
                        select(func.pg_sleep_until(initial_expiry + timedelta(milliseconds=100)))
                    )
                    job = await session.get(ProcessingJob, result.processing_job_id)
                    assert job.last_heartbeat_at > job.leased_at
                    assert job.lease_expires_at > await session.scalar(
                        select(func.clock_timestamp())
                    )
                assert (
                    await _worker(factory, storage, identity="worker-b").coordinator.claim_next(
                        worker_id="worker-b", stage=ProcessingStage.EXTRACTION
                    )
                    is None
                )
                return await original(run_id, **kwargs)

            monkeypatch.setattr(worker.coordinator, "heartbeat", observed_heartbeat)
            monkeypatch.setattr(worker.ingestion, "ingest", slow)
            assert await worker.run_once()
            await _assert_handoff(factory, result)
            assert len(asyncio.all_tasks()) == 1

    asyncio.run(exercise())


def test_retryable_storage_failure_recovers(
    migrated_database_url, test_storage_settings, monkeypatch
):
    async def exercise():
        async with _environment(migrated_database_url, test_storage_settings) as (
            factory,
            storage,
            keys,
        ):
            result = await _upload(factory, storage, keys)
            worker = _worker(factory, storage)
            original = storage.inspect

            async def unavailable(*args, **kwargs):
                raise ObjectStorageError("private storage error")

            monkeypatch.setattr(storage, "inspect", unavailable)
            assert await worker.run_once()
            async with factory() as session, session.begin():
                job = await session.get(ProcessingJob, result.processing_job_id)
                assert job.status == "retry_wait" and job.attempt_count == 1
                assert job.failure_reason == "storage_unavailable" and job.lease_token is None
                assert await session.scalar(select(func.count(CanonicalChunk.id))) == 0
                await session.execute(
                    text(
                        "UPDATE processing_jobs SET next_retry_at = clock_timestamp() "
                        "- INTERVAL '1 second' WHERE id = :id"
                    ),
                    {"id": job.id},
                )
            monkeypatch.setattr(storage, "inspect", original)
            assert await worker.run_once()
            await _assert_handoff(factory, result)

    asyncio.run(exercise())


def test_permanent_failure_does_not_stop_unrelated_jobs(
    migrated_database_url, test_storage_settings
):
    async def exercise():
        async with _environment(migrated_database_url, test_storage_settings) as (
            factory,
            storage,
            keys,
        ):
            bad = await _upload(factory, storage, keys, corrupt=True)
            good = await _upload(factory, storage, keys)
            worker = _worker(factory, storage)
            assert await worker.run_once()
            async with factory() as session:
                job = await session.get(ProcessingJob, bad.processing_job_id)
                run = await session.get(ProcessingRun, job.processing_run_id)
                assert job.status == "failed" and run.status == "failed"
                assert job.attempt_count == 1 and job.max_attempts == 3
                assert job.failure_reason == "extraction_failed" and job.finished_at is not None
                assert job.lease_token is None
                assert await session.scalar(select(func.count(CanonicalChunk.id))) == 0
            assert await worker.run_once()
            await _assert_handoff(factory, good)

    asyncio.run(exercise())


def test_handoff_requires_chunks_and_keeps_sequential_advance_semantics(
    migrated_database_url, test_storage_settings
):
    async def exercise():
        async with _environment(migrated_database_url, test_storage_settings) as (
            factory,
            storage,
            keys,
        ):
            await _upload(factory, storage, keys)
            worker = _worker(factory, storage)
            claim = await worker.coordinator.claim_next(worker_id=worker.worker_id)
            ownership = dict(
                job_id=claim.job_id, worker_id=worker.worker_id, lease_token=claim.lease_token
            )
            with pytest.raises(InvalidProcessingTransition):
                await worker.coordinator.handoff_extraction(**ownership)
            with pytest.raises(InvalidProcessingTransition):
                await worker.coordinator.advance_stage(**ownership, stage=ProcessingStage.EMBEDDING)
            await worker.ingestion.ingest(
                claim.processing_run_id, worker_id=worker.worker_id, lease_token=claim.lease_token
            )
            await worker.coordinator.advance_stage(**ownership, stage=ProcessingStage.CLEANING)
            with pytest.raises(InvalidProcessingTransition):
                await worker.coordinator.handoff_extraction(**ownership)
            async with factory() as session:
                job = await session.get(ProcessingJob, claim.job_id)
                run = await session.get(ProcessingRun, claim.processing_run_id)
                assert job.status == "leased" and job.lease_token == claim.lease_token
                assert run.stage == ProcessingStage.CLEANING and run.status == "processing"

    asyncio.run(exercise())

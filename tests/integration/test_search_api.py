from __future__ import annotations

import asyncio
import io
import os
from contextlib import asynccontextmanager
from datetime import timedelta
from uuid import UUID, uuid4

import pytest
from docx import Document as DocxDocument
from sqlalchemy import delete, select

from doclens_ai.db.session import Database
from doclens_ai.models import User, WorkspaceMembership
from doclens_ai.services.index_worker import IndexWorker
from doclens_ai.services.indexing import IndexingService, collection_name, embedding_identity
from doclens_ai.services.processing import ProcessingCoordinator, ProcessingPolicy
from doclens_ai.services.retrieval import QueryEncoder, RetrievalService
from doclens_ai.services.upload import DOCX_CONTENT_TYPE
from doclens_ai.vector.qdrant import ProductionQdrant
from tests.integration.test_authenticated_upload import _api_context
from tests.integration.test_production_indexing import _extraction_worker
from tests.unit.test_indexing import FakeEmbedder
from tests.unit.test_retrieval import QueryFixtureEmbedder

pytestmark = [pytest.mark.integration, pytest.mark.storage]


@asynccontextmanager
async def _search_api_context(url, settings, *, initialize=True):
    endpoint = os.getenv("TEST_QDRANT_URL")
    if not endpoint:
        pytest.skip("TEST_QDRANT_URL is not configured")
    provider = FakeEmbedder()
    name = collection_name(embedding_identity(provider)[1], "testapi_" + uuid4().hex)
    vector = ProductionQdrant(endpoint, name)
    async with _api_context(url, settings) as context:
        keys = []
        try:
            if initialize:
                request = await context.client.post(
                    f"/api/v1/workspaces/{context.workspace_id}/documents",
                    headers=context.authorization(context.owner_subject),
                )
                assert request.status_code == 201
                document_id = request.json()["id"]
                stream = io.BytesIO()
                document = DocxDocument()
                document.add_paragraph("Controlled authenticated search fixture paragraph.")
                document.add_table(rows=1, cols=2).cell(0, 0).text = "Controlled authorized table"
                document.save(stream)
                response = await context.client.post(
                    f"/api/v1/workspaces/{context.workspace_id}/documents/{document_id}/versions",
                    headers=context.authorization(context.owner_subject)
                    | {"Idempotency-Key": uuid4().hex},
                    files={"file": ("fixture.docx", stream.getvalue(), DOCX_CONTENT_TYPE)},
                )
                assert response.status_code == 201
                from doclens_ai.models import DocumentVersion

                async with context.database.session_factory() as session:
                    version = await session.get(
                        DocumentVersion, UUID(response.json()["document_version_id"])
                    )
                    keys.append(version.object_storage_key)
                assert await _extraction_worker(
                    context.database.session_factory, context.storage
                ).run_once()
                worker = IndexWorker(
                    ProcessingCoordinator(
                        context.database.session_factory,
                        policy=ProcessingPolicy(lease_duration=timedelta(seconds=120)),
                    ),
                    IndexingService(context.database.session_factory, provider, vector),
                )
                assert await worker.run_once()
                context.application.state.retrieval_service = RetrievalService(
                    context.database.session_factory, QueryEncoder(QueryFixtureEmbedder()), vector
                )
            yield context, vector
        finally:
            for key in keys:
                await context.storage.delete(key)
            await vector.close()
            import httpx

            async with httpx.AsyncClient(base_url=endpoint, timeout=5) as client:
                assert (await client.delete("/collections/" + name)).status_code in {200, 404}


def _body(context, **kwargs):
    return {"workspace_id": str(context.workspace_id), "query": "Controlled table query", **kwargs}


@pytest.mark.anyio
async def test_search_authentication_default_fail_closed_and_valid_results(
    migrated_database_url, test_storage_settings
):
    async with _search_api_context(
        migrated_database_url, test_storage_settings, initialize=False
    ) as (ctx, _):
        assert (await ctx.client.post("/api/v1/search", json=_body(ctx))).status_code == 401
        assert (
            await ctx.client.post(
                "/api/v1/search", json=_body(ctx), headers={"Authorization": "Bearer invalid"}
            )
        ).status_code == 401
        response = await ctx.client.post(
            "/api/v1/search", json=_body(ctx), headers=ctx.authorization(ctx.viewer_subject)
        )
        assert response.status_code == 503 and response.json() == {
            "detail": "search service is unavailable"
        }
        assert (await ctx.client.get("/api/v1/health/ready")).status_code == 200


@pytest.mark.anyio
async def test_search_verified_source_locators_and_workspace_denial(
    migrated_database_url, test_storage_settings
):
    async with _search_api_context(migrated_database_url, test_storage_settings) as (ctx, _):
        response = await ctx.client.post(
            "/api/v1/search", json=_body(ctx), headers=ctx.authorization(ctx.viewer_subject)
        )
        assert response.status_code == 200
        results = response.json()["results"]
        assert len(results) == 2
        assert all(
            set(hit)
            == {"chunk_id", "document_id", "document_version_id", "score", "text", "source_locator"}
            for hit in results
        )
        assert all(hit["source_locator"]["document_format"] == "docx" for hit in results)
        assert any(
            hit["source_locator"]["spans"][0]["source"].get("table_index") == 0 for hit in results
        )
        assert "object_storage_key" not in response.text and "model_dir" not in response.text
        denied = await ctx.client.post(
            "/api/v1/search", json=_body(ctx), headers=ctx.authorization(ctx.outsider_subject)
        )
        assert denied.status_code == 403 and "Controlled" not in denied.text
        wrong_workspace = await ctx.client.post(
            "/api/v1/search",
            json=_body(ctx, workspace_id=str(ctx.other_workspace_id)),
            headers=ctx.authorization(ctx.viewer_subject),
        )
        assert wrong_workspace.status_code == 403
        missing = await ctx.client.post(
            "/api/v1/search",
            json=_body(ctx, document_ids=[str(uuid4())]),
            headers=ctx.authorization(ctx.viewer_subject),
        )
        assert missing.status_code == 200 and missing.json() == {"results": []}
        concurrent = await asyncio.gather(
            *(
                ctx.client.post(
                    "/api/v1/search", json=_body(ctx), headers=ctx.authorization(ctx.viewer_subject)
                )
                for _ in range(6)
            )
        )
        assert all(r.status_code == 200 and len(r.json()["results"]) == 2 for r in concurrent)


@pytest.mark.anyio
async def test_api_strict_bounds_and_dependency_error_sanitization(
    migrated_database_url, test_storage_settings, monkeypatch
):
    async with _search_api_context(migrated_database_url, test_storage_settings) as (ctx, vector):
        headers = ctx.authorization(ctx.viewer_subject)
        for changes in (
            {"query": ""},
            {"query": " "},
            {"query": "x" * 2001},
            {"top_k": 21},
            {"top_k": True},
            {"limit": 0},
            {"document_ids": []},
            {"document_ids": [str(uuid4()) for _ in range(51)]},
            {"filter": {}},
            {"top_k": 5, "limit": 10},
        ):
            assert (
                await ctx.client.post("/api/v1/search", json=_body(ctx, **changes), headers=headers)
            ).status_code == 422
        service = ctx.application.state.retrieval_service

        def broken(*args, **kwargs):
            raise RuntimeError("private query model path credentials")

        with monkeypatch.context() as patch:
            patch.setattr(service.encoder.provider, "encode", broken)
            response = await ctx.client.post("/api/v1/search", json=_body(ctx), headers=headers)
            assert response.status_code == 503 and response.json() == {
                "detail": "search service is unavailable"
            }
        unavailable = ProductionQdrant("http://127.0.0.1:1", vector.collection, timeout_seconds=0.1)
        try:
            service.vector = unavailable
            response = await ctx.client.post("/api/v1/search", json=_body(ctx), headers=headers)
            assert response.status_code == 503 and response.json() == {
                "detail": "search service is unavailable"
            }
        finally:
            service.vector = vector
            await unavailable.close()
        bad = Database("postgresql+asyncpg://doclens:doclens-local@127.0.0.1:1/unused")
        try:
            service.factory = bad.session_factory
            response = await ctx.client.post("/api/v1/search", json=_body(ctx), headers=headers)
            assert response.status_code == 503 and response.json() == {
                "detail": "search service is unavailable"
            }
        finally:
            service.factory = ctx.database.session_factory
            await bad.dispose()


@pytest.mark.anyio
async def test_api_membership_revocation_before_release(
    migrated_database_url, test_storage_settings, monkeypatch
):
    async with _search_api_context(migrated_database_url, test_storage_settings) as (ctx, vector):
        original = vector.search

        async def revoke(*args, **kwargs):
            hits = await original(*args, **kwargs)
            async with ctx.database.session_factory() as session, session.begin():
                user_id = await session.scalar(
                    select(User.id).where(User.subject == ctx.viewer_subject)
                )
                await session.execute(
                    delete(WorkspaceMembership).where(WorkspaceMembership.user_id == user_id)
                )
            return hits

        monkeypatch.setattr(vector, "search", revoke)
        response = await ctx.client.post(
            "/api/v1/search", json=_body(ctx), headers=ctx.authorization(ctx.viewer_subject)
        )
        assert response.status_code == 403 and "Controlled" not in response.text


@pytest.mark.anyio
async def test_configured_missing_model_is_sanitized_startup_failure(
    migrated_database_url, test_storage_settings, monkeypatch, caplog
):
    monkeypatch.setenv("SEARCH_ENABLED", "true")
    monkeypatch.setenv("SEARCH_MODEL_DIR", "/missing-private-model-location")
    async with _search_api_context(
        migrated_database_url, test_storage_settings, initialize=False
    ) as (ctx, _):
        assert ctx.application.state.retrieval_service is None
        response = await ctx.client.post(
            "/api/v1/search", json=_body(ctx), headers=ctx.authorization(ctx.viewer_subject)
        )
        assert response.status_code == 503 and response.json() == {
            "detail": "search service is unavailable"
        }
        assert "missing-private-model-location" not in caplog.text

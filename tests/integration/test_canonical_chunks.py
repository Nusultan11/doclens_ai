from __future__ import annotations

import asyncio
import hashlib
import io
from collections.abc import AsyncIterator, Callable
from dataclasses import dataclass, replace
from uuid import UUID, uuid4

import pytest
from botocore.exceptions import ClientError
from docx import Document as DocxDocument
from openpyxl import Workbook
from openpyxl.worksheet.table import Table as XlsxTable
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

import doclens_ai.services.ingestion as ingestion_module
from doclens_ai.models import (
    CanonicalChunk,
    Document,
    DocumentVersion,
    ProcessingJob,
    ProcessingJobStatus,
    ProcessingOperationKind,
    ProcessingRun,
    ProcessingStage,
    ProcessingStatus,
    Workspace,
)
from doclens_ai.services.ingestion import (
    IngestionFailure,
    IngestionFailureCategory,
    IngestionResultStatus,
    IngestionService,
    canonical_chunk_uuid,
)
from doclens_ai.services.processing import (
    PROCESSING_CONFIGURATION_VERSION,
    ClaimedProcessingJob,
    InvalidProcessingTransition,
    ProcessingCoordinator,
    default_processing_configuration_fingerprint,
)
from doclens_ai.services.upload import (
    DOCX_CONTENT_TYPE,
    PDF_CONTENT_TYPE,
    XLSX_CONTENT_TYPE,
    UploadService,
)
from doclens_ai.storage.config import StorageSettings
from doclens_ai.storage.s3 import ObjectMetadata, ObjectStorageError, S3ObjectStorage


@dataclass
class _StoredObject:
    content: bytes
    content_type: str
    checksum: str


class MemoryStorage:
    def __init__(self) -> None:
        self.objects: dict[str, _StoredObject] = {}
        self.available = True

    def _require(self, key: str) -> _StoredObject:
        if not self.available:
            raise ObjectStorageError("storage unavailable")
        try:
            return self.objects[key]
        except KeyError:
            try:
                raise ClientError(
                    {"Error": {"Code": "NoSuchKey", "Message": "missing"}},
                    "GetObject",
                )
            except ClientError as exc:
                raise ObjectStorageError("object unavailable") from exc

    async def put_stream(
        self,
        key: str,
        stream: io.BytesIO,
        *,
        content_type: str,
        sha256_checksum: str,
    ) -> None:
        if not self.available:
            raise ObjectStorageError("storage unavailable")
        stream.seek(0)
        self.objects[key] = _StoredObject(
            content=stream.read(),
            content_type=content_type,
            checksum=sha256_checksum,
        )

    async def inspect(self, key: str) -> ObjectMetadata:
        stored = self._require(key)
        return ObjectMetadata(
            size=len(stored.content),
            content_type=stored.content_type,
            metadata={"sha256": stored.checksum},
        )

    async def sha256(self, key: str, *, chunk_size: int = 1024 * 1024) -> str:
        return hashlib.sha256(self._require(key).content).hexdigest()

    async def stream_read(
        self,
        key: str,
        *,
        chunk_size: int = 1024 * 1024,
    ) -> AsyncIterator[bytes]:
        content = self._require(key).content
        for offset in range(0, len(content), chunk_size):
            yield content[offset : offset + chunk_size]

    async def delete(self, key: str) -> None:
        self.objects.pop(key, None)


def _pdf_bytes(page_texts: list[str]) -> bytes:
    output = io.BytesIO()
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_reference = writer._add_object(font)  # noqa: SLF001
    for page_text in page_texts:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_reference})}
        )
        escaped = page_text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1"))
        page[NameObject("/Contents")] = writer._add_object(stream)  # noqa: SLF001
    writer.write(output)
    return output.getvalue()


def _docx_bytes(*, empty: bool = False, text: str | None = None) -> bytes:
    output = io.BytesIO()
    document = DocxDocument()
    if not empty:
        document.add_paragraph(text or "A production document paragraph with useful content.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "left value"
        table.cell(0, 1).text = "right value"
    document.save(output)
    return output.getvalue()


def _xlsx_bytes() -> bytes:
    output = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Metrics"
    sheet.append(["name", "value"])
    sheet.append(["alpha", 7])
    sheet.add_table(XlsxTable(displayName="MetricsTable", ref="A1:B2"))
    workbook.save(output)
    return output.getvalue()


async def _create_document(
    session_factory: async_sessionmaker[AsyncSession],
) -> tuple[UUID, UUID]:
    async with session_factory() as session, session.begin():
        workspace = Workspace(name=f"Canonical chunks {uuid4()}")
        document = Document(workspace=workspace)
        session.add(document)
        await session.flush()
        return workspace.id, document.id


async def _upload(
    session_factory: async_sessionmaker[AsyncSession],
    storage: MemoryStorage | S3ObjectStorage,
    *,
    workspace_id: UUID,
    document_id: UUID,
    filename: str,
    content_type: str,
    content: bytes,
) -> tuple[UUID, UUID, str]:
    result = await UploadService(
        session_factory,
        storage,
        max_upload_size_bytes=50 * 1024 * 1024,
    ).upload_and_confirm(
        workspace_id=workspace_id,
        document_id=document_id,
        idempotency_key=str(uuid4()),
        original_filename=filename,
        content_type=content_type,
        stream=io.BytesIO(content),
    )
    async with session_factory() as session:
        job = await session.get(ProcessingJob, result.processing_job_id)
        assert job is not None
        version = await session.get(DocumentVersion, result.document_version_id)
        assert version is not None
        return version.id, job.processing_run_id, version.object_storage_key


async def _claim_run(
    session_factory: async_sessionmaker[AsyncSession],
    run_id: UUID,
    *,
    worker_id: str = "canonical-worker",
) -> ClaimedProcessingJob:
    claim = await ProcessingCoordinator(session_factory).claim_next(worker_id=worker_id)
    assert claim is not None and claim.processing_run_id == run_id
    return claim


async def _ingest_claimed(service, session_factory, run_id: UUID):
    claim = await _claim_run(session_factory, run_id)
    return await service.ingest(run_id, worker_id="canonical-worker", lease_token=claim.lease_token)


async def _add_queued_job(session: AsyncSession, run_id: UUID) -> None:
    await session.flush()
    session.add(ProcessingJob(processing_run_id=run_id, max_attempts=3))


async def _legacy_completed_job(factory, claim) -> None:
    """Fixture for historical terminal rows, not a production completion entrypoint."""
    async with factory() as session, session.begin():
        await session.execute(
            text(
                "UPDATE processing_jobs SET status = 'completed', worker_id = NULL, "
                "lease_token = NULL, lease_expires_at = NULL, next_retry_at = NULL, "
                "failure_reason = NULL, finished_at = clock_timestamp() WHERE id = :id"
            ),
            {"id": claim.job_id},
        )
        await session.execute(
            text("UPDATE processing_runs SET status = 'completed' WHERE id = :id"),
            {"id": claim.processing_run_id},
        )


@pytest.mark.integration
def test_successful_pdf_text_is_not_published_as_complete_evidence(
    migrated_database_url: str,
) -> None:
    async def exercise() -> None:
        engine = create_async_engine(migrated_database_url)
        factory = async_sessionmaker(engine, expire_on_commit=False)
        storage = MemoryStorage()
        try:
            workspace, document = await _create_document(factory)
            _, run_id, _ = await _upload(
                factory,
                storage,
                workspace_id=workspace,
                document_id=document,
                filename="unverified-table.pdf",
                content_type=PDF_CONTENT_TYPE,
                content=_pdf_bytes(["Amount kg 2026 entity A -10. Native text is recoverable."]),
            )
            with pytest.raises(IngestionFailure) as failure:
                await _ingest_claimed(IngestionService(factory, storage), factory, run_id)
            assert failure.value.category is IngestionFailureCategory.EXTRACTION_NEEDS_REVIEW
            assert not failure.value.retryable
            async with factory() as session:
                assert await session.scalar(select(func.count(CanonicalChunk.id))) == 0
        finally:
            await engine.dispose()

    asyncio.run(exercise())


@pytest.mark.integration
def test_docx_xlsx_persistence_idempotency_and_reprocessing_isolation(
    migrated_database_url: str,
) -> None:
    async def exercise() -> None:
        engine = create_async_engine(migrated_database_url)
        session_factory = async_sessionmaker(engine, expire_on_commit=False)
        storage = MemoryStorage()
        workspace_id, document_id = await _create_document(session_factory)
        sources = (
            (
                "source-text.docx",
                DOCX_CONTENT_TYPE,
                _docx_bytes(text="First version of the source with extractable text."),
            ),
            ("source.docx", DOCX_CONTENT_TYPE, _docx_bytes()),
            ("source.xlsx", XLSX_CONTENT_TYPE, _xlsx_bytes()),
        )
        uploaded = [
            await _upload(
                session_factory,
                storage,
                workspace_id=workspace_id,
                document_id=document_id,
                filename=filename,
                content_type=content_type,
                content=content,
            )
            for filename, content_type, content in sources
        ]
        service = IngestionService(session_factory, storage)
        try:
            claim = await _claim_run(session_factory, uploaded[0][1])
            competing = await asyncio.gather(
                service.ingest(
                    uploaded[0][1], worker_id="canonical-worker", lease_token=claim.lease_token
                ),
                service.ingest(
                    uploaded[0][1], worker_id="canonical-worker", lease_token=claim.lease_token
                ),
            )
            assert {item.status for item in competing} == {
                IngestionResultStatus.PERSISTED,
                IngestionResultStatus.ALREADY_PERSISTED,
            }
            for _, run_id, _ in uploaded[1:]:
                assert (
                    await _ingest_claimed(service, session_factory, run_id)
                ).status is IngestionResultStatus.PERSISTED

            async with session_factory() as session:
                chunks_by_run: dict[UUID, tuple[CanonicalChunk, ...]] = {}
                for version_id, run_id, _ in uploaded:
                    chunks = tuple(
                        (
                            await session.scalars(
                                select(CanonicalChunk)
                                .where(CanonicalChunk.processing_run_id == run_id)
                                .order_by(CanonicalChunk.ordinal)
                            )
                        ).all()
                    )
                    assert chunks
                    assert [chunk.ordinal for chunk in chunks] == list(range(len(chunks)))
                    assert all(
                        chunk.id == canonical_chunk_uuid(run_id, chunk.chunk_identity)
                        and chunk.content_sha256
                        == hashlib.sha256(chunk.chunk_text.encode()).hexdigest()
                        and chunk.character_count == len(chunk.chunk_text)
                        and chunk.token_count is None
                        for chunk in chunks
                    )
                    chunks_by_run[run_id] = chunks
                    resolved_workspace = await session.scalar(
                        select(Document.workspace_id)
                        .join(DocumentVersion, DocumentVersion.document_id == Document.id)
                        .join(
                            ProcessingRun,
                            ProcessingRun.document_version_id == DocumentVersion.id,
                        )
                        .where(ProcessingRun.id == run_id)
                    )
                    assert resolved_workspace == workspace_id
                    assert (
                        version_id == (await session.get(ProcessingRun, run_id)).document_version_id
                    )

                first_chunks = chunks_by_run[uploaded[0][1]]
                assert {
                    span["source"]["part_uri"]
                    for chunk in first_chunks
                    for span in chunk.source_provenance
                } == {"/word/document.xml"}
                docx_chunks = chunks_by_run[uploaded[1][1]]
                assert any(
                    span["source"].get("table_index") == 0
                    for chunk in docx_chunks
                    for span in chunk.source_provenance
                )
                xlsx_chunks = chunks_by_run[uploaded[2][1]]
                assert all(
                    span["source"]["sheet_index"] == 0
                    for chunk in xlsx_chunks
                    for span in chunk.source_provenance
                )
                assert any(
                    span["source"].get("cell_coordinate") == "B2"
                    for chunk in xlsx_chunks
                    for span in chunk.source_provenance
                )

                original_run = await session.get(ProcessingRun, uploaded[0][1])
                original_job = await session.scalar(
                    select(ProcessingJob).where(ProcessingJob.processing_run_id == uploaded[0][1])
                )
                assert original_run is not None
                assert original_job is not None
                assert original_run.status == ProcessingStatus.PROCESSING
                assert original_run.stage == ProcessingStage.EXTRACTION
                assert original_job.status == ProcessingJobStatus.LEASED

            reprocess_run_id = uuid4()
            async with session_factory() as session, session.begin():
                session.add(
                    ProcessingRun(
                        id=reprocess_run_id,
                        document_version_id=uploaded[0][0],
                        operation_id=uuid4(),
                        operation_kind=ProcessingOperationKind.REPROCESS,
                        configuration_version=PROCESSING_CONFIGURATION_VERSION,
                        configuration_fingerprint=default_processing_configuration_fingerprint(),
                        stage=ProcessingStage.EXTRACTION,
                        status=ProcessingStatus.QUEUED,
                    )
                )
                await _add_queued_job(session, reprocess_run_id)
            assert (
                await _ingest_claimed(service, session_factory, reprocess_run_id)
            ).status is IngestionResultStatus.PERSISTED
            async with session_factory() as session:
                original = tuple(
                    (
                        await session.scalars(
                            select(CanonicalChunk)
                            .where(CanonicalChunk.processing_run_id == uploaded[0][1])
                            .order_by(CanonicalChunk.ordinal)
                        )
                    ).all()
                )
                reprocessed = tuple(
                    (
                        await session.scalars(
                            select(CanonicalChunk)
                            .where(CanonicalChunk.processing_run_id == reprocess_run_id)
                            .order_by(CanonicalChunk.ordinal)
                        )
                    ).all()
                )
                assert [item.chunk_identity for item in original] == [
                    item.chunk_identity for item in reprocessed
                ]
                assert {item.id for item in original}.isdisjoint({item.id for item in reprocessed})
        finally:
            await engine.dispose()

    asyncio.run(exercise())


@pytest.mark.integration
def test_stale_worker_cannot_publish_after_lease_reclaimed_during_extraction(
    migrated_database_url: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def exercise() -> None:
        engine = create_async_engine(migrated_database_url)
        session_factory = async_sessionmaker(engine, expire_on_commit=False)
        storage = MemoryStorage()
        workspace_id, document_id = await _create_document(session_factory)
        _, run_id, _ = await _upload(
            session_factory,
            storage,
            workspace_id=workspace_id,
            document_id=document_id,
            filename="fenced.docx",
            content_type=DOCX_CONTENT_TYPE,
            content=_docx_bytes(text="Lease and atomic publication regression source."),
        )
        coordinator = ProcessingCoordinator(session_factory)
        first = await coordinator.claim_next(worker_id="worker-a")
        assert first is not None and first.processing_run_id == run_id
        service = IngestionService(session_factory, storage)
        extracted = asyncio.Event()
        release = asyncio.Event()
        original = service._extract_and_chunk
        captured = []

        async def paused(source, *, suffix):
            chunks = await original(source, suffix=suffix)
            captured.append(
                (
                    source,
                    tuple(ingestion_module.canonicalize_chunk(run_id, chunk) for chunk in chunks),
                )
            )
            extracted.set()
            await release.wait()
            return chunks

        monkeypatch.setattr(service, "_extract_and_chunk", paused)
        task = asyncio.create_task(
            service.ingest(run_id, worker_id="worker-a", lease_token=first.lease_token)
        )
        try:
            await asyncio.wait_for(extracted.wait(), 10)
            async with session_factory() as session, session.begin():
                await session.execute(
                    text(
                        "UPDATE processing_jobs SET lease_expires_at = clock_timestamp() "
                        "- INTERVAL '1 second' WHERE id = :id"
                    ),
                    {"id": first.job_id},
                )
            second = await asyncio.wait_for(coordinator.claim_next(worker_id="worker-b"), 10)
            assert second is not None and second.job_id == first.job_id
            assert second.lease_token != first.lease_token
            release.set()
            with pytest.raises(IngestionFailure, match="lease_ownership_lost"):
                await asyncio.wait_for(task, 10)
            async with session_factory() as session:
                assert await session.scalar(select(func.count(CanonicalChunk.id))) == 0
                job = await session.get(ProcessingJob, first.job_id)
                assert job.worker_id == "worker-b" and job.lease_token == second.lease_token
            result = await service.ingest(
                run_id, worker_id="worker-b", lease_token=second.lease_token
            )
            assert result.status is IngestionResultStatus.PERSISTED
            assert result.chunk_count > 0
            assert (
                await service.ingest(run_id, worker_id="worker-b", lease_token=second.lease_token)
            ).status is IngestionResultStatus.ALREADY_PERSISTED
            with pytest.raises(IngestionFailure, match="lease_ownership_lost"):
                await service.ingest(run_id, worker_id="worker-a", lease_token=first.lease_token)
            # Even equivalent already-published chunks cannot authorize the old lease.
            source, chunks = captured[0]
            with pytest.raises(IngestionFailure, match="lease_ownership_lost"):
                await service._persist(
                    source, chunks, worker_id="worker-a", lease_token=first.lease_token
                )
            async with session_factory() as session:
                job = await session.get(ProcessingJob, second.job_id)
                run = await session.get(ProcessingRun, run_id)
                assert job.status == ProcessingJobStatus.LEASED
                assert job.worker_id == "worker-b" and job.lease_token == second.lease_token
                assert run.status == ProcessingStatus.PROCESSING
                assert run.stage == ProcessingStage.EXTRACTION
                assert (
                    await session.scalar(select(func.count(CanonicalChunk.id)))
                    == result.chunk_count
                )
        finally:
            release.set()
            if not task.done():
                task.cancel()
            await asyncio.gather(task, return_exceptions=True)
            await engine.dispose()

    asyncio.run(exercise())


@pytest.mark.integration
@pytest.mark.parametrize(
    "case",
    [
        "queued",
        "missing_job",
        "wrong_worker",
        "wrong_token",
        "blank_worker",
        "expired",
        "retry_wait",
        "completed",
        "failed",
        "run_stage",
        "run_status",
        "missing_run",
    ],
)
def test_ownership_preflight_prevents_storage_and_extraction(
    migrated_database_url: str,
    monkeypatch: pytest.MonkeyPatch,
    case: str,
) -> None:
    async def exercise() -> None:
        engine = create_async_engine(migrated_database_url)
        factory = async_sessionmaker(engine, expire_on_commit=False)
        storage = MemoryStorage()
        try:
            workspace, document = await _create_document(factory)
            version_id, run_id, _ = await _upload(
                factory,
                storage,
                workspace_id=workspace,
                document_id=document,
                filename="preflight.pdf",
                content_type=PDF_CONTENT_TYPE,
                content=_pdf_bytes(["Preflight source content must never reach the extractor."]),
            )
            coordinator = ProcessingCoordinator(factory)
            worker, token = "canonical-worker", uuid4()
            if case not in {"queued", "missing_job", "missing_run"}:
                claim = await _claim_run(factory, run_id)
                token = claim.lease_token
                if case == "wrong_worker":
                    worker = "other-worker"
                elif case == "wrong_token":
                    token = uuid4()
                elif case == "blank_worker":
                    worker = " "
                elif case == "expired":
                    async with factory() as session, session.begin():
                        await session.execute(
                            text(
                                "UPDATE processing_jobs SET lease_expires_at = clock_timestamp() "
                                "- INTERVAL '1 second' WHERE id = :id"
                            ),
                            {"id": claim.job_id},
                        )
                elif case in {"retry_wait", "failed"}:
                    if case == "failed":
                        async with factory() as session, session.begin():
                            job = await session.get(ProcessingJob, claim.job_id)
                            job.max_attempts = 1
                    await coordinator.fail_attempt(
                        job_id=claim.job_id,
                        worker_id=worker,
                        lease_token=token,
                        failure_reason="test_failure",
                    )
                elif case in {"completed", "run_stage"}:
                    stages = (
                        tuple(ProcessingStage)[1:]
                        if case == "completed"
                        else (ProcessingStage.CLEANING,)
                    )
                    for stage in stages:
                        await coordinator.advance_stage(
                            job_id=claim.job_id, worker_id=worker, lease_token=token, stage=stage
                        )
                    if case == "completed":
                        with pytest.raises(InvalidProcessingTransition, match="verified index"):
                            await coordinator.complete(
                                job_id=claim.job_id, worker_id=worker, lease_token=token
                            )
                        await _legacy_completed_job(factory, claim)
                elif case == "run_status":
                    async with factory() as session, session.begin():
                        run = await session.get(ProcessingRun, run_id)
                        run.status = ProcessingStatus.FAILED
            if case == "missing_job":
                async with factory() as session, session.begin():
                    await session.execute(
                        text("DELETE FROM processing_jobs WHERE processing_run_id = :id"),
                        {"id": run_id},
                    )
            service = IngestionService(factory, storage)

            async def forbidden(*args, **kwargs):
                pytest.fail("Invalid ownership reached storage or extraction")

            monkeypatch.setattr(storage, "inspect", forbidden)
            monkeypatch.setattr(service, "_extract_and_chunk", forbidden)
            async with factory() as session:
                before = (
                    await session.execute(
                        text(
                            "SELECT worker_id, lease_token, status, attempt_count, "
                            "lease_expires_at, "
                            "failure_reason FROM processing_jobs WHERE processing_run_id = :id"
                        ),
                        {"id": run_id},
                    )
                ).first()
                original_version = (
                    await session.execute(
                        text("SELECT * FROM document_versions WHERE id = :id"), {"id": version_id}
                    )
                ).one()
            with pytest.raises(TypeError):
                await service.ingest(run_id)  # type: ignore[call-arg]
            with pytest.raises(IngestionFailure) as failure:
                await service.ingest(
                    uuid4() if case == "missing_run" else run_id,
                    worker_id=worker,
                    lease_token=token,
                )
            expected = IngestionFailureCategory.LEASE_OWNERSHIP_LOST
            if case in {"run_stage", "run_status"}:
                expected = IngestionFailureCategory.PROCESSING_STATE_MISMATCH
            elif case == "missing_run":
                expected = IngestionFailureCategory.PROCESSING_RUN_NOT_FOUND
            assert failure.value.category is expected and not failure.value.retryable
            assert str(failure.value) == expected.value
            async with factory() as session:
                after = (
                    await session.execute(
                        text(
                            "SELECT worker_id, lease_token, status, attempt_count, "
                            "lease_expires_at, "
                            "failure_reason FROM processing_jobs WHERE processing_run_id = :id"
                        ),
                        {"id": run_id},
                    )
                ).first()
                assert after == before
                assert (
                    await session.execute(
                        text("SELECT * FROM document_versions WHERE id = :id"), {"id": version_id}
                    )
                ).one() == original_version
                assert await session.scalar(select(func.count(CanonicalChunk.id))) == 0
        finally:
            await engine.dispose()

    asyncio.run(exercise())


@pytest.mark.integration
@pytest.mark.parametrize("transition", ["retry_wait", "failed", "completed", "stage"])
def test_revocation_or_transition_during_extraction_rolls_back(
    migrated_database_url: str,
    monkeypatch: pytest.MonkeyPatch,
    transition: str,
) -> None:
    async def exercise() -> None:
        engine = create_async_engine(migrated_database_url)
        factory = async_sessionmaker(engine, expire_on_commit=False)
        storage = MemoryStorage()
        coordinator = ProcessingCoordinator(factory)
        extracted, release = asyncio.Event(), asyncio.Event()
        task = None
        try:
            workspace, document = await _create_document(factory)
            _, run_id, _ = await _upload(
                factory,
                storage,
                workspace_id=workspace,
                document_id=document,
                filename="transition.docx",
                content_type=DOCX_CONTENT_TYPE,
                content=_docx_bytes(text="Lease and atomic publication regression source."),
            )
            claim = await _claim_run(factory, run_id)
            service = IngestionService(factory, storage)
            original = service._extract_and_chunk

            async def paused(source, *, suffix):
                chunks = await original(source, suffix=suffix)
                extracted.set()
                await release.wait()
                return chunks

            monkeypatch.setattr(service, "_extract_and_chunk", paused)
            task = asyncio.create_task(
                service.ingest(run_id, worker_id="canonical-worker", lease_token=claim.lease_token)
            )
            await asyncio.wait_for(extracted.wait(), 10)
            if transition in {"failed", "retry_wait"}:
                if transition == "failed":
                    async with factory() as session, session.begin():
                        job = await session.get(ProcessingJob, claim.job_id)
                        job.max_attempts = 1
                await asyncio.wait_for(
                    coordinator.fail_attempt(
                        job_id=claim.job_id,
                        worker_id="canonical-worker",
                        lease_token=claim.lease_token,
                        failure_reason="test_failure",
                    ),
                    10,
                )
            else:
                stages = (
                    tuple(ProcessingStage)[1:]
                    if transition == "completed"
                    else (ProcessingStage.CLEANING,)
                )
                for stage in stages:
                    await asyncio.wait_for(
                        coordinator.advance_stage(
                            job_id=claim.job_id,
                            worker_id="canonical-worker",
                            lease_token=claim.lease_token,
                            stage=stage,
                        ),
                        10,
                    )
                if transition == "completed":
                    with pytest.raises(InvalidProcessingTransition, match="verified index"):
                        await asyncio.wait_for(
                            coordinator.complete(
                                job_id=claim.job_id,
                                worker_id="canonical-worker",
                                lease_token=claim.lease_token,
                            ),
                            10,
                        )
                    await _legacy_completed_job(factory, claim)
            # Coordinator transitions finish while extraction is paused: no Job/Run locks held.
            release.set()
            expected = (
                "processing_state_mismatch" if transition == "stage" else "lease_ownership_lost"
            )
            with pytest.raises(IngestionFailure, match=expected) as failure:
                await asyncio.wait_for(task, 10)
            assert not failure.value.retryable
            async with factory() as session:
                assert await session.scalar(select(func.count(CanonicalChunk.id))) == 0
                job = await session.get(ProcessingJob, claim.job_id)
                expected_status = "leased" if transition == "stage" else transition
                assert job.status == expected_status
        finally:
            release.set()
            if task is not None:
                if not task.done():
                    task.cancel()
                await asyncio.gather(task, return_exceptions=True)
            await engine.dispose()

    asyncio.run(exercise())


@pytest.mark.integration
def test_lease_expiration_after_row_lock_wait_uses_postgres_wall_clock(
    migrated_database_url: str,
) -> None:
    async def exercise() -> None:
        engine = create_async_engine(migrated_database_url)
        factory = async_sessionmaker(engine, expire_on_commit=False)
        storage = MemoryStorage()
        task = None
        try:
            workspace, document = await _create_document(factory)
            _, run_id, _ = await _upload(
                factory,
                storage,
                workspace_id=workspace,
                document_id=document,
                filename="lock-wait.docx",
                content_type=DOCX_CONTENT_TYPE,
                content=_docx_bytes(text="Lease and atomic publication regression source."),
            )
            claim = await _claim_run(factory, run_id)
            service = IngestionService(factory, storage)
            source = await service._resolve_source(
                run_id, worker_id="canonical-worker", lease_token=claim.lease_token
            )
            chunks = tuple(
                ingestion_module.canonicalize_chunk(run_id, chunk)
                for chunk in await service._extract_and_chunk(source, suffix=".docx")
            )
            async with factory() as blocker, blocker.begin():
                await blocker.scalar(
                    select(ProcessingRun).where(ProcessingRun.id == run_id).with_for_update()
                )
                blocker_pid = await blocker.scalar(select(func.pg_backend_pid()))
                async with factory() as session, session.begin():
                    deadline = await session.scalar(
                        text(
                            "UPDATE processing_jobs SET lease_expires_at = clock_timestamp() "
                            "+ INTERVAL '2 seconds' WHERE id = :id RETURNING lease_expires_at"
                        ),
                        {"id": claim.job_id},
                    )
                task = asyncio.create_task(
                    service._persist(
                        source, chunks, worker_id="canonical-worker", lease_token=claim.lease_token
                    )
                )
                # Synchronize on an actual PostgreSQL lock waiter, not a guessed sleep.
                async with asyncio.timeout(10), factory() as monitor:
                    while not await monitor.scalar(
                        text(
                            "SELECT EXISTS(SELECT 1 FROM pg_locks "
                            "WHERE NOT granted AND :blocker = ANY(pg_blocking_pids(pid)))"
                        ),
                        {"blocker": blocker_pid},
                    ):
                        assert not task.done()
                        await asyncio.sleep(0.01)
                    await monitor.scalar(select(func.pg_sleep_until(deadline)))
                    assert await monitor.scalar(select(func.clock_timestamp())) >= deadline
                # A owns Job before waiting on Run, so reclaim skips that locked job.
                assert (
                    await asyncio.wait_for(
                        ProcessingCoordinator(factory).claim_next(worker_id="worker-b"), 10
                    )
                    is None
                )
            with pytest.raises(IngestionFailure, match="lease_ownership_lost"):
                await asyncio.wait_for(task, 10)
            async with factory() as session:
                assert await session.scalar(select(func.count(CanonicalChunk.id))) == 0
            second = await _claim_run(factory, run_id, worker_id="worker-b")
            assert second.lease_token != claim.lease_token
            assert (
                await service.ingest(run_id, worker_id="worker-b", lease_token=second.lease_token)
            ).status is IngestionResultStatus.PERSISTED
        finally:
            if task is not None:
                if not task.done():
                    task.cancel()
                await asyncio.gather(task, return_exceptions=True)
            await engine.dispose()

    asyncio.run(exercise())


@pytest.mark.integration
@pytest.mark.parametrize("already_persisted", [False, True])
def test_expiration_at_final_check_rolls_back_or_rejects_idempotence(
    migrated_database_url: str,
    monkeypatch: pytest.MonkeyPatch,
    already_persisted: bool,
) -> None:
    async def exercise() -> None:
        engine = create_async_engine(migrated_database_url)
        factory = async_sessionmaker(engine, expire_on_commit=False)
        storage = MemoryStorage()
        try:
            workspace, document = await _create_document(factory)
            version_id, run_id, _ = await _upload(
                factory,
                storage,
                workspace_id=workspace,
                document_id=document,
                filename="expiry.docx",
                content_type=DOCX_CONTENT_TYPE,
                content=_docx_bytes(text="Lease and atomic publication regression source."),
            )
            claim = await _claim_run(factory, run_id)
            service = IngestionService(factory, storage)
            source = await service._resolve_source(
                run_id, worker_id="canonical-worker", lease_token=claim.lease_token
            )
            chunks = tuple(
                ingestion_module.canonicalize_chunk(run_id, chunk)
                for chunk in await service._extract_and_chunk(source, suffix=".docx")
            )
            assert len(chunks) == 2
            if already_persisted:
                await service._persist(
                    source, chunks, worker_id="canonical-worker", lease_token=claim.lease_token
                )
            async with factory() as session, session.begin():
                await session.execute(
                    text(
                        "UPDATE processing_jobs SET lease_expires_at = clock_timestamp() "
                        "+ INTERVAL '1 second' WHERE id = :id"
                    ),
                    {"id": claim.job_id},
                )
                original_version = (
                    await session.execute(
                        text("SELECT * FROM document_versions WHERE id = :id"), {"id": version_id}
                    )
                ).one()
            original = ingestion_module._require_ingestion_ownership
            calls = 0

            async def expire_before_final_check(session, job, run, **ownership):
                nonlocal calls
                calls += 1
                if calls == 2:
                    # INSERTs have really flushed, but are not yet committed.
                    assert await session.scalar(select(func.count(CanonicalChunk.id))) == len(
                        chunks
                    )
                    await session.scalar(select(func.pg_sleep_until(job.lease_expires_at)))
                await original(session, job, run, **ownership)

            monkeypatch.setattr(
                ingestion_module, "_require_ingestion_ownership", expire_before_final_check
            )
            with pytest.raises(IngestionFailure, match="lease_ownership_lost"):
                await asyncio.wait_for(
                    service._persist(
                        source, chunks, worker_id="canonical-worker", lease_token=claim.lease_token
                    ),
                    10,
                )
            assert calls == 2
            async with factory() as session:
                assert await session.scalar(select(func.count(CanonicalChunk.id))) == (
                    len(chunks) if already_persisted else 0
                )
                assert (
                    await session.execute(
                        text("SELECT * FROM document_versions WHERE id = :id"), {"id": version_id}
                    )
                ).one() == original_version
                job = await session.get(ProcessingJob, claim.job_id)
                assert (
                    job.status == ProcessingJobStatus.LEASED
                    and job.lease_token == claim.lease_token
                )
        finally:
            await engine.dispose()

    asyncio.run(exercise())


@pytest.mark.integration
def test_independent_runs_isolate_lease_failure(migrated_database_url: str) -> None:
    async def exercise() -> None:
        engine = create_async_engine(migrated_database_url)
        factory = async_sessionmaker(engine, expire_on_commit=False)
        storage = MemoryStorage()
        try:
            workspace, document = await _create_document(factory)
            runs = [
                (
                    await _upload(
                        factory,
                        storage,
                        workspace_id=workspace,
                        document_id=document,
                        filename=f"independent-{index}.docx",
                        content_type=DOCX_CONTENT_TYPE,
                        content=_docx_bytes(text="Lease and atomic publication regression source."),
                    )
                )[1]
                for index in range(2)
            ]
            claims = [
                await _claim_run(factory, run, worker_id=f"worker-{index}")
                for index, run in enumerate(runs)
            ]
            service = IngestionService(factory, storage)
            results = await asyncio.gather(
                service.ingest(
                    runs[0], worker_id="wrong-worker", lease_token=claims[0].lease_token
                ),
                service.ingest(runs[1], worker_id="worker-1", lease_token=claims[1].lease_token),
                return_exceptions=True,
            )
            assert isinstance(results[0], IngestionFailure)
            assert results[0].category is IngestionFailureCategory.LEASE_OWNERSHIP_LOST
            assert results[1].status is IngestionResultStatus.PERSISTED
            async with factory() as session:
                assert set(await session.scalars(select(CanonicalChunk.processing_run_id))) == {
                    runs[1]
                }
                for index, claim in enumerate(claims):
                    job = await session.get(ProcessingJob, claim.job_id)
                    assert (
                        job.worker_id == f"worker-{index}" and job.lease_token == claim.lease_token
                    )
        finally:
            await engine.dispose()

    asyncio.run(exercise())


@pytest.mark.integration
def test_ingestion_failure_gates_publish_no_chunks(migrated_database_url: str) -> None:
    async def exercise() -> None:
        engine = create_async_engine(migrated_database_url)
        session_factory = async_sessionmaker(engine, expire_on_commit=False)
        storage = MemoryStorage()
        workspace_id, document_id = await _create_document(session_factory)
        try:
            _, missing_run, missing_key = await _upload(
                session_factory,
                storage,
                workspace_id=workspace_id,
                document_id=document_id,
                filename="missing.pdf",
                content_type=PDF_CONTENT_TYPE,
                content=_pdf_bytes(["A complete source document for the missing-object test."]),
            )
            await storage.delete(missing_key)
            with pytest.raises(IngestionFailure) as missing:
                await _ingest_claimed(
                    IngestionService(session_factory, storage), session_factory, missing_run
                )
            assert (missing.value.category, missing.value.retryable) == (
                IngestionFailureCategory.SOURCE_OBJECT_MISSING,
                False,
            )

            _, mismatch_run, mismatch_key = await _upload(
                session_factory,
                storage,
                workspace_id=workspace_id,
                document_id=document_id,
                filename="mismatch.pdf",
                content_type=PDF_CONTENT_TYPE,
                content=_pdf_bytes(["Original source identity with sufficient extracted text."]),
            )
            stored = storage.objects[mismatch_key]
            storage.objects[mismatch_key] = replace(
                stored,
                content=stored.content[:-1] + bytes([stored.content[-1] ^ 1]),
            )
            with pytest.raises(IngestionFailure) as mismatch:
                await _ingest_claimed(
                    IngestionService(session_factory, storage), session_factory, mismatch_run
                )
            assert mismatch.value.category is IngestionFailureCategory.SOURCE_IDENTITY_MISMATCH

            _, unavailable_run, _ = await _upload(
                session_factory,
                storage,
                workspace_id=workspace_id,
                document_id=document_id,
                filename="unavailable.pdf",
                content_type=PDF_CONTENT_TYPE,
                content=_pdf_bytes(["Infrastructure failure source with sufficient text."]),
            )
            storage.available = False
            with pytest.raises(IngestionFailure) as unavailable:
                await _ingest_claimed(
                    IngestionService(session_factory, storage), session_factory, unavailable_run
                )
            assert (unavailable.value.category, unavailable.value.retryable) == (
                IngestionFailureCategory.STORAGE_UNAVAILABLE,
                True,
            )
            storage.available = True

            _, corrupt_run, _ = await _upload(
                session_factory,
                storage,
                workspace_id=workspace_id,
                document_id=document_id,
                filename="corrupt.pdf",
                content_type=PDF_CONTENT_TYPE,
                content=b"%PDF-1.7\nnot a valid pdf\n%%EOF\n",
            )
            with pytest.raises(IngestionFailure) as corrupt:
                await _ingest_claimed(
                    IngestionService(session_factory, storage), session_factory, corrupt_run
                )
            assert corrupt.value.category is IngestionFailureCategory.EXTRACTION_FAILED
            assert corrupt.value.document_error_category == "corrupt_document"

            _, empty_run, _ = await _upload(
                session_factory,
                storage,
                workspace_id=workspace_id,
                document_id=document_id,
                filename="empty.docx",
                content_type=DOCX_CONTENT_TYPE,
                content=_docx_bytes(empty=True),
            )
            with pytest.raises(IngestionFailure) as empty:
                await _ingest_claimed(
                    IngestionService(session_factory, storage), session_factory, empty_run
                )
            assert empty.value.category is IngestionFailureCategory.EMPTY_CONTENT

            _, review_run, _ = await _upload(
                session_factory,
                storage,
                workspace_id=workspace_id,
                document_id=document_id,
                filename="review.pdf",
                content_type=PDF_CONTENT_TYPE,
                content=_pdf_bytes(["short"]),
            )
            with pytest.raises(IngestionFailure) as review:
                await _ingest_claimed(
                    IngestionService(session_factory, storage), session_factory, review_run
                )
            assert review.value.category is IngestionFailureCategory.EXTRACTION_NEEDS_REVIEW

            async with session_factory() as session:
                assert await session.scalar(select(func.count(CanonicalChunk.id))) == 0
        finally:
            await engine.dispose()

    asyncio.run(exercise())


@pytest.mark.integration
def test_canonical_chunk_set_rolls_back_on_persistence_failure(
    migrated_database_url: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def exercise() -> None:
        engine = create_async_engine(migrated_database_url)
        session_factory = async_sessionmaker(engine, expire_on_commit=False)
        storage = MemoryStorage()
        workspace_id, document_id = await _create_document(session_factory)
        _, run_id, _ = await _upload(
            session_factory,
            storage,
            workspace_id=workspace_id,
            document_id=document_id,
            filename="rollback.docx",
            content_type=DOCX_CONTENT_TYPE,
            content=_docx_bytes(text="Lease and atomic publication regression source."),
        )
        original = ingestion_module.canonicalize_chunk

        def invalid_second(processing_run_id: UUID, chunk: object):
            payload = original(processing_run_id, chunk)  # type: ignore[arg-type]
            return replace(payload, character_count=0) if payload.ordinal == 1 else payload

        monkeypatch.setattr(ingestion_module, "canonicalize_chunk", invalid_second)
        try:
            with pytest.raises(IngestionFailure) as failure:
                await _ingest_claimed(
                    IngestionService(session_factory, storage), session_factory, run_id
                )
            assert failure.value.category is IngestionFailureCategory.PERSISTENCE_CONFLICT
            async with session_factory() as session:
                assert await session.scalar(select(func.count(CanonicalChunk.id))) == 0
        finally:
            await engine.dispose()

    asyncio.run(exercise())


@pytest.mark.integration
def test_configuration_mismatch_and_existing_chunk_conflict_fail_safely(
    migrated_database_url: str,
) -> None:
    async def exercise() -> None:
        engine = create_async_engine(migrated_database_url)
        session_factory = async_sessionmaker(engine, expire_on_commit=False)
        storage = MemoryStorage()
        workspace_id, document_id = await _create_document(session_factory)
        version_id, run_id, _ = await _upload(
            session_factory,
            storage,
            workspace_id=workspace_id,
            document_id=document_id,
            filename="conflict.docx",
            content_type=DOCX_CONTENT_TYPE,
            content=_docx_bytes(text="Lease and atomic publication regression source."),
        )
        conflicting_text = "pre-existing conflicting canonical text"
        conflicting_id = f"chunk-{'d' * 64}"
        async with session_factory() as session, session.begin():
            session.add(
                CanonicalChunk(
                    id=canonical_chunk_uuid(run_id, conflicting_id),
                    processing_run_id=run_id,
                    ordinal=0,
                    chunk_identity=conflicting_id,
                    chunk_text=conflicting_text,
                    content_sha256=hashlib.sha256(conflicting_text.encode()).hexdigest(),
                    character_count=len(conflicting_text),
                    token_count=None,
                    document_format="pdf",
                    chunking_strategy="structure_aware",
                    chunker_version="1.0",
                    chunker_configuration_fingerprint="e" * 64,
                    source_provenance=[
                        {
                            "unit_id": "page:1:block:0",
                            "source": {"page_number": 1, "block_index": 0},
                            "source_start": 0,
                            "source_end": len(conflicting_text),
                            "source_length": len(conflicting_text),
                        }
                    ],
                    split_oversized_element=False,
                )
            )

        service = IngestionService(session_factory, storage)
        try:
            with pytest.raises(IngestionFailure) as conflict:
                await _ingest_claimed(service, session_factory, run_id)
            assert conflict.value.category is IngestionFailureCategory.PERSISTENCE_CONFLICT
            async with session_factory() as session:
                persisted = tuple(
                    (
                        await session.scalars(
                            select(CanonicalChunk).where(CanonicalChunk.processing_run_id == run_id)
                        )
                    ).all()
                )
                assert len(persisted) == 1
                assert persisted[0].chunk_text == conflicting_text

            mismatched_run_id = uuid4()
            async with session_factory() as session, session.begin():
                session.add(
                    ProcessingRun(
                        id=mismatched_run_id,
                        document_version_id=version_id,
                        operation_id=uuid4(),
                        operation_kind=ProcessingOperationKind.REPROCESS,
                        configuration_version=PROCESSING_CONFIGURATION_VERSION,
                        configuration_fingerprint="f" * 64,
                        stage=ProcessingStage.EXTRACTION,
                        status=ProcessingStatus.QUEUED,
                    )
                )
                await _add_queued_job(session, mismatched_run_id)
            with pytest.raises(IngestionFailure) as mismatch:
                await _ingest_claimed(service, session_factory, mismatched_run_id)
            assert mismatch.value.category is IngestionFailureCategory.CONFIGURATION_MISMATCH
            async with session_factory() as session:
                assert (
                    await session.scalar(
                        select(func.count(CanonicalChunk.id)).where(
                            CanonicalChunk.processing_run_id == mismatched_run_id
                        )
                    )
                    == 0
                )
        finally:
            await engine.dispose()

    asyncio.run(exercise())


@pytest.mark.integration
def test_canonical_chunk_migration_upgrade_check_and_downgrade(
    disposable_database_url: str,
    alembic_runner: Callable[..., None],
) -> None:
    alembic_runner(disposable_database_url, "upgrade", "0005_processing_runs_jobs")
    alembic_runner(disposable_database_url, "upgrade", "head")
    alembic_runner(disposable_database_url, "check")

    async def table_exists() -> bool:
        engine = create_async_engine(disposable_database_url)
        try:
            async with engine.connect() as connection:
                return bool(
                    await connection.scalar(
                        text("SELECT to_regclass('public.canonical_chunks') IS NOT NULL")
                    )
                )
        finally:
            await engine.dispose()

    assert asyncio.run(table_exists()) is True
    alembic_runner(disposable_database_url, "downgrade", "0005_processing_runs_jobs")
    assert asyncio.run(table_exists()) is False


@pytest.mark.integration
@pytest.mark.storage
def test_real_minio_ingestion_persists_canonical_chunks(
    migrated_database_url: str,
    test_storage_settings: StorageSettings,
) -> None:
    async def exercise() -> None:
        engine = create_async_engine(migrated_database_url)
        session_factory = async_sessionmaker(engine, expire_on_commit=False)
        storage = S3ObjectStorage(test_storage_settings)
        await storage.probe()
        workspace_id, document_id = await _create_document(session_factory)
        _, run_id, object_key = await _upload(
            session_factory,
            storage,
            workspace_id=workspace_id,
            document_id=document_id,
            filename="real-minio.docx",
            content_type=DOCX_CONTENT_TYPE,
            content=_docx_bytes(),
        )
        try:
            result = await _ingest_claimed(
                IngestionService(session_factory, storage), session_factory, run_id
            )
            assert result.status is IngestionResultStatus.PERSISTED
            assert result.chunk_count > 0
            async with session_factory() as session:
                assert (
                    await session.scalar(
                        select(func.count(CanonicalChunk.id)).where(
                            CanonicalChunk.processing_run_id == run_id
                        )
                    )
                    == result.chunk_count
                )
        finally:
            await storage.delete(object_key)
            await engine.dispose()

    asyncio.run(exercise())
